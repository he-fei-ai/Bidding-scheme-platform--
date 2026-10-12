# -*- coding: utf-8 -*-
r"""《装饰装修专项施工方案》文档质量分析报告 —— 格式修复验证（2026-10-10）。

覆盖报告中的格式问题修复：
- 修复 3/4：Markdown 标题符号残留（含 `##` 无空格形态）
- 修复 4/13：HTML 表格 → DOCX 表格（含 colspan/rowspan、多行表格）
- 修复 5：LaTeX 公式 → 可读 Unicode 文本（降级方案）
- 修复 8：导出端防御性重复标题清理（prepare_for_export）
- 修复 10：`(1）` 括号列表归一
- 修复 11：Markdown 转义字符（`\~`）清理
"""

from types import SimpleNamespace

from docx import Document

from app.services.content_cleaning import (
    ContentCleaner,
    clean_generated_content,
    grid_to_gfm_table,
    parse_content,
    parse_html_table_grid,
)
from app.services.latex_text import convert_latex_in_text, latex_expr_to_text


# ============================================================================
# 修复 5：LaTeX 公式
# ============================================================================


class TestLatexConversion:
    def test_block_formula_report_case(self):
        """报告格式问题 7 的原始案例：\\[N = 1.2(...)\\] 不再泄漏转义命令。"""
        raw = (
            r"\[N = 1.2(N\dot{\iota}\dot{\iota} G + N_{Gk}) "
            r"+ 1.4\sum_{i = 1}^{n}N Q k\iota\]"
        )
        out = convert_latex_in_text(raw)
        assert "\\dot" not in out, f"转义命令残留: {out}"
        assert "\\sum" not in out, f"求和命令残留: {out}"
        assert "∑" in out, f"应渲染为 Unicode 求和符号: {out}"
        # 强调命令保留语义：\dot{\iota} → ι + 上点组合符（而非裸 ιι）
        assert "ι" + "\u0307" in out, f"上点语义丢失: {out}"

    def test_accent_commands_keep_semantics(self):
        """\\dot / \\ddot / \\hat / \\vec / \\tilde → 字符 + Unicode 组合附加符。"""
        out = convert_latex_in_text(r"\[\dot{t}\ddot{q}\hat{r}\vec{V}\tilde{x}\]")
        assert out == "t\u0307q\u0308r\u0302V\u20d7x\u0303", f"强调命令语义丢失: {out}"

    def test_text_wrapping_commands_strip_only(self):
        """\\text/\\mathrm/\\boxed 等纯包装命令只剥壳、不加附加符。"""
        out = convert_latex_in_text(r"\[\mathrm{V}\boxed{\Xi}\overline{a}\]")
        assert out == "VΞa", f"包装命令应只保留内容: {out}"

    def test_greek_before_plain_char_kept(self):
        """希腊字母后紧跟普通字符时不得退化为 \\Xi→Xi（命令 token 完整替换）。"""
        out = convert_latex_in_text(r"\[\boxed{\Xi}\overline{a}\]")
        assert out == "Ξa", f"希腊字母被误切分: {out}"

    def test_boxed_formula_report_case(self):
        """报告格式问题 7 的第二个案例：\\[\\boxed {\\Xi}\\leq f\\]。"""
        out = convert_latex_in_text(r"\[\boxed {\Xi}\leq f\]")
        assert "\\boxed" not in out
        assert "Ξ" in out
        assert "≤" in out

    def test_inline_dollar_math(self):
        out = convert_latex_in_text(r"由 \(a^2+b^2=c^2\) 可得")
        assert out == "由 a²+b²=c² 可得", out

    def test_dollar_amount_not_converted(self):
        """金额 `$500` 不含 LaTeX 特征，不得被当作公式转换。"""
        raw = "造价 $500 元，质保金 $100 元"
        assert convert_latex_in_text(raw) == raw

    def test_double_dollar_block(self):
        assert convert_latex_in_text(r"$$E = mc^2$$") == "E = mc²"

    def test_none_safe(self):
        assert latex_expr_to_text(None) is None
        assert latex_expr_to_text("") is None
        assert convert_latex_in_text("") == ""

    def test_clean_pipeline_applies_latex(self):
        """clean_generated_content 全链路生效，且在字面量 \\n 还原之前。"""
        raw = r"\[\alpha + \nu = \beta\]"
        out = clean_generated_content(raw)
        assert "\\alpha" not in out and "\\nu" not in out
        assert "α" in out and "ν" in out

    def test_malformed_latex_does_not_raise(self):
        """病态输入只降级不抛异常（导出不能因公式失败）。"""
        assert isinstance(convert_latex_in_text(r"\[ \frac{1}{ \sqrt{ \dot{ "), str)
        assert isinstance(convert_latex_in_text(r"\[a{b{c}d}e\]"), str)


# ============================================================================
# 修复 11：Markdown 转义字符 / 修复 3：# 残留
# ============================================================================


class TestEscapeAndHeadingCleanup:
    def test_tilde_escape_report_case(self):
        """报告格式问题 13：`·普工：8\\~10人` → `8~10人`。"""
        out = clean_generated_content(r"·普工：8\~10人，架子工：8\~12人")
        assert "\\~" not in out
        assert "8~10人" in out and "8~12人" in out

    def test_hash_heading_without_space(self):
        """报告格式问题 4：`##1.1工程概况`（无空格）也必须剥掉 #。"""
        out = clean_generated_content("##1.1工程概况\n正文内容")
        assert out.startswith("1.1工程概况"), out
        assert "#" not in out.split("\n")[0]

    def test_hash_heading_with_space(self):
        out = clean_generated_content("### 标题\n正文")
        assert out == "标题\n正文"

    def test_bold_still_stripped(self):
        assert clean_generated_content("**粗体**内容") == "粗体内容"


# ============================================================================
# 修复 4/13：HTML 表格
# ============================================================================


class TestHtmlTableParsing:
    def test_single_line_table_report_case(self):
        """报告格式问题 5：单行 <table> 不得再连成一片纯文本。"""
        html = (
            "<table><tr><td>项目要素</td><td>具体内容</td></tr>"
            "<tr><td>工程名称</td><td>上海市市级机关第二幼儿园</td></tr>"
            "<tr><td>工期安排</td><td>2026.05.08-2026.11.18</td></tr></table>"
        )
        blocks = parse_content(f"工程基本情况如下：{html}详见下文。")
        tables = [b for b in blocks if b.type == "table"]
        assert len(tables) == 1, f"应解析出 1 个 table 块: {[b.type for b in blocks]}"
        gfm = tables[0].text
        assert "工程名称" in gfm and "2026.05.08-2026.11.18" in gfm
        assert gfm.split("\n")[0].count("|") >= 3  # 真正的管道表
        body_text = "\n".join(b.text for b in blocks if b.type == "paragraph")
        assert "<table>" not in body_text and "<td>" not in body_text
        paragraphs = [b.text for b in blocks if b.type == "paragraph"]
        assert any("工程基本情况如下" in t for t in paragraphs)
        assert any("详见下文" in t for t in paragraphs)

    def test_colspan_rowspan_grid(self):
        """报告格式问题 6：合并单元格展开为均匀网格，不粘连、不错位。"""
        html = (
            "<table>"
            "<tr><td rowspan='2'>控制项目</td><td colspan='2'>限值要求</td></tr>"
            "<tr><td>指标A</td><td>指标B</td></tr>"
            "<tr><td>人造板</td><td>≤0.05</td><td>mg/m³</td></tr>"
            "</table>"
        )
        grid = parse_html_table_grid(html)
        assert len(grid) == 3
        assert all(len(row) == 3 for row in grid), grid
        assert grid[0][0] == "控制项目"
        assert grid[0][1] == "限值要求" and grid[0][2] == ""  # colspan 占位
        assert grid[1][0] == "控制项目"  # rowspan 下延
        assert grid[2] == ["人造板", "≤0.05", "mg/m³"]

    def test_multiline_table(self):
        lines = [
            "材料控制表如下：",
            "<table>",
            "<tr><th>材料</th><th>标准</th></tr>",
            "<tr><td>水泥</td><td>PO42.5</td></tr>",
            "</table>",
            "表后说明。",
        ]
        blocks = parse_content("\n".join(lines))
        types = [b.type for b in blocks]
        assert "table" in types, types
        table = next(b for b in blocks if b.type == "table")
        assert "PO42.5" in table.text
        assert any(b.type == "paragraph" and b.text == "表后说明。" for b in blocks)
        assert not any("<tr>" in b.text for b in blocks)

    def test_cell_pipe_does_not_break_columns(self):
        """单元格内的竖线必须替换，否则会把一列拆成两列。"""
        html = (
            "<table><tr><td>a|b</td><td>c</td></tr>"
            "<tr><td>d</td><td>e</td></tr></table>"
        )
        grid = parse_html_table_grid(html)
        assert grid[0][0] == "a/b"
        gfm = grid_to_gfm_table(grid)
        rows = [ln for ln in gfm.split("\n") if ln.startswith("|")]
        # 2 列表格每行应恰好 3 个竖线（首尾 + 1 个列分隔）
        assert all(ln.count("|") == 3 for ln in rows), rows
        assert "| a/b | c |" in rows

    def test_unclosed_table_does_not_leak_html(self):
        blocks = parse_content("<table>\n<tr><td>孤表</td><td>数据</td></tr>")
        joined = "\n".join(b.text for b in blocks)
        assert "<table>" not in joined
        assert "孤表" in joined

    def test_table_renders_to_docx(self):
        """HTML → GFM → 既有导出渲染链路端到端可用。"""
        from app.models import ExportConfig
        from app.services.export import _add_table_from_markup

        html = (
            "<table><tr><td>岗位</td><td>人数</td></tr>"
            "<tr><td>项目经理</td><td>1</td></tr></table>"
        )
        grid = parse_html_table_grid(html)
        gfm = grid_to_gfm_table(grid)
        doc = Document()
        _add_table_from_markup(doc, gfm, ExportConfig())
        assert len(doc.tables) == 1
        cells = [c.text for c in doc.tables[0].rows[0].cells]
        assert cells == ["岗位", "人数"], cells
        row2 = [c.text for c in doc.tables[0].rows[1].cells]
        assert row2 == ["项目经理", "1"], row2


# ============================================================================
# 修复 8：导出端防御性重复标题清理
# ============================================================================


class TestPrepareForExport:
    def test_duplicate_title_removed(self):
        cleaner = ContentCleaner()
        out = cleaner.prepare_for_export("施工部署", "施工部署\n\n本项目施工部署如下……")
        assert not out.startswith("施工部署")
        assert out.startswith("本项目施工部署如下")

    def test_numbered_duplicate_title_removed(self):
        cleaner = ContentCleaner()
        out = cleaner.prepare_for_export("施工部署", "2、施工部署\n正文")
        assert out.startswith("正文")

    def test_markdown_headings_removed(self):
        cleaner = ContentCleaner()
        out = cleaner.prepare_for_export("任意", "## 一级残留\n正文")
        assert out.startswith("一级残留")

    def test_numbered_first_line_preserved(self):
        """导出清洗不得误伤首行是合法编号列表的正文。"""
        cleaner = ContentCleaner()
        content = "1. 《建设工程质量管理条例》执行\n第二条说明。"
        assert cleaner.prepare_for_export("编制依据", content) == content

    def test_idempotent(self):
        cleaner = ContentCleaner()
        once = cleaner.prepare_for_export("T", "T\n正文")
        assert cleaner.prepare_for_export("T", once) == once

    def test_export_add_parsed_content_applies_defense(self):
        """_add_parsed_content 必须调用防御清洗（重复标题不进 DOCX）。"""
        from app.models import ExportConfig
        from app.services.export import _add_parsed_content

        doc = Document()
        detail = SimpleNamespace(id=1, title="施工部署", content="")
        content = "施工部署\n\n本项目严格按图施工，确保质量。"
        _add_parsed_content(doc, content, ExportConfig(), section_detail=detail)
        texts = [p.text for p in doc.paragraphs]
        assert any("本项目严格按图施工" in t for t in texts)
        assert not any(t == "施工部署" for t in texts), f"重复标题未清洗: {texts}"


# ============================================================================
# 修复 10：括号列表归一
# ============================================================================


class TestBracketList:
    def test_mixed_bracket_list_detected(self):
        """报告格式问题 12：`(1）外墙铲除…` / `（2）外墙抹灰…` → 有序列表。"""
        blocks = parse_content("(1）外墙铲除与基层处理。\n（2）外墙抹灰修复。")
        assert [b.type for b in blocks] == ["list_ordered", "list_ordered"], [
            b.type for b in blocks
        ]
        assert blocks[0].text == "外墙铲除与基层处理。"
        assert blocks[1].text == "外墙抹灰修复。"

    def test_year_not_detected_as_list(self):
        """`(2026)年…` 这类年份开头正文不得误判为列表。"""
        blocks = parse_content("(2026)年度检查情况说明如下。")
        assert blocks and blocks[0].type == "paragraph"


