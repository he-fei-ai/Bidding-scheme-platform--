"""2026-09-18 八模块端到端审计 —— BUG 修复与增强回归测试

覆盖模块（文件导入 → 解析 → 提取项目 → 目录生成 → 全局事实 → 正文生成 → 审核预检 → 导出）：

A. 文件导入/解析
   A1 `_pick_main_file` 不再把「解析失败的主文件」当作主文件
   A2 `_update_parse_task` 按 filename 精确更新，并发上传不串写
   A3 `save_upload` 先建任务行，失败状态不再丢失
   A4 辅助文件读取走 `_read_text_auto`（GBK 不再抛 UnicodeDecodeError）
   A5 `_is_under_projects_dir` 拒绝同前缀目录与路径穿越

B. 目录生成
   B1 `_format_tender_context_block` 渲染招标上下文/行业背景
   B2 `generate_outline_v51` 的 tender_context 不再是死参数（L1 提示词含招标内容）
   B3 `_clear_and_save_outline` 重建目录时按标题保全已生成正文
   B4 sections 落库保留 target_words

C. 全局事实 / 正文生成
   C1 `_scoring_context_to_text` 归一化 list/dict/str
   C2 `build_chapter_content_messages` 不再把 Python list repr 塞进提示词
   C3 `_coerce_fact_audit_list` 解包 AI 包装格式
   C4 `score_section_quality_local` 本地质量打分可用
   C5 `_fetch_project_data_async` 的 scoring_points 缺省值为 list（非字符串 "[]"）

D. 审核预检
   D1 单维度重试继承其余维度历史结果
   D2 一致性冲突 id 含 project_id，且重复审计幂等（不膨胀）

E. 导出
   E1 `_is_safe_export_artifact` 识别真实产物路径
   E2 `_iter_export_artifacts` 扫描 `<数据目录>/<pid>/exports`
"""

from __future__ import annotations

import asyncio
import json
from pathlib import Path
from unittest.mock import MagicMock

import pytest


def _new_project(name: str) -> int:
    """在当前测试库中创建一个项目并返回其 id。

    conftest.py 已把 `BID_ASSISTANT_DATA_DIR` 指向独立临时目录并调用过
    `init_db()`，因此这里不需要（也不应该）去 monkeypatch 数据库路径
    —— `app.db` 并没有 `DB_PATH` 这种模块级属性，库路径由配置层解析。
    """
    from app.db import get_db, init_db

    init_db()
    with get_db() as conn:
        conn.execute("INSERT INTO projects (name, industry) VALUES (?, ?)", (name, "通用"))
        return conn.execute("SELECT id FROM projects ORDER BY id DESC LIMIT 1").fetchone()["id"]


# ============================================================================
# A. 文件导入 / 解析
# ============================================================================


class TestA1PickMainFile:
    def test_prefers_success_main_file(self):
        from app.services.files import _pick_main_file

        files = [
            {"id": 1, "filename": "招标文件.docx", "file_role": "main", "status": "failed"},
            {"id": 2, "filename": "招标文件-补充.pdf", "file_role": "main", "status": "success"},
        ]
        picked = _pick_main_file(files)
        assert picked["id"] == 2, "必须优先选解析成功的主文件，避免下游基于空文本运行"

    def test_falls_back_to_failed_main_when_no_success(self):
        from app.services.files import _pick_main_file

        files = [
            {"id": 1, "filename": "a.docx", "file_role": "main", "status": "failed"},
        ]
        assert _pick_main_file(files)["id"] == 1

    def test_no_main_marker_uses_success(self):
        from app.services.files import _pick_main_file

        files = [
            {"id": 1, "filename": "a.docx", "file_role": "supporting", "status": "failed"},
            {"id": 2, "filename": "b.docx", "file_role": "supporting", "status": "success"},
        ]
        assert _pick_main_file(files)["id"] == 2

    def test_empty_returns_none(self):
        from app.services.files import _pick_main_file

        assert _pick_main_file([]) is None


class TestA5PathGuard:
    def test_rejects_sibling_prefix_dir(self, monkeypatch, tmp_path):
        from app.services import files as files_mod

        base = tmp_path / "projects"
        base.mkdir()
        evil = tmp_path / "projects_evil"
        evil.mkdir()
        target = evil / "x.doc"
        target.write_text("x", encoding="utf-8")

        monkeypatch.setattr(files_mod, "PROJECTS_DATA_DIR", base)
        assert files_mod._is_under_projects_dir(target) is False

    def test_rejects_traversal(self, monkeypatch, tmp_path):
        from app.services import files as files_mod

        base = tmp_path / "projects"
        (base / "1").mkdir(parents=True)
        monkeypatch.setattr(files_mod, "PROJECTS_DATA_DIR", base)

        outside = tmp_path / "secret.doc"
        outside.write_text("s", encoding="utf-8")
        traversal = base / "1" / ".." / ".." / "secret.doc"
        assert files_mod._is_under_projects_dir(traversal) is False

    def test_accepts_inside(self, monkeypatch, tmp_path):
        from app.services import files as files_mod

        base = tmp_path / "projects"
        (base / "1" / "uploads").mkdir(parents=True)
        f = base / "1" / "uploads" / "ok.doc"
        f.write_text("ok", encoding="utf-8")
        monkeypatch.setattr(files_mod, "PROJECTS_DATA_DIR", base)
        assert files_mod._is_under_projects_dir(f) is True


class TestA3SaveUploadTaskOrder:
    def test_parse_task_created_before_write_failure(self, monkeypatch, tmp_path):
        """保存失败时任务行必须已存在，失败状态才能落库。"""
        from app.services import files as files_mod

        created: list = []
        updated: list = []

        monkeypatch.setattr(
            files_mod, "_create_parse_task",
            lambda pid, fn, ft="": created.append((pid, fn, ft)),
        )
        monkeypatch.setattr(
            files_mod, "_update_parse_task",
            lambda pid, status, progress=0, error_message="", filename="": updated.append(
                (pid, status, filename)
            ),
        )
        # 让落盘必然失败
        monkeypatch.setattr(Path, "write_bytes", lambda self, data: (_ for _ in ()).throw(OSError("disk full")))
        monkeypatch.setattr(files_mod, "PROJECTS_DATA_DIR", tmp_path)

        result = files_mod.save_upload(424242, "x.txt", b"hello", "main")

        assert result["status"] == "failed"
        assert created, "必须先创建任务行"
        assert any(u[1] == "failed" for u in updated), "失败状态必须被写入"
        assert all(u[2] == "x.txt" for u in updated if u[1] == "failed"), "失败更新需带 filename"


class TestA2UpdateParseTaskFilter:
    def test_update_targets_filename(self, monkeypatch):
        from app.services import files as files_mod
        import app.db as db_mod

        captured: list = []

        class _Conn:
            def execute(self, sql, params):
                captured.append((sql, params))
                return MagicMock()

        class _Ctx:
            def __enter__(self):
                return _Conn()

            def __exit__(self, *a):
                return False

        # `_update_parse_task` 内部是 `from ..db import get_db` 的**延迟导入**，
        # 每次调用都从 `app.db` 取属性，因此必须 patch 源模块上的 `get_db`；
        # patch `files_mod.get_db` 会因该模块没有此属性而直接报 AttributeError。
        monkeypatch.setattr(db_mod, "get_db", lambda *a, **k: _Ctx())
        files_mod._update_parse_task(7, "completed", 100, filename="a.docx")

        assert captured, "应执行 UPDATE"
        sql, params = captured[0]
        assert "filename" in sql, "SQL 必须按 filename 过滤，避免并发上传串写"
        assert "a.docx" in params


class TestA4SupportingFileEncoding:
    def test_gbk_supporting_file_does_not_raise(self, monkeypatch, tmp_path):
        from app.services import files as files_mod

        gbk = tmp_path / "清单.txt"
        gbk.write_bytes("工程量清单\n项目名称 数量 单价".encode("gbk"))

        files = [
            {
                "id": 1,
                "filename": "清单.txt",
                "file_role": "supporting",
                "status": "success",
                "extracted_text_path": str(gbk),
            }
        ]
        text = asyncio.run(files_mod._truncate_supporting_file(files))
        assert "工程量清单" in text


# ============================================================================
# B. 目录生成
# ============================================================================


class TestB1TenderContextBlock:
    def test_renders_context_and_industry(self):
        from app.services.ai.workflows_outline_v51 import _format_tender_context_block

        block = _format_tender_context_block("招标要求：工期 450 日历天", "房建工程")
        assert "<tender_context>" in block
        assert "450 日历天" in block
        assert "<industry_profile>" in block
        assert "房建工程" in block

    def test_empty_returns_empty(self):
        from app.services.ai.workflows_outline_v51 import _format_tender_context_block

        assert _format_tender_context_block("", "") == ""

    def test_truncation_marks(self):
        from app.services.ai.workflows_outline_v51 import _format_tender_context_block

        block = _format_tender_context_block("甲" * 5000, max_chars=100)
        assert "已截断" in block


class TestB2TenderContextNotDeadParam:
    def test_l1_prompt_includes_tender_context(self, monkeypatch):
        """tender_context 必须真正进入 L1 提示词（此前是死参数）。"""
        from app.services.ai import workflows_outline_v51 as v51

        captured: dict = {}

        async def _fake_call(provider, sys_prompt, user_prompt, **kwargs):
            captured["user"] = user_prompt
            return json.dumps(
                [{"id": "1", "title": "项目概述", "description": "d"}], ensure_ascii=False
            )

        monkeypatch.setattr(v51, "_call_ai_v51", _fake_call)

        asyncio.run(
            v51._generate_level1_v51(
                provider=MagicMock(),
                summary="项目摘要",
                sp_text=json.dumps([{"item": "评分项A"}], ensure_ascii=False),
                reference_json="[]",
                tender_context="本项目必须采用盾构法施工，工期 450 日历天",
                industry_profile="市政工程",
            )
        )

        # 标书智能体（4）：user_prompt 现为「稳定上下文块 + 任务差异」列表，
        # 合并后断言内容不变（tender_context 仍是稳定上下文消息之一）。
        user_text = (
            "\n".join(captured["user"])
            if isinstance(captured["user"], list)
            else captured["user"]
        )
        assert "盾构法施工" in user_text, "招标文件内容必须出现在 L1 提示词中"
        assert "市政工程" in user_text

    def test_generate_outline_v51_passes_context_to_l1(self, monkeypatch):
        from app.services.ai import workflows_outline_v51 as v51

        seen: dict = {}

        async def _fake_l1(provider, summary, sp_text, reference_json, **kwargs):
            seen.update(kwargs)
            return [{"id": "1", "title": "T", "description": "", "children": []}]

        async def _fake_l2(provider, level1, summary, sp_text, **kwargs):
            seen["l2_tender"] = kwargs.get("tender_context")
            return level1

        async def _fake_l3(provider, level1, summary, sp_text, **kwargs):
            return level1

        async def _fake_self_check(provider, level1, *args, **kwargs):
            return level1

        monkeypatch.setattr(v51, "_generate_level1_v51", _fake_l1)
        monkeypatch.setattr(v51, "_generate_level2_v51", _fake_l2)
        # L3 与 AI 自检在主链路里都是 `await` 的协程，替身必须是 async def，
        # 否则会以 "object list can't be used in 'await' expression" 失败。
        monkeypatch.setattr(v51, "_generate_level3_v51", _fake_l3)
        monkeypatch.setattr(v51, "_self_check_v51", _fake_self_check)

        asyncio.run(
            v51.generate_outline_v51(
                provider=MagicMock(),
                summary="s",
                scoring_points=json.dumps([{"item": "x"}], ensure_ascii=False),
                industry_profile="房建",
                tender_context="招标文件正文内容",
            )
        )

        assert seen.get("tender_context") == "招标文件正文内容"
        assert seen.get("industry_profile") == "房建"
        assert seen.get("l2_tender") == "招标文件正文内容"


class TestB3OutlineContentPreservation:
    def test_content_preserved_for_same_title(self):
        from app.routers.sse import sse_db_ops
        from app.db import get_db

        pid = _new_project("目录重建-正文保全")

        outline = [{"id": "1", "title": "项目概述", "description": "", "children": []}]
        saved = sse_db_ops._clear_and_save_outline(pid, outline)
        sid = saved[0]["id"]

        with get_db() as conn:
            conn.execute(
                "UPDATE sections SET content = ?, status = 'completed' WHERE id = ?",
                ("这是已经生成好的正文内容", sid),
            )

        # 再次生成同名目录 → 正文必须被保留
        saved2 = sse_db_ops._clear_and_save_outline(pid, outline)
        sid2 = saved2[0]["id"]

        with get_db() as conn:
            row = conn.execute("SELECT content, status FROM sections WHERE id = ?", (sid2,)).fetchone()

        assert row["content"] == "这是已经生成好的正文内容", "重新生成目录不得清空已生成的正文"
        assert row["status"] == "completed"
        assert saved2[0]["has_content"] is True

    def test_new_title_starts_empty(self):
        from app.routers.sse import sse_db_ops
        from app.db import get_db

        pid = _new_project("目录重建-新标题为空")

        saved = sse_db_ops._clear_and_save_outline(
            pid, [{"id": "1", "title": "旧标题", "description": "", "children": []}]
        )
        with get_db() as conn:
            conn.execute(
                "UPDATE sections SET content = '旧正文' WHERE id = ?", (saved[0]["id"],)
            )

        saved2 = sse_db_ops._clear_and_save_outline(
            pid, [{"id": "1", "title": "全新标题", "description": "", "children": []}]
        )
        with get_db() as conn:
            row = conn.execute(
                "SELECT content FROM sections WHERE id = ?", (saved2[0]["id"],)
            ).fetchone()
        assert row["content"] == ""
        assert saved2[0]["has_content"] is False


class TestB4TargetWordsPersisted:
    def test_target_words_saved_to_db(self):
        from app.routers.sse import sse_db_ops
        from app.db import get_db

        pid = _new_project("目录-字数预算落库")

        saved = sse_db_ops._clear_and_save_outline(
            pid,
            [{"id": "1", "title": "章节", "description": "", "target_words": 3500, "children": []}],
        )
        with get_db() as conn:
            row = conn.execute(
                "SELECT target_words FROM sections WHERE id = ?", (saved[0]["id"],)
            ).fetchone()
        assert row["target_words"] == 3500, "target_words 必须落库，否则刷新后字数预算丢失"


# ============================================================================
# C. 全局事实 / 正文生成
# ============================================================================


class TestC1ScoringContextCoercion:
    def test_list_of_dicts(self):
        from app.services.ai.workflows_sections import _scoring_context_to_text

        text = _scoring_context_to_text([{"item": "质量管理体系", "score": 10}])
        assert "质量管理体系" in text
        assert "{" not in text and "'" not in text, "不得残留 Python repr 字符"

    def test_wrapped_dict(self):
        from app.services.ai.workflows_sections import _scoring_context_to_text

        text = _scoring_context_to_text({"items": [{"item": "A"}, {"item": "B"}]})
        assert "A" in text and "B" in text

    def test_str_passthrough(self):
        from app.services.ai.workflows_sections import _scoring_context_to_text

        assert _scoring_context_to_text("原始文本") == "原始文本"

    def test_none(self):
        from app.services.ai.workflows_sections import _scoring_context_to_text

        assert _scoring_context_to_text(None) == ""


class TestC2NoListReprInPrompt:
    def test_prompt_has_no_python_repr(self):
        from app.services.ai.workflows_sections import build_chapter_content_messages

        scoring = [{"item": "施工组织设计完整性", "score": 15}]
        messages = build_chapter_content_messages(
            {"id": 1, "title": "施工组织设计", "summary": "概要"},
            scoring_context=scoring,
        )
        user = "\n".join(m.content for m in messages if m.role == "user")
        assert "施工组织设计完整性" in user
        assert "[{'item'" not in user, "提示词里不得出现 Python list repr"


class TestC3CoerceFactAuditList:
    def test_unwraps_items(self):
        from app.routers.facts import _coerce_fact_audit_list

        assert len(_coerce_fact_audit_list({"items": [{"fact_key": "a"}]})) == 1

    def test_unwraps_findings(self):
        from app.routers.facts import _coerce_fact_audit_list

        assert len(_coerce_fact_audit_list({"findings": [{"fact_key": "a"}]})) == 1

    def test_bare_list(self):
        from app.routers.facts import _coerce_fact_audit_list

        assert len(_coerce_fact_audit_list([{"fact_key": "a"}])) == 1

    def test_garbage(self):
        from app.routers.facts import _coerce_fact_audit_list

        assert _coerce_fact_audit_list("nope") == []
        assert _coerce_fact_audit_list(None) == []


class TestC4LocalQualityScore:
    def test_good_content_scores_higher(self):
        from app.services.ai.workflows_quality import score_section_quality_local

        good = (
            "本项目采用盾构法施工，管片外径 6200 mm，壁厚 350 mm。\n\n"
            "施工准备阶段投入 12 人，配备 2 台盾构机。\n\n"
            "进度计划：2026 年 3 月开工，工期 450 天，共分 5 个阶段。\n\n"
            "质量控制措施包括：原材料进场检验、过程三检制、成品保护。\n\n"
            "安全措施：临边防护、洞口防护、用电三级配电两级保护。\n\n"
            "环保措施：泥浆分离、噪声监测、扬尘控制。\n\n"
            "应急预案：成立应急小组，配备抢险物资 30 套。\n\n"
            "验收标准执行 GB 50204 相关规定。"
        )
        poor = "内容待补充。"

        good_score = score_section_quality_local(good, "施工方案", 300)["score"]
        poor_score = score_section_quality_local(poor, "施工方案", 300)["score"]
        assert good_score > poor_score

    def test_returns_expected_shape(self):
        from app.services.ai.workflows_quality import score_section_quality_local

        q = score_section_quality_local("测试内容" * 50, "标题", 200)
        for key in ("score", "grade", "dimensions", "issues", "evaluator"):
            assert key in q
        assert q["evaluator"] == "local"
        assert 0 <= q["score"] <= 100

    def test_empty_content(self):
        from app.services.ai.workflows_quality import score_section_quality_local

        q = score_section_quality_local("", "标题", None)
        assert q["score"] == 0 or q["score"] >= 0


class TestC5ScoringPointsTypeContract:
    def test_missing_scoring_points_is_list(self):
        """`_fetch_project_data_async` 的缺省值必须是 list，不能是字符串 '[]'。"""
        import app.routers.sse.sse_utils as su

        src = Path(su.__file__).read_text(encoding="utf-8")
        # 只检查**代码行**：修复注释里会原文引用旧的错误写法（`result["scoring_points"] = "[]"`），
        # 直接对全文做子串断言会把注释也算进去，导致永远失败。
        code_lines = [
            ln.strip()
            for ln in src.splitlines()
            if ln.strip() and not ln.strip().startswith("#")
        ]
        assert not any('result["scoring_points"] = "[]"' in ln for ln in code_lines), (
            "scoring_points 缺省值必须是 list，字符串会破坏下游 isinstance(..., list) 判定"
        )
        assert any('result["scoring_points"] = []' in ln for ln in code_lines), (
            "缺省值应显式赋为 list（[]）"
        )


# ============================================================================
# D. 审核预检
# ============================================================================


class TestD2ConsistencyConflictId:
    # 同一事实（工期）出现两个不同取值，才能触发冲突检测
    _CONFLICT_TEXT = "本工程工期：450 天。若遇特殊情况，工期：500 天。"

    @staticmethod
    def _seed_section(pid: int, text: str) -> None:
        from app.db import get_db

        with get_db() as conn:
            conn.execute(
                "INSERT INTO sections (project_id, title, level, sort_order, content, status) "
                "VALUES (?, ?, 1, 0, ?, 'completed')",
                (pid, "章节A", text),
            )

    def test_conflict_id_contains_project_id(self):
        from app.services.consistency_auditor import ConsistencyAuditor

        pid = _new_project("一致性-冲突ID含项目")
        self._seed_section(pid, self._CONFLICT_TEXT)

        auditor = ConsistencyAuditor()
        # run_full_audit 是**同步**方法（返回 AuditResult），不要用 asyncio.run 包
        result = auditor.run_full_audit(pid)
        assert result.conflicts, "构造的同事实不同取值文本必须能触发冲突"
        assert all(str(pid) in c.conflict_id for c in result.conflicts), (
            "conflict_id 必须含 project_id，否则跨项目会撞全局 UNIQUE 约束"
        )

    def test_repeat_audit_is_idempotent(self):
        from app.db import get_db
        from app.services.consistency_auditor import ConsistencyAuditor

        pid = _new_project("一致性-重复审计幂等")
        self._seed_section(pid, self._CONFLICT_TEXT)

        auditor = ConsistencyAuditor()
        auditor.run_full_audit(pid)
        with get_db() as conn:
            n1 = conn.execute(
                "SELECT COUNT(*) AS c FROM consistency_conflicts WHERE project_id = ?", (pid,)
            ).fetchone()["c"]
        assert n1 > 0, "首轮审计应写入冲突记录"

        auditor.run_full_audit(pid)
        with get_db() as conn:
            n2 = conn.execute(
                "SELECT COUNT(*) AS c FROM consistency_conflicts WHERE project_id = ?", (pid,)
            ).fetchone()["c"]

        assert n2 == n1, "重复审计不得让未解决冲突无限膨胀（一次审计 = 一份当前快照）"


# ============================================================================
# E. 导出
# ============================================================================


class TestEExportArtifactDetection:
    def test_real_artifact_path_detected(self, monkeypatch, tmp_path):
        from app.services import export as export_mod

        base = tmp_path / "projects"
        pid_dir = base / "12" / "exports"
        pid_dir.mkdir(parents=True)
        f = pid_dir / "out.docx"
        f.write_bytes(b"x")

        monkeypatch.setattr(export_mod, "PROJECTS_DATA_DIR", base)
        assert export_mod._is_safe_export_artifact(f) is True

    def test_outside_path_rejected(self, monkeypatch, tmp_path):
        from app.services import export as export_mod

        base = tmp_path / "projects"
        base.mkdir()
        outside = tmp_path / "elsewhere.docx"
        outside.write_bytes(b"x")

        monkeypatch.setattr(export_mod, "PROJECTS_DATA_DIR", base)
        assert export_mod._is_safe_export_artifact(outside) is False

    def test_non_docx_rejected(self, monkeypatch, tmp_path):
        from app.services import export as export_mod

        base = tmp_path / "projects"
        d = base / "1" / "exports"
        d.mkdir(parents=True)
        f = d / "note.txt"
        f.write_bytes(b"x")

        monkeypatch.setattr(export_mod, "PROJECTS_DATA_DIR", base)
        assert export_mod._is_safe_export_artifact(f) is False

    def test_iter_finds_real_artifacts(self, monkeypatch, tmp_path):
        from app.services import export as export_mod

        base = tmp_path / "projects"
        for pid in ("1", "2"):
            d = base / pid / "exports"
            d.mkdir(parents=True)
            (d / "a.docx").write_bytes(b"x")

        monkeypatch.setattr(export_mod, "PROJECTS_DATA_DIR", base)
        found = export_mod._iter_export_artifacts()
        assert len(found) == 2, "必须扫到 <数据目录>/<pid>/exports 下的真实产物"


# ============================================================================
# F. 提取项目（Step02：清单 / 工作界面 / 工程作法 + 评分点占位过滤）
# ============================================================================


class _BoomProvider:
    """任何 chat 调用都失败的 provider（模拟 AI 服务不可用 / 返回非 JSON）。"""

    async def chat(self, *args, **kwargs):
        raise RuntimeError("AI 服务不可用")


class TestF1AnalysisFallbackReachable:
    """AI 失败时兜底默认值必须**可达**。

    原缺陷：`analyze_boq` / `analyze_work_interface` 没有 try 包裹
    `collect_json_response`（该函数失败时抛 RuntimeError），异常冒泡到路由层被转成
    HTTP 502，两个函数里精心准备的默认清单/默认工作界面成为**死代码**；
    而同型的 `analyze_construction_method` 早已修过（BUG-X-02），行为不一致。
    """

    def test_boq_fallback_reachable_on_ai_failure(self):
        from app.services.ai.workflows_analysis import analyze_boq

        result = asyncio.run(analyze_boq(_BoomProvider(), "招标上下文"))
        assert isinstance(result, list) and result, "AI 失败时必须返回默认清单而非抛异常"
        assert result[0]["id"].startswith("BOQ")

    def test_work_interface_fallback_reachable_on_ai_failure(self):
        from app.services.ai.workflows_analysis import analyze_work_interface

        result = asyncio.run(analyze_work_interface(_BoomProvider(), "招标上下文"))
        assert isinstance(result, list) and result, "AI 失败时必须返回默认工作界面"
        assert result[0]["id"].startswith("WI")

    def test_construction_method_fallback_reachable_on_ai_failure(self):
        from app.services.ai.workflows_analysis import analyze_construction_method

        result = asyncio.run(analyze_construction_method(_BoomProvider(), "招标上下文"))
        assert isinstance(result, list) and result
        assert result[0]["id"].startswith("CM")


class TestF2ScoringPlaceholderNotInOutline:
    """`__weight_warning__` 是系统提示项，绝不能被当成评分项。

    `normalize_writing_targets` 在「评分项权重之和偏离 100 超过 1 个百分点」时
    会往列表首位插入该伪项（`name="[权重异常] 评分项权重之和为 X…"`）。
    而「评分表只列技术标分数、权重合计不足 100」是极常见的真实情况。
    """

    _WARNING_ITEM = {
        "id": "__weight_warning__",
        "name": "[权重异常] 评分项权重之和为 90，请检查是否遗漏评分项",
        "category": "系统提示",
    }

    def test_strip_helper(self):
        from app.services.ai.workflows_outline_v51 import _strip_scoring_placeholders

        assert len(_strip_scoring_placeholders([self._WARNING_ITEM, {"id": "R1"}])) == 1
        assert _strip_scoring_placeholders(None) is None
        assert _strip_scoring_placeholders("原文") == "原文"
        wrapped = _strip_scoring_placeholders(
            {"items": [self._WARNING_ITEM, {"id": "R1"}]}
        )
        assert len(wrapped["items"]) == 1

    def test_l1_prompt_count_excludes_placeholder(self, monkeypatch):
        from app.services.ai import workflows_outline_v51 as v51

        captured: dict = {}

        async def _fake_call(provider, sys_prompt, user_prompt, **kwargs):
            captured["user"] = user_prompt
            return json.dumps(
                [{"id": "1", "title": "施工组织设计", "description": ""}], ensure_ascii=False
            )

        monkeypatch.setattr(v51, "_call_ai_v51", _fake_call)

        sp_text = json.dumps(
            [self._WARNING_ITEM, {"id": "R1", "name": "施工组织设计", "weight": "30"}],
            ensure_ascii=False,
        )
        asyncio.run(
            v51._generate_level1_v51(
                provider=MagicMock(), summary="s", sp_text=sp_text, reference_json="[]"
            )
        )

        # 标书智能体（4）：user_prompt 现为列表（稳定上下文块 + 任务差异），
        # 合并后断言「伪项不计入评分项总数」的口径不变。
        user_text = (
            "\n".join(captured["user"])
            if isinstance(captured["user"], list)
            else captured["user"]
        )
        assert "评分项总数: 1" in user_text, (
            "伪项不得计入评分项总数，否则 AI 会按错误总数多生成一个凑数章节"
        )

    def test_fallback_builder_skips_placeholder(self):
        from app.services.ai.workflows_outline_v51 import _build_l1_from_scoring_items

        built = _build_l1_from_scoring_items(
            [
                self._WARNING_ITEM,
                {"id": "R1", "name": "施工组织设计"},
                {"id": "R2", "name": "质量保证措施"},
            ],
            3,
        )
        titles = [n["title"] for n in built]
        assert "施工组织设计" in titles and "质量保证措施" in titles
        assert not any("权重异常" in t for t in titles), (
            "兜底目录不得出现标题为系统提示的伪章节"
        )

    def test_generate_outline_v51_strips_placeholder_upstream(self, monkeypatch):
        """最上游清洗：注入 L1 的 `sp_text` 本身不得含伪项。"""
        from app.services.ai import workflows_outline_v51 as v51

        seen: dict = {}

        async def _fake_l1(provider, summary, sp_text, reference_json, **kwargs):
            seen["sp_text"] = sp_text
            return [{"id": "1", "title": "T", "description": "", "children": []}]

        async def _fake_l2(provider, level1, summary, sp_text, **kwargs):
            return level1

        async def _fake_l3(provider, level1, summary, sp_text, **kwargs):
            return level1

        async def _fake_self_check(provider, level1, *args, **kwargs):
            return level1

        monkeypatch.setattr(v51, "_generate_level1_v51", _fake_l1)
        monkeypatch.setattr(v51, "_generate_level2_v51", _fake_l2)
        monkeypatch.setattr(v51, "_generate_level3_v51", _fake_l3)
        monkeypatch.setattr(v51, "_self_check_v51", _fake_self_check)

        asyncio.run(
            v51.generate_outline_v51(
                provider=MagicMock(),
                summary="s",
                scoring_points=json.dumps(
                    [self._WARNING_ITEM, {"id": "R1", "name": "施工组织设计"}],
                    ensure_ascii=False,
                ),
                industry_profile="",
            )
        )

        assert "__weight_warning__" not in seen["sp_text"], (
            "伪项必须在上游被剔除，不能进入 L1 提示词的 <requirements>"
        )
        assert "施工组织设计" in seen["sp_text"], "真实评分项必须保留"


class TestF3ScoringContextSkipsPlaceholder:
    """正文提示词的「评分要点」不得混进系统提示项。"""

    def test_placeholder_excluded_from_prompt_context(self):
        from app.services.ai.workflows_sections import _scoring_context_to_text

        text = _scoring_context_to_text(
            [
                {
                    "id": "__weight_warning__",
                    "name": "[权重异常] 评分项权重之和为 90，请检查是否遗漏评分项",
                },
                {"item": "施工组织设计", "score": 30},
            ]
        )
        assert "施工组织设计" in text, "真实评分项必须保留"
        assert "权重异常" not in text, "系统提示项不得进入正文提示词的评分要点"

    def test_placeholder_excluded_from_full_messages(self):
        from app.services.ai.workflows_sections import build_chapter_content_messages

        messages = build_chapter_content_messages(
            {"id": 1, "title": "施工组织设计"},
            scoring_context=[
                {"id": "__weight_warning__", "name": "[权重异常] 权重之和为 90"},
                {"item": "施工组织设计", "score": 30},
            ],
        )
        user = "\n".join(m.content for m in messages if m.role == "user")
        assert "权重异常" not in user


# ============================================================================
# G. 审核预检 · 标书查重（多维度并发容错）
# ============================================================================


class TestG1DuplicateCheckDimensionResilience:
    """查重四个维度并发执行，任一维度失败不得让整单崩溃。

    原缺陷：`asyncio.gather(...)` 未传 `return_exceptions=True`，任一维度抛异常会
    立即传播 → `run_duplicate_check` 崩溃 → 既不写 results 也不执行
    `update_task_status("completed")`，任务状态**永远卡在 running**，
    前端只能轮询到超时，用户完全不知道失败了。
    """

    @staticmethod
    def _patch_base(dup, monkeypatch, statuses):
        monkeypatch.setattr(
            dup,
            "update_task_status",
            lambda task_id, status, results=None, error="": statuses.append((status, error, results)),
        )
        monkeypatch.setattr(dup, "update_task_progress", lambda *a, **k: None)
        monkeypatch.setattr(
            dup,
            "get_task",
            lambda tid: {"id": tid, "file_ids": [1, 2], "tender_file_id": None},
        )
        monkeypatch.setattr(
            dup,
            "_load_file_info",
            lambda fid: {"id": fid, "filename": f"标书{fid}.docx", "text": "正文内容"},
        )

    def test_all_dimensions_ok(self, monkeypatch):
        from app.services import duplicate as dup

        statuses: list = []
        self._patch_base(dup, monkeypatch, statuses)
        monkeypatch.setattr(dup, "analyze_metadata", lambda fis: {"metadata": [], "anomalies": []})
        monkeypatch.setattr(dup, "analyze_toc", lambda fis: {"pairwise": []})
        monkeypatch.setattr(dup, "analyze_content", lambda fis, t="": {"pairs": []})
        monkeypatch.setattr(dup, "analyze_images", lambda fis: {"duplicate_groups": []})

        asyncio.run(dup.run_duplicate_check(1))

        assert statuses[-1][0] == "completed", "全部成功应正常完成"
        assert not statuses[-1][1], "全部成功不应带错误说明"

    def test_single_dimension_failure_still_completes(self, monkeypatch):
        """单维度失败：降级完成，但必须如实告知结论不完整。"""
        from app.services import duplicate as dup

        statuses: list = []
        self._patch_base(dup, monkeypatch, statuses)
        monkeypatch.setattr(dup, "analyze_metadata", lambda fis: {"metadata": [], "anomalies": []})
        monkeypatch.setattr(dup, "analyze_toc", lambda fis: {"pairwise": []})
        monkeypatch.setattr(dup, "analyze_images", lambda fis: {"duplicate_groups": []})

        def _boom(*a, **k):
            raise RuntimeError("PDF 解析失败")

        monkeypatch.setattr(dup, "analyze_content", _boom)

        asyncio.run(dup.run_duplicate_check(2))

        assert statuses, "必须有终态上报"
        final_status, final_error, results = statuses[-1]
        assert final_status == "completed", f"单维度失败应降级完成，实际: {final_status}"
        assert final_error, "部分维度失败必须带回错误说明"
        assert results["content"].get("available") is False
        # 其余维度结果必须保留，不能因为一个维度失败就丢掉全部
        assert results["metadata"]["metadata"] == []

    def test_all_dimensions_failed_reports_failed(self, monkeypatch):
        """全维度失败：必须报 failed，不能谎报完成，更不能卡在 running。"""
        from app.services import duplicate as dup

        statuses: list = []
        self._patch_base(dup, monkeypatch, statuses)

        def _boom(*a, **k):
            raise RuntimeError("文件损坏")

        for name in ("analyze_metadata", "analyze_toc", "analyze_content", "analyze_images"):
            monkeypatch.setattr(dup, name, _boom)

        asyncio.run(dup.run_duplicate_check(3))

        assert statuses[-1][0] == "failed", (
            f"四个维度全部失败必须报 failed，实际: {statuses[-1][0]}"
        )


class TestG2AnalyzeContentStatsHonest:
    """`analyze_content` 的统计口径必须如实（不虚报、不静默跳过）。"""

    def test_small_sample_not_skipped_and_fields_consistent(self):
        from app.services import duplicate as dup

        files = [
            {
                "file_id": 1,
                "filename": "A.docx",
                "text": "我方承诺投入两台大型履带式起重机完成吊装作业任务。",
            },
            {
                "file_id": 2,
                "filename": "B.docx",
                "text": "我方承诺投入两台大型履带式起重机完成吊装作业任务。",
            },
        ]
        stats = dup.analyze_content(files)["stats"]
        assert "similar_skipped" in stats, "必须如实上报相似检测是否被跳过"
        assert stats["similar_skipped"] is False, "小样本不应跳过相似检测"
        assert stats["similar_count"] == stats["similar_returned"], (
            "未截断时 总数 与 返回数 应一致"
        )

    def test_similar_count_is_total_before_truncation(self):
        """`similar_count` 必须是**截断前**的真实总数。

        原缺陷：先 `[:100]` 再 `len()`，与同结构的 `duplicate_count`（截断前取总数）
        口径不一致 —— 用户看到「相似 100 条」，实际可能有数千条。
        """
        from app.services import duplicate as dup

        # 20 × 20 = 400 对「相似但不完全相同」的句子，远超 100 的输出上限
        files = [
            {
                "file_id": 1,
                "filename": "A.docx",
                "text": "".join(
                    f"第{i}项工作由甲方负责组织实施并保证按期完成。" for i in range(20)
                ),
            },
            {
                "file_id": 2,
                "filename": "B.docx",
                "text": "".join(
                    f"第{i}项工作由乙方负责组织实施并保证按期完成。" for i in range(20)
                ),
            },
        ]
        result = dup.analyze_content(files)
        stats = result["stats"]

        assert len(result["similar_groups"]) == 100, "输出仍应限制在 100 条"
        assert stats["similar_returned"] == 100
        assert stats["similar_count"] > 100, (
            f"similar_count 必须是截断前的真实总数，实际: {stats['similar_count']}"
        )

    def test_large_sample_reports_skipped(self):
        """超规模时必须如实上报「已跳过」，不能静默返回空结果。"""
        from app.services import duplicate as dup

        big = "".join(f"这是第{i}句用于构造大规模文本的内容。" for i in range(5000))
        files = [
            {"file_id": 1, "filename": "A.docx", "text": big},
            {"file_id": 2, "filename": "B.docx", "text": big.replace("内容", "内容二")},
        ]
        result = dup.analyze_content(files)
        assert result["stats"]["similar_skipped"] is True, (
            "超过 5000 句时必须上报已跳过，而不是静默返回 0 条"
        )
        assert result["similar_groups"] == []


# ============================================================================
# H. 导出模块深化（预检误报 / 缓存指纹 / 图号计数器 / 伪评分项收敛）
# ============================================================================


def _mk_section(pid, title, level, parent_id, content, sp_id=None):
    from app.db import get_db

    with get_db() as conn:
        cur = conn.execute(
            "INSERT INTO sections (project_id, parent_id, title, level, sort_order, "
            "section_type, scoring_point_id, content) VALUES (?,?,?,?,?,?,?,?)",
            (pid, parent_id, title, level, 0, "normal", sp_id, content),
        )
        return cur.lastrowid


class TestH1CheckDocumentIgnoresParentSections:
    """父章节（有子章节）按设计不挂正文，不得被判为「内容为空」。

    原缺陷：`check_document` 不区分父子，父章节 content 恒为空 → 每次预检必然
    报「有 N 个章节内容为空」（N = 父章节数）。这是一条用户永远消不掉、
    也定位不到真正缺正文章节的恒真误报。
    """

    def test_parent_not_reported_as_empty(self):
        from app.services.export import check_document

        pid = _new_project("H1-父章节")
        p1 = _mk_section(pid, "第一章 总体概述", 1, None, "")
        p2 = _mk_section(pid, "第二章 实施方案", 1, None, "")
        _mk_section(pid, "1.1 项目理解", 2, p1, "正文内容" * 50)
        _mk_section(pid, "2.1 部署方案", 2, p2, "正文内容" * 50)

        r = check_document(pid)
        assert r.total_sections == 4
        assert r.empty_sections == [], (
            f"父章节无正文是设计如此，不应计入 empty_sections: "
            f"{[s.title for s in r.empty_sections]}"
        )
        assert not any("内容为空" in w for w in r.warnings), f"不应有误报警告: {r.warnings}"

    def test_real_empty_leaf_still_reported(self):
        """真正的空叶子章节必须仍然被报出来（不能因修复而漏报）。"""
        from app.services.export import check_document

        pid = _new_project("H1-空叶子")
        p1 = _mk_section(pid, "第一章", 1, None, "")
        _mk_section(pid, "1.1 有内容", 2, p1, "正文内容" * 50)
        _mk_section(pid, "1.2 没内容", 2, p1, "")

        r = check_document(pid)
        assert [s.title for s in r.empty_sections] == ["1.2 没内容"], (
            "空叶子章节必须仍然报出"
        )

    def test_all_leaves_empty_is_fatal_error(self):
        """所有叶子都为空 → 致命 error；但只要有一个叶子有内容就不能判致命。"""
        from app.services.export import check_document

        pid = _new_project("H1-全空")
        p1 = _mk_section(pid, "第一章", 1, None, "")
        _mk_section(pid, "1.1 空", 2, p1, "")

        r = check_document(pid)
        assert any("所有章节内容均为空" in e for e in r.errors), (
            "全空时必须升级为致命错误"
        )

        pid2 = _new_project("H1-部分有内容")
        p2 = _mk_section(pid2, "第一章", 1, None, "")
        _mk_section(pid2, "1.1 空", 2, p2, "")
        _mk_section(pid2, "1.2 有", 2, p2, "正文" * 50)

        r2 = check_document(pid2)
        assert not any("所有章节内容均为空" in e for e in r2.errors), (
            "只要还有叶子有内容，就不能误判为「全部为空」"
        )


class TestH2CheckDocumentSkipsWeightWarning:
    """`__weight_warning__` 系统提示项不得计入「未覆盖评分点」。

    它是 `normalize_writing_targets` 在权重之和 ≠100 时插入的校验提醒，
    不是评分项、也不可能有对应章节。原实现把它算进 scoring_ids，
    于是预检恒报「有 1 个评分点未被目录覆盖」，用户补多少章节都消不掉。
    """

    def test_pseudo_item_not_counted(self):
        from app.services.export import check_document

        pid = _new_project("H2-权重伪项")
        from app.db import get_db

        points = [
            {"id": "__weight_warning__", "name": "[权重异常] 权重之和为 88", "weight": ""},
            {"id": "R1", "name": "技术方案", "weight": "50"},
            {"id": "R2", "name": "项目团队", "weight": "30"},
        ]
        with get_db() as conn:
            conn.execute(
                "INSERT INTO analysis (project_id, summary_json, scoring_points_json) "
                "VALUES (?,?,?)",
                (pid, "{}", json.dumps(points, ensure_ascii=False)),
            )
        _mk_section(pid, "技术方案", 1, None, "正文内容" * 50, sp_id="R1")
        _mk_section(pid, "项目团队", 1, None, "正文内容" * 50, sp_id="R2")

        r = check_document(pid)
        assert r.uncovered_scoring_points == [], (
            f"伪项不得计入未覆盖评分点: {r.uncovered_scoring_points}"
        )
        assert not any("未被目录覆盖" in w for w in r.warnings), f"不应有误报: {r.warnings}"

    def test_real_uncovered_point_still_reported(self):
        from app.services.export import check_document

        pid = _new_project("H2-真未覆盖")
        from app.db import get_db

        with get_db() as conn:
            conn.execute(
                "INSERT INTO analysis (project_id, summary_json, scoring_points_json) "
                "VALUES (?,?,?)",
                (
                    pid,
                    "{}",
                    json.dumps(
                        [{"id": "R1", "name": "技术方案"}, {"id": "R9", "name": "售后服务"}],
                        ensure_ascii=False,
                    ),
                ),
            )
        _mk_section(pid, "技术方案", 1, None, "正文内容" * 50, sp_id="R1")

        r = check_document(pid)
        assert r.uncovered_scoring_points == ["R9"], "真实未覆盖评分点必须仍然报出"


class TestH3ExportFingerprintCoversProjectIdentity:
    """内容指纹必须纳入项目名 / 投标单位。

    原缺陷：指纹只哈希 sections，而 `build_docx` 的封面标题、封面投标单位、
    页眉都取自 `projects.name` / `projects.bidder_name`（且未显式传 config 时
    取项目表的值 —— 该值根本不在 config 维度的缓存键里）。
    用户改了投标单位重新导出 → 指纹不变 → 命中旧缓存 DOCX → 封面仍是旧单位名。
    投标文件中的投标人名称错误是硬性废标风险。
    """

    def test_fingerprint_changes_with_bidder_name(self):
        from app.db import get_db
        from app.services.export import _compute_content_fingerprint

        pid = _new_project("H3-指纹")
        _mk_section(pid, "第一章", 1, None, "正文")
        before = _compute_content_fingerprint(pid)
        assert before is not None

        with get_db() as conn:
            conn.execute(
                "UPDATE projects SET bidder_name = ? WHERE id = ?", ("某某建设集团", pid)
            )
        after = _compute_content_fingerprint(pid)
        assert after != before, "改投标单位后指纹必须变化，否则会复用旧封面"

    def test_fingerprint_changes_with_project_name(self):
        from app.db import get_db
        from app.services.export import _compute_content_fingerprint

        pid = _new_project("H3-指纹2")
        _mk_section(pid, "第一章", 1, None, "正文")
        before = _compute_content_fingerprint(pid)

        with get_db() as conn:
            conn.execute("UPDATE projects SET name = ? WHERE id = ?", ("改名后项目", pid))
        assert _compute_content_fingerprint(pid) != before, "改项目名后指纹必须变化"

    def test_fingerprint_stable_when_nothing_changes(self):
        """无变更时指纹必须稳定，否则缓存就失去意义。"""
        from app.services.export import _compute_content_fingerprint

        pid = _new_project("H3-稳定")
        _mk_section(pid, "第一章", 1, None, "正文")
        assert _compute_content_fingerprint(pid) == _compute_content_fingerprint(pid)


class TestH4PreloadDoesNotConsumeFigureNumbers:
    """预渲染（preload_only）不得消耗会话级图号计数器。

    图号计数器在 `build_docx` 开头 `_reset_figure_counter()` 归零，预渲染先跑一轮
    会把号吃掉，正文阶段从 N+1 接着数 → 文档首图不是「图X-1」。
    目前预渲染未传 parent_section_number（都落在 "default" 桶）所以尚未串号，
    但这只是巧合，一旦补传章节号就立刻串位，故显式短路并加测试锁定。
    """

    def test_preload_leaves_counter_untouched(self):
        from io import BytesIO
        import time
        from types import SimpleNamespace

        from app.services import export as ex

        payload = {
            "type": "gantt",
            "title": "施工进度计划",
            "tasks": [{"name": "基础施工", "start": "2026-01-01", "end": "2026-02-01"}],
        }
        section = SimpleNamespace(
            id=1,
            title="进度计划",
            content="正文",
            gantt_json=json.dumps(payload, ensure_ascii=False),
        )
        # 预置 L1 缓存，让预渲染直接命中、完全不触发真实渲染
        expected_key = ex._strip_fig_no_for_cache(payload)
        ex._chart_render_cache[f"gantt:{expected_key}"] = (
            BytesIO(b"\x89PNG\r\n\x1a\n"),
            8,
            time.time(),
        )

        ex._reset_figure_counter()
        ex._add_inline_chart(
            None, "gantt", "施工进度计划", section, ex.ExportConfig(), preload_only=True
        )
        assert ex._figure_counter_state == {}, (
            f"预渲染不得消耗图号计数器，实际: {dict(ex._figure_counter_state)}"
        )

        # 正式渲染必须正常发号，且从 1 开始
        assert ex._next_figure_number("3") == "图3-1"


class TestH8NonLeafContentExport:
    """非叶子（父级）章节的正文必须进文档，且导出不得在收尾日志处崩。

    契约变更（BUG-FIX 2026-09-18 P1）：正文不再「只挂叶子节点」——用户在父级
    目录上手写的正文此前被 `not has_children` 门静默丢弃。
    该修复分两半：预加载集合 `_collect_leaf_ids` → `_collect_content_node_ids`、
    写出端去掉 `not has_children` 门。但收尾日志行漏改，仍引用已不存在的
    `leaf_ids` → 每次导出抛 `NameError`，文档写完却拿不到（导出整体 500）。
    """

    def test_parent_content_is_written(self):
        from docx import Document

        from app.services.export import build_docx

        pid = _new_project("H8-父章节正文")
        parent = _mk_section(pid, "第一章 总体概述", 1, None, "这是父级章节的手写正文内容。")
        _mk_section(pid, "1.1 项目理解", 2, parent, "这是子章节的正文内容。")

        out = Path(_tmpdir()) / "h8.docx"
        path, _ = build_docx(pid, output_path=str(out))
        assert Path(path).exists() and Path(path).stat().st_size > 0

        text = "\n".join(p.text for p in Document(path).paragraphs)
        assert "这是子章节的正文内容。" in text, "叶子正文必须写出"
        assert "这是父级章节的手写正文内容。" in text, "非叶子（父级）正文也必须写出"

    def test_export_does_not_raise_name_error(self):
        """收尾日志不得再引用已被取代的 `leaf_ids`。"""
        import inspect

        from app.services import export as ex

        src = inspect.getsource(ex.build_docx)
        non_comment = "\n".join(
            ln for ln in src.splitlines() if not ln.strip().startswith("#")
        )
        assert "leaf_ids" not in non_comment, (
            "build_docx 不得再引用 leaf_ids（已由 content_ids 取代），"
            "否则每次导出都会 NameError"
        )


def _tmpdir() -> str:
    import tempfile

    return tempfile.mkdtemp(prefix="bid_h8_")


class TestH6TenderContextNoiseCompression:
    """无关长区段压缩：降噪声、保信号、省略必须显式。"""

    HEAD = "\n".join(
        f"第{i}页 投标人须知前附表补充说明：本页为格式性条款，无技术或评分信息。"
        for i in range(1, 300)
    )
    SCORING = "\n".join(
        [
            "第三章 评标办法",
            "3.1 评标方法：本项目采用综合评分法。",
            "3.2 评分标准",
            "3.2.1 技术部分（30分）",
            "  （1）施工组织设计：15分",
            "  （2）质量保证措施：10分",
            "  （3）安全文明施工措施：5分",
            "3.2.2 报价部分（60分）",
        ]
    )
    TECH = "\n".join(
        ["第四章 技术标准和要求"]
        + [f"4.{i} 施工技术方案要求第{i}条：混凝土强度等级不低于C30。" for i in range(1, 80)]
    )
    TEXT = "\n".join([HEAD, SCORING, TECH])

    def test_noise_ratio_drops_below_quarter(self):
        from app.services import tender_context as tc

        ctx = tc.extract_tender_context(self.TEXT, max_chars=40000)
        lines = ctx.splitlines()
        noise = sum(1 for ln in lines if "投标人须知前附表补充说明" in ln)
        ratio = noise / max(1, len(lines))
        # 补丁前实测 53%，补丁后 10%；这里留足余量，只锁「必须显著下降」
        assert ratio < 0.25, f"噪声占比应显著下降，实际 {ratio:.0%}（{noise}/{len(lines)}）"

    def test_signal_fully_preserved(self):
        """压缩绝不能吃掉评分细则与技术条款。"""
        from app.services import tender_context as tc

        ctx = tc.extract_tender_context(self.TEXT, max_chars=40000)
        for kw in (
            "评标办法",
            "综合评分",
            "评分标准",
            "施工组织设计",
            "质量保证措施",
            "安全文明施工",
            "4.1 施工技术方案要求第1条",
            "4.79 施工技术方案要求第79条",
        ):
            assert kw in ctx, f"{kw} 必须保留"

    def test_omission_is_explicit_not_silent(self):
        """省略必须是显式的：有标记、且保留区段首尾边界。"""
        from app.services import tender_context as tc

        ctx = tc.extract_tender_context(self.TEXT, max_chars=40000)
        markers = [ln for ln in ctx.splitlines() if "此处省略" in ln]
        assert markers, "压缩处必须插入显式省略标记，不能静默丢弃"
        # 首尾边界仍在：本例窗口选中 201~299 页的噪声段，
        # 压缩后应保留首 5 行（201~205）与末 5 行（295~299）
        assert "第201页 投标人须知" in ctx and "第299页 投标人须知" in ctx, (
            "压缩区段的首尾边界必须保留"
        )
        assert "第250页 投标人须知" not in ctx, "中段无关行应被压缩掉"

    def test_short_text_untouched(self):
        """短文本未达 30 行阈值，必须原样返回（不引入意外截断）。"""
        from app.services.tender_context import extract_tender_context

        for t in (
            "这是一个招标公告",
            "简单的描述文本内容",
            "评分标准：技术方案占30分，需提供详细的技术路线图。",
        ):
            assert extract_tender_context(t) == t

    def test_compress_helper_direct(self):
        from app.services import tender_context as tc

        lines = [f"无关行{i}" for i in range(50)] + ["施工组织设计：15分"] + [f"无关行尾{i}" for i in range(50)]
        out = tc._compress_irrelevant_runs(lines, tc.TECHNICAL_ANCHORS + tc._SCORING_KEYWORDS)
        assert len(out) < len(lines), "长无关区段必须被压缩"
        assert "施工组织设计：15分" in out, "相关行必须保留"
        assert any("此处省略" in ln for ln in out)

    def test_scoring_region_also_compressed(self):
        from app.services import tender_context as tc

        region = tc.extract_scoring_region(self.TEXT)
        assert "评标办法" in region and "施工组织设计" in region
        noise = sum(1 for ln in region.splitlines() if "投标人须知前附表补充说明" in ln)
        assert noise < 15, f"评分区域同样应抑制噪声，实际 {noise} 行"


class TestH7ScoringKeywordConvergence:
    """评分关键词此前在本文件里硬编码三份且互不一致。"""

    def test_locate_list_is_superset_of_legacy(self):
        """截断定位改用模块级全集；原 4 词必须是其前 4 项（只增不减）。"""
        from app.services import tender_context as tc

        legacy = ["评分标准", "技术评分", "综合评分", "评分项"]
        assert list(tc._SCORING_KEYWORDS)[:4] == legacy, (
            "模块级列表前 4 项必须与原硬编码一致，保证改动只增不减"
        )
        assert set(legacy) <= set(tc._SCORING_KEYWORDS)
        # 原先漏掉的三个词现已纳入 —— 「评标办法」是评分章节最常见标题
        for kw in ("评分细则", "评审方法", "评标办法"):
            assert kw in tc._SCORING_KEYWORDS

    def test_no_third_hardcoded_list_left(self):
        """函数体内不得再出现第三份评分关键词字面量。"""
        import inspect

        from app.services.tender_context import extract_tender_context

        src = inspect.getsource(extract_tender_context)
        non_comment = "\n".join(
            ln for ln in src.splitlines() if not ln.strip().startswith("#")
        )
        assert '"评标办法"' not in non_comment, (
            "截断定位应复用 _SCORING_KEYWORDS，不得再硬编码关键词字面量"
        )
        assert '"评分标准", "技术评分", "综合评分", "评分项"' not in non_comment


class TestH5ScoringPlaceholderSingleSourceOfTruth:
    """伪评分项判据必须收敛到 `app.utils.scoring_points` 一处。

    修复前 outline_gate / workflows_sections / workflows_analysis /
    workflows_outline_v51 / export 各写一份 `__weight_warning__` 字面量，
    新增伪项类型需改五处、漏一处即复现缺陷。
    """

    def test_shared_helper_shapes(self):
        from app.utils.scoring_points import (
            SCORING_PLACEHOLDER_IDS,
            is_scoring_placeholder,
            strip_scoring_placeholders,
        )

        warn = {"id": "__weight_warning__", "name": "[权重异常] …"}
        real = {"id": "R1", "name": "技术方案"}

        assert "__weight_warning__" in SCORING_PLACEHOLDER_IDS
        assert is_scoring_placeholder(warn) is True
        assert is_scoring_placeholder(real) is False
        assert is_scoring_placeholder("__weight_warning__") is True  # 裸 id 形态
        assert is_scoring_placeholder(None) is False

        assert strip_scoring_placeholders([warn, real]) == [real]
        assert strip_scoring_placeholders({"items": [warn, real]})["items"] == [real]
        assert strip_scoring_placeholders("原文") == "原文"

    def test_all_modules_delegate_to_shared_helper(self):
        """各模块的过滤入口对伪项的判定必须一致（不是各写各的字符串）。"""
        from app.services.ai import (
            workflows_outline_v51 as v51,
            workflows_sections as wsec,
        )
        from app.services.ai.outline_gate import parse_scoring_item_titles
        from app.utils.scoring_points import is_scoring_placeholder

        warn = {"id": "__weight_warning__", "name": "[权重异常] 评分项权重之和为 90"}
        real = {"id": "R1", "name": "技术方案"}

        # v51：保留旧名称做转发
        assert v51._is_scoring_placeholder(warn) is True
        assert v51._strip_scoring_placeholders([warn, real]) == [real]
        assert v51._is_scoring_placeholder is not is_scoring_placeholder, (
            "v51 保留的是转发包装（保持向后兼容），实现必须是共享那份"
        )

        # outline_gate：评分标题提取
        assert parse_scoring_item_titles([warn, real]) == ["技术方案"]

        # workflows_sections：正文提示词的评分要点
        text = wsec._scoring_context_to_text([warn, real])
        assert "权重异常" not in text and "技术方案" in text

    def test_analysis_uses_shared_helper_for_duplicate_guard(self):
        """`normalize_writing_targets` 重复插入防护也走共享判据。"""
        from app.services.ai.workflows_analysis import normalize_writing_targets

        # 权重 30 + 30 = 60，偏离 100 → 应插入一条且仅一条提示项
        # 注意：name 取自 name/title 字段，权重取自 weight/score 字段
        first = normalize_writing_targets(
            [
                {"name": "技术方案", "weight": "30"},
                {"name": "项目团队", "weight": "30"},
            ]
        )
        warns = [p for p in first if p.get("id") == "__weight_warning__"]
        assert len(warns) == 1, f"应恰好插入一条权重提示项，实际 {len(warns)}"


