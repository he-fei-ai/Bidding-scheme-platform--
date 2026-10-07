"""回归测试（2026-09-16 第三轮）：深度探索后的系统性 BUG 修复

覆盖的缺陷与改动文件（只测行为，不测实现细节）：

A. `app/db.py`
   A1. GET_DB 重入计数泄漏：取连接失败会在 try/finally 之外抛出，线程本地的
       `_db_ctx.depth` 永远回不去 → 该线程后续所有写操作都「不提交、不归还连接」。
   A2. `get_db_with_retry` 新增 `readonly`：纯读不再抢全局写锁。
   A3. `query_retry` 不再吞掉被包装函数的 `max_retries` / `retry_delay` 实参。

B. 资源 / 安全
   B1. `routers/ai_config.py`：提交脱敏占位符时保持库内原值（密文仍是密文），
       不再「解密后回写」把加密降级成明文；无法解析时报错而不是写占位符。
   B2. `routers/api_keys.py`：列表 / 用量端点补管理员校验（原先仅登录即可）。
   B3. `routers/websocket.py`：REST 端点不再信任 query 里的 user_id。
   B4. `services/files.py`：openpyxl 异常路径必须 close workbook。

C. 计数 / 并发
   C1. `services/generation_manager.py`：失败/取消计数幂等（重试不重复计数）。
   C2. `services/ai/usage_stats.py`：`truncate_stats` 返回本次 DELETE 行数。
   C3. `routers/sse/sse_utils.py`：`regenerate_ids` 含非法段不再让 int() 穿透。

D. 静态守卫
   D1. `routers/export.py`：SSE 导出的 background task 必须被强引用。
"""

from __future__ import annotations

import asyncio
import inspect
import sqlite3

import pytest

from app import db as db_mod
from app.config import DATABASE_PATH


def _connect():
    """返回**需要显式关闭**的连接（sqlite3 的 with 只提交不关闭，勿改成 with）。"""
    conn = sqlite3.connect(str(DATABASE_PATH))
    conn.row_factory = sqlite3.Row
    return conn


# ---------------------------------------------------------------------------
# A. db.py
# ---------------------------------------------------------------------------


def test_get_db_restores_nesting_depth_when_connection_fails(monkeypatch):
    """A1：取连接耗尽重试后，depth 必须恢复，否则该线程后续写入静默丢失。"""

    def _boom():
        raise sqlite3.OperationalError("disk I/O error")

    monkeypatch.setattr(db_mod, "_get_connection", _boom)
    monkeypatch.setattr(db_mod.time, "sleep", lambda *_a: None)

    depth_before = getattr(db_mod._db_ctx, "depth", 0)

    with pytest.raises(Exception):
        with db_mod.get_db() as _conn:  # noqa: F841
            pass

    assert getattr(db_mod._db_ctx, "depth", 0) == depth_before


def test_get_db_with_retry_restores_nesting_depth_when_connection_fails(monkeypatch):
    """A1（对称场景）：get_db_with_retry 同样不能泄漏 depth。"""

    def _boom():
        raise sqlite3.OperationalError("disk I/O error")

    monkeypatch.setattr(db_mod, "_get_connection", _boom)
    monkeypatch.setattr(db_mod.time, "sleep", lambda *_a: None)

    depth_before = getattr(db_mod._db_ctx, "depth", 0)

    with pytest.raises(Exception):
        with db_mod.get_db_with_retry() as _conn:  # noqa: F841
            pass

    assert getattr(db_mod._db_ctx, "depth", 0) == depth_before


class _CountingLock:
    """替代全局写锁，统计被获取次数（RLock 语义）。"""

    def __init__(self):
        self.acquired = 0
        self.released = 0

    def acquire(self, timeout=None):  # noqa: D102
        self.acquired += 1
        return True

    def release(self):  # noqa: D102
        self.released += 1


def test_get_db_with_retry_readonly_skips_write_lock(monkeypatch):
    """A2：readonly=True 不抢写锁，默认仍然抢（语义保持一致）。"""
    fake_lock = _CountingLock()
    monkeypatch.setattr(db_mod, "_db_write_lock", fake_lock)

    with db_mod.get_db_with_retry(readonly=True) as conn:
        conn.execute("SELECT 1")

    assert fake_lock.acquired == 0, "纯读连接不应获取全局写锁"

    with db_mod.get_db_with_retry() as conn:
        conn.execute("SELECT 1")

    assert fake_lock.acquired == 1
    assert fake_lock.released == 1


def test_query_retry_keeps_target_kwargs():
    """A3：装饰器不能吞掉被包装函数的同名实参。"""

    @db_mod.query_retry
    def _plain(value, max_retries=5, retry_delay=0.01):
        return (value, max_retries, retry_delay)

    assert _plain(1, max_retries=9, retry_delay=0.02)[1:] == (9, 0.02)

    @db_mod.query_retry(max_retries=2, retry_delay=0.01)
    def _configured(value, max_retries=5):
        return (value, max_retries)

    assert _configured(1, max_retries=7) == (1, 7)


def test_query_retry_retries_io_errors_then_gives_up(monkeypatch):
    """A3（配套）：I/O 错误仍需重试 max_retries 次后原样抛最后异常。"""
    monkeypatch.setattr(db_mod.time, "sleep", lambda *_a: None)

    calls = {"n": 0}

    @db_mod.query_retry(max_retries=3, retry_delay=0)
    def _flaky():
        calls["n"] += 1
        raise sqlite3.OperationalError("database is locked")

    with pytest.raises(sqlite3.OperationalError):
        _flaky()

    assert calls["n"] == 3


# ---------------------------------------------------------------------------
# B1. ai_config 加密不降级
# ---------------------------------------------------------------------------


def _read_ai_config_row():
    conn = _connect()
    try:
        row = conn.execute("SELECT * FROM ai_config WHERE id = 1").fetchone()
        return dict(row) if row else None
    finally:
        conn.close()


@pytest.fixture
def _restore_ai_config():
    snapshot = _read_ai_config_row()
    yield
    conn = _connect()
    try:
        conn.execute("DELETE FROM ai_config WHERE id = 1")
        if snapshot is not None:
            cols = ", ".join(snapshot.keys())
            marks = ", ".join("?" for _ in snapshot)
            conn.execute(
                f"INSERT INTO ai_config ({cols}) VALUES ({marks})",
                tuple(snapshot.values()),
            )
        conn.commit()
    finally:
        conn.close()


async def test_saving_masked_key_keeps_stored_value(_restore_ai_config):
    """B1：提交占位符时保持库内原值，且不得把 ENC: 密文降级成明文。"""
    from app.routers import ai_config as ai_config_router
    from app.routers.ai_config import AIConfigData
    from app.services.crypto import encrypt_api_key, is_encrypted

    conn = _connect()
    conn.execute("DELETE FROM ai_config WHERE id = 1")
    conn.commit()
    conn.close()

    plain = "sk-round3-plain-key"
    await ai_config_router.save_ai_config(AIConfigData(provider_name="t", api_key=plain))
    # 同步 2026-10-04（F-1 测试钉更新）：原断言「新密钥明文入库」钉的是旧实现；
    # 新密钥统一 ENC: 加密落库，消费侧透明解密后语义不变。
    stored_first = (_read_ai_config_row() or {}).get("api_key") or ""
    assert is_encrypted(stored_first) and plain not in stored_first

    # ① 加密入库 → 提交掩码后库内值不变（不是解密后的重复往返，也不二次加密）
    await ai_config_router.save_ai_config(AIConfigData(provider_name="t", api_key="***已配置***"))
    assert (_read_ai_config_row() or {}).get("api_key") == stored_first

    # ② 密文入库 → 提交掩码后必须**仍是密文**（旧实现会解密后回写，加密被抹掉）
    conn = _connect()
    conn.execute("UPDATE ai_config SET api_key = ? WHERE id = 1", (encrypt_api_key(plain),))
    conn.commit()
    conn.close()

    await ai_config_router.save_ai_config(AIConfigData(provider_name="t", api_key="***已配置***"))
    stored = (_read_ai_config_row() or {}).get("api_key") or ""
    assert is_encrypted(stored), "保存掩码不得把加密密钥降级为明文"
    assert plain not in stored


async def test_saving_masked_key_without_existing_row_is_rejected(_restore_ai_config):
    """B1（连带）：解析不到已有密钥时必须报错，不能把占位符写成密钥。"""
    from fastapi import HTTPException

    from app.routers import ai_config as ai_config_router

    from app.routers.ai_config import AIConfigData

    with _connect() as conn:
        conn.execute("DELETE FROM ai_config WHERE id = 1")

    with pytest.raises(HTTPException) as info:
        await ai_config_router.save_ai_config(AIConfigData(provider_name="t", api_key="***已配置***"))

    assert info.value.status_code == 400
    assert _read_ai_config_row() is None  # 占位符绝不能落库


# ---------------------------------------------------------------------------
# B2. api_keys 鉴权
# ---------------------------------------------------------------------------


async def test_list_and_usage_endpoints_require_admin(monkeypatch):
    """B2：只读端点也必须有管理员校验（否则普通用户可枚举全部 Key）。"""
    from fastapi import HTTPException

    from app.routers import api_keys as ak
    from app.security.auth import User

    class _Req:
        @property
        def state(self):
            return self

    original = ak.get_current_user
    ak.get_current_user = lambda request: User(user_id=2, username="p", role="user")
    try:
        with pytest.raises(HTTPException) as e1:
            await ak.list_api_keys(_Req())
        assert e1.value.status_code == 403

        with pytest.raises(HTTPException) as e2:
            await ak.get_api_key_usage(1, _Req())
        assert e2.value.status_code == 403
    finally:
        ak.get_current_user = original


# ---------------------------------------------------------------------------
# B3. websocket REST user_id
# ---------------------------------------------------------------------------


def test_trusted_user_id_prefers_authenticated_user(monkeypatch):
    """B3：AUTH_ENABLED 下不得再信任 query 里的 user_id。"""
    from app.routers.websocket import _trusted_user_id

    class _Req:
        def __init__(self, current_user):
            self.state = type("_S", (), {"current_user": current_user})()

    monkeypatch.setattr("app.config.AUTH_ENABLED", False)
    assert _trusted_user_id(_Req({"username": "alice"}), "bob") == "bob"

    monkeypatch.setattr("app.config.AUTH_ENABLED", True)
    assert _trusted_user_id(_Req({"username": "alice", "id": 7}), "bob") == "alice"
    # 已鉴权的 User 对象形态（中间件可能塞对象而非 dict）
    assert _trusted_user_id(_Req(type("_U", (), {"username": "carol", "id": 3})()), "bob") == "carol"
    # 无中间件信息时退化为旧行为，不阻断
    assert _trusted_user_id(_Req(None), "bob") == "bob"


# ---------------------------------------------------------------------------
# B4. openpyxl 句柄
# ---------------------------------------------------------------------------


def test_xlsx_extraction_closes_workbook_on_error(monkeypatch, tmp_path):
    """B4：解析中途报错也必须 close，避免句柄泄漏 / Windows 删不掉文件。"""
    import openpyxl

    from app.services import files as files_mod

    class _WS:
        def iter_rows(self, values_only=True):
            raise RuntimeError("模拟读取失败")

    class _WB:
        sheetnames = ["Sheet1"]
        closed = False

        def __getitem__(self, _name):
            return _WS()

        def close(self):
            self.closed = True

    wb = _WB()
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *a, **k: wb)

    valid = tmp_path / "sample.xlsx"
    openpyxl.Workbook().save(valid)  # 真实文件，保证回退路径能跑完

    files_mod._extract_xlsx_openpyxl(str(valid))
    assert wb.closed, "异常路径必须关闭 workbook"


# ---------------------------------------------------------------------------
# C1. generation_manager 计数幂等
# ---------------------------------------------------------------------------


@pytest.fixture
def _cleanup_generation_rows():
    yield
    conn = _connect()
    try:
        conn.execute("DELETE FROM generation_tasks WHERE generation_id LIKE 'round3-%'")
        conn.execute("DELETE FROM section_generation_progress WHERE generation_id LIKE 'round3-%'")
        conn.commit()
    finally:
        conn.close()


async def test_failed_and_cancelled_counts_are_idempotent(_cleanup_generation_rows):
    """C1：同一章节多次重试失败只计一次；脱离终态要撤销，避免任务提前进入终态。"""
    from app.services.generation_manager import GenerationManager, GenerationTask

    mgr = GenerationManager()
    task = GenerationTask(generation_id="round3-1", project_id=990301, total_sections=3)
    mgr._tasks[990301] = task

    await mgr.update_section_progress(990301, 11, "failed", error="第一次失败")
    # SSE 重试：先回调 generating，再回调 failed
    await mgr.update_section_progress(990301, 11, "generating", progress=10)
    await mgr.update_section_progress(990301, 11, "failed", error="第二次失败")
    assert task.failed_count == 1, "重试不得重复计入失败"

    await mgr.update_section_progress(990301, 11, "generating", progress=0)
    assert task.failed_count == 0, "离开 failed 态应撤销计数"

    body = "正文" * 20
    await mgr.update_section_progress(990301, 11, "completed", content=body)
    await mgr.update_section_progress(990301, 11, "completed", content=body)
    assert task.completed_count == 1
    assert task.failed_count == 0

    await mgr.update_section_progress(990301, 12, "cancelled")
    await mgr.update_section_progress(990301, 12, "generating", progress=5)
    await mgr.update_section_progress(990301, 12, "cancelled")
    assert task.cancelled_count == 1

    # 计数没有因重试而虚增，任务不会提前达到终态
    assert task.completed_count + task.failed_count + task.cancelled_count < task.total_sections
    assert task.status == "running"


# ---------------------------------------------------------------------------
# C2. usage_stats 清理行数
# ---------------------------------------------------------------------------


def test_truncate_stats_returns_deleted_rows():
    """C2：必须返回本次 DELETE 的行数，而不是连接的累计变更数。"""
    from datetime import datetime, timedelta

    from app.services.ai.usage_stats import truncate_stats

    marker = "round3-truncate"
    old = "2020-01-01 00:00:00"
    rows = [(f"call-{marker}-{i}", marker, old) for i in range(3)]
    fresh = (f"call-{marker}-keep", marker, "2999-01-01 00:00:00")

    conn = _connect()
    try:
        for call_id, profile, created in rows + [fresh]:
            conn.execute(
                "INSERT OR REPLACE INTO ai_usage_stats "
                "(call_id, model, base_url, created_at, profile_name) VALUES (?, 'm', 'http://x', ?, ?)",
                (call_id, created, profile),
            )
        conn.commit()
    finally:
        conn.close()

    cutoff = (datetime.now() - timedelta(days=1)).strftime("%Y-%m-%d %H:%M:%S")
    conn = _connect()
    try:
        expected = conn.execute(
            "SELECT COUNT(*) c FROM ai_usage_stats WHERE created_at < ?", (cutoff,)
        ).fetchone()["c"]
    finally:
        conn.close()

    deleted = truncate_stats(keep_days=1)

    conn = _connect()
    try:
        left = [
            r["call_id"]
            for r in conn.execute(
                "SELECT call_id FROM ai_usage_stats WHERE profile_name = ?", (marker,)
            ).fetchall()
        ]
        conn.execute("DELETE FROM ai_usage_stats WHERE profile_name = ?", (marker,))
        conn.commit()
    finally:
        conn.close()

    assert expected >= 3
    assert deleted == expected, f"应返回本次删除的 {expected} 行，实际 {deleted}"
    assert left == [fresh[0]]


# ---------------------------------------------------------------------------
# C3. sse_utils 非法 regenerate_ids
# ---------------------------------------------------------------------------


def test_build_leaf_dicts_tolerates_mixed_regenerate_ids():
    """C3：`"abc,5"` 这类混合输入不应让统计段的 int() 穿透导致 SSE 断流。"""
    from app.routers.sse.sse_utils import _build_leaf_dicts

    sections = {
        5: {
            "id": 5,
            "parent_id": 4,
            "level": 3,
            "title": "施工工艺",
            "section_type": "text",
            "status": "pending",
            "content": "",
            "sort_order": 1,
        },
        4: {"id": 4, "parent_id": None, "level": 2, "title": "章", "status": "pending", "content": ""},
    }

    # 旧实现在日志统计行抛 ValueError（未被任何 except 覆盖）
    leaves = _build_leaf_dicts(
        project_id=990302,
        sections_by_id=sections,
        children_map={4: [5]},
        root_section_id=None,
        regenerate_ids="abc,5",
        regenerate_missing=False,
        level2_default=1000,
    )
    assert isinstance(leaves, list)


# ---------------------------------------------------------------------------
# D1. 导出后台任务强引用（静态守卫）
# ---------------------------------------------------------------------------


def test_export_build_task_is_strongly_referenced():
    """D1：`asyncio.create_task` 结果必须入组 + done 回调清理，否则会被 GC。"""
    from app.routers.export import export_docx_stream

    src = inspect.getsource(export_docx_stream)
    assert "_BUILD_TASKS.add(" in src, "导出构建任务缺少强引用"
    assert "add_done_callback(" in src, "导出构建任务缺少完成清理"
