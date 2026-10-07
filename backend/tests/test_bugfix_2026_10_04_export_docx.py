"""导出文档模块 —— BUG 修复回归测试（2026-10-04）。

覆盖本轮识别并修复的缺陷（每项均有修复前复现依据）：

1. BUG-A（P1 超长内容）：`_add_paragraph_with_format` 只清洗不截断，
   AI 超长单段（>3.2 万字）原样落盘 → Word 打开成品提示「文档需要修复」。
   修复前实测：40000 字段落完整写入段落（para_len=40000）。
   修复口径：与表格单元格同用 MAX_DOCX_TEXT_CHARS（32000）截断。

2. BUG-B（P1 特殊字符）：build_docx 封面项目名/投标单位与页眉文本直接取
   projects 表原文写 add_run，未走 sanitize。含 \\x0b 时 ValueError
   （修复前实测复现「All strings must be XML compatible」），被 try/except
   整块吞掉 → 封面/页眉整体静默消失。

3. BUG-C（P1 特殊字符·图题链路）：`_add_diagram_picture` / `_add_image_placeholder`
   / 渲染失败说明的 caption 源于章节 *_json 图表标题（未走正文清洗链路），
   add_run 炸掉会让渲染成功的好图按失败处理、连兜底说明都写不进去。

4. P1 图号占号回归（配合验证）：渲染失败回滚图号 + 失败说明不携带编号
   （图号虚跳缺陷 2026-09-29 修复的并行守卫，本轮以「失败说明含控制字符」
   新形态一并钉住）。

5. P2 口径对齐：GET /export/latest 与 /download/{id} 同口径限制 .docx/.dotx。

6. 路由契约：导出模块全部端点（同步/SSE/预检/下载/最新/历史/预设）注册齐全，
   与 frontend/src/api/index.ts 的调用路径一一对应。

测试规范：包含真实 docx 渲染断言（保存 → 重新打开 → 校验落盘文本）。
"""
from __future__ import annotations

import base64
import io
import uuid

import pytest
from docx import Document
from fastapi import HTTPException

from app.db import get_db, init_db
from app.models import ExportConfig
from app.security.auth import User
from app.services import export as export_mod

OWNER = 100


def _new_project(name: str, created_by: int | None = OWNER) -> int:
    init_db()
    with get_db() as conn:
        conn.execute(
            "INSERT INTO projects (name, industry, created_by) VALUES (?, ?, ?)",
            (name, "通用", created_by),
        )
        return conn.execute(
            "SELECT id FROM projects ORDER BY id DESC LIMIT 1"
        ).fetchone()["id"]


def _make_png_stream() -> io.BytesIO:
    PilImage = pytest.importorskip("PIL.Image", reason="PIL 不可用时跳过图片用例")
    buf = io.BytesIO()
    PilImage.new("RGB", (40, 40), (10, 120, 200)).save(buf, format="PNG")
    buf.seek(0)
    return buf


# ============================================================
# 1. BUG-A：超长段落截断（Word 单段 32767 硬上限）
# ============================================================


class TestOverlongParagraphTruncation:
    def test_paragraph_truncated_to_word_limit(self):
        """40000 字段落必须被截断到 MAX_DOCX_TEXT_CHARS 以内（含省略号）。"""
        cfg = ExportConfig()
        doc = Document()
        long_text = "测" * 40000
        p = export_mod._add_paragraph_with_format(doc, long_text, cfg, "paragraph")
        assert len(p.text) <= export_mod.MAX_DOCX_TEXT_CHARS + 1, (
            f"段落长度 {len(p.text)} 超过 Word 单段上限，成品会触发「文档需要修复」"
        )
        assert p.text.endswith("…"), "截断处必须携带省略号，提示读者内容被裁剪"

    def test_real_docx_roundtrip_keeps_paragraph_under_limit(self):
        """真实渲染断言：保存 → 重新打开，落盘段落仍在安全长度内。"""
        cfg = ExportConfig()
        doc = Document()
        export_mod._add_paragraph_with_format(doc, "文" * 50000, cfg, "paragraph")
        buf = io.BytesIO()
        doc.save(buf)
        reopened = Document(io.BytesIO(buf.getvalue()))
        assert len(reopened.paragraphs[0].text) <= export_mod.MAX_DOCX_TEXT_CHARS + 1

    def test_normal_paragraph_unchanged(self):
        """向后兼容：未超长的普通段落行为不变（不截断、不加省略号）。"""
        cfg = ExportConfig()
        doc = Document()
        text = "根据招标文件要求，所有施工材料须符合国家标准。"
        p = export_mod._add_paragraph_with_format(doc, text, cfg, "paragraph")
        assert p.text == text


# ============================================================
# 2. BUG-B：封面/页眉项目名特殊字符不再整块丢失
# ============================================================


class TestCoverHeaderSanitization:
    def test_build_docx_cover_survives_control_char_project_name(self):
        """项目名含 \\x0b 时封面与页眉必须正常落盘（修复前整块静默消失）。"""
        pid = _new_project("智慧园区\x0b建设工程", created_by=OWNER)
        with get_db() as conn:
            conn.execute(
                "INSERT INTO sections (project_id, title, level, sort_order, content, status) "
                "VALUES (?, '第一节 施工组织设计', 1, 0, ?, 'completed')",
                (pid, "本项目施工组织设计依据国家标准编制，内容包括总体部署与进度计划安排。"),
            )
        cfg = ExportConfig(show_title_page=True, show_header_footer=True)
        output_path, export_id = export_mod.build_docx(pid, config=cfg)
        assert output_path

        reopened = Document(output_path)
        all_text = "\n".join(p.text for p in reopened.paragraphs)
        # 修复判据：封面标题段存在（清洗后项目名仍可在文档中找到）
        assert "智慧园区建设工程" in all_text.replace("\u3000", ""), (
            "封面项目名丢失（修复前控制字符导致 add_run 抛错、封面整块被吞）"
        )
        assert "\x0b" not in all_text
        # 页眉落盘且不含控制字符
        header_text = reopened.sections[0].header.paragraphs[0].text
        assert header_text and "\x0b" not in header_text

    def test_sanitize_is_idempotent_for_clean_names(self):
        """向后兼容：无特殊字符的项目名不受影响。"""
        assert export_mod.sanitize_for_docx("正常项目名") == "正常项目名"


# ============================================================
# 3. BUG-C：图题/占位符/失败说明的 caption 清洗
# ============================================================


class TestCaptionSpecialCharResilience:
    def test_add_diagram_picture_valid_image_survives_bad_caption(self):
        """渲染成功的好图不能因图题含控制字符而被按失败处理（修复前返回 False）。"""
        cfg = ExportConfig()
        doc = Document()
        ok = export_mod._add_diagram_picture(
            doc, _make_png_stream(), cfg, caption="图1-1 施工进度\x0b计划图"
        )
        assert ok is True, "caption 炸掉会让整张图按插入失败处理（修复前缺陷）"
        last_para_text = doc.paragraphs[-1].text
        assert "\x0b" not in last_para_text
        assert "施工进度计划图" in last_para_text

    def test_image_placeholder_survives_bad_caption(self):
        """降级占位符的最后防线：caption 含 XML 不安全字符也不得抛错。"""
        cfg = ExportConfig()
        doc = Document()
        export_mod._add_image_placeholder(
            doc, cfg, caption="图2-1 组\ufffethe织架构", reason="图片未加载"
        )
        texts = [p.text for p in doc.paragraphs]
        assert any("[图片未加载]" in t for t in texts)
        assert all("\ufffe" not in t and "\x0b" not in t for t in texts)

    def test_render_fail_label_is_sanitized_and_fig_no_rewound(self, monkeypatch):
        """渲染失败：失败说明写入且不带图号，图号计数器回滚（不占号）。

        CROSS-FIX 2026-10-04：图表模块新增 CHART_FAIL_PLACEHOLDER 开关后默认
        改为「静默跳过」，历史「始终插入失败说明」行为等价于开关=true。
        本用例钉的是兜底分支行为，显式开启开关（图号回滚与开关无关，
        由调用方 _add_inline_chart 无条件执行）。
        """
        # 补丁：让真实渲染恒失败，走「渲染失败说明」兜底分支
        monkeypatch.setattr(export_mod, "render_mermaid_to_bytes", lambda *a, **k: None)
        monkeypatch.setattr(export_mod, "CHART_FAIL_PLACEHOLDER", True)
        pid = 987_654_321  # 独立项目分桶，避免与其他用例串扰
        export_mod._reset_figure_counter(pid)
        cfg = ExportConfig()
        doc = Document()
        detail = type("S", (), {"id": 1, "project_id": pid, "title": "进度节", "content": ""})()
        marker = {
            "type": "gantt",
            "title": f"进度\x0b计划{uuid.uuid4().hex[:8]}",  # 唯一载荷，绕开渲染缓存
            "tasks": [{"name": "土方开挖", "start": "2025-01-01", "days": 10}],
        }
        rendered = export_mod._add_inline_chart(
            doc,
            "gantt",
            __import__("json").dumps(marker, ensure_ascii=False),
            detail,
            cfg,
            parent_section_number="7",
        )
        assert rendered is False
        fail_texts = [p.text for p in doc.paragraphs if "图表渲染失败" in p.text]
        assert fail_texts, "渲染失败必须留下事实性说明，文档不得空白"
        assert "\x0b" not in fail_texts[0], "失败说明（源于 caption）必须清洗"
        assert not fail_texts[0].startswith("图7-"), "失败说明不得携带已回滚的图号"
        # 图号回滚：失败不占号，下一张图仍是 图7-1（消除图号虚跳）
        nxt = export_mod._next_figure_number("7", pid)
        assert nxt == "图7-1"
        export_mod._reset_figure_counter(pid)


# ============================================================
# 4. P2：/latest 与 /download/{id} 文件类型口径对齐
# ============================================================


class TestDownloadLatestGuard:
    def test_latest_rejects_non_docx_file_path(self, tmp_path, monkeypatch):
        import asyncio
        from pathlib import Path

        from app.config import PROJECTS_DATA_DIR

        pid = _new_project("下载口径-项目", created_by=OWNER)
        # 伪造历史脏行：file_path 指向项目数据目录内的 .txt 文件
        rogue_dir = PROJECTS_DATA_DIR / str(pid) / "exports"
        rogue_dir.mkdir(parents=True, exist_ok=True)
        rogue = rogue_dir / "notes.txt"
        rogue.write_text("not a docx", encoding="utf-8")
        with get_db() as conn:
            conn.execute(
                "INSERT INTO exports (project_id, filename, file_path, file_size) "
                "VALUES (?, ?, ?, 10)",
                (pid, "notes.txt", str(rogue)),
            )
        from app.routers.export import download_latest

        with pytest.raises(HTTPException) as exc:
            asyncio.run(download_latest(pid, identity=User(OWNER, "o", "user")))
        assert exc.value.status_code == 400
        assert not Path(rogue).is_absolute() or rogue.exists()  # 文件未被消费也未被删


# ============================================================
# 5. 路由契约：端点齐全，与前端 api/index.ts 调用一一对应
# ============================================================


class TestSseCompletedContract:
    """ENHANCE 2026-10-04：SSE completed 的 file_path 与同步响应口径统一。

    修复前：同步 POST /docx 的 file_path 是下载 URL，SSE completed 的
    file_path 是磁盘相对路径 —— 同一字段名两种语义，前端 ExportResult
    类型声明按 URL 口径，跨模块消费易误用。
    """

    def test_sse_completed_file_path_aligned_to_download_url(self):
        import asyncio
        import json

        from app.routers.export import export_docx_stream

        pid = _new_project("SSE口径-项目", created_by=OWNER)
        with get_db() as conn:
            conn.execute(
                "INSERT INTO sections (project_id, title, level, sort_order, content, status) "
                "VALUES (?, '第一节 总体部署', 1, 0, ?, 'completed')",
                (pid, "本项目总体部署依据招标文件与国家标准编制，含进度计划与资源配置安排。"),
            )
        cfg = ExportConfig()

        class _FakeRequest:
            async def is_disconnected(self):
                return False

        async def _collect():
            resp = await export_docx_stream(
                _FakeRequest(), pid, cfg, identity=User(OWNER, "o", "user")
            )
            events = []
            async for chunk in resp.body_iterator:
                for line in chunk.splitlines():
                    if line.startswith("data: "):
                        events.append(json.loads(line[6:]))
                if any(e.get("event") in ("completed", "error") for e in events):
                    break
            return events

        events = asyncio.run(_collect())
        completed = [e for e in events if e.get("event") == "completed"]
        assert completed, f"SSE 终态缺失: {[e.get('event') for e in events]}"
        ev = completed[0]
        # 口径统一：file_path == download_url，且为下载 URL 而非磁盘路径
        assert ev["file_path"] == ev["download_url"]
        assert ev["file_path"].startswith(f"/api/v1/projects/{pid}/export/download/")
        # 向后兼容：原始磁盘相对路径仍以 relative_path 透出（非绝对路径）
        assert ev["relative_path"] and not ev["relative_path"].startswith(("/", "\\"))
        assert ev["export_id"] and not ev["record_failed"]

    def test_sync_and_sse_same_field_semantics(self, monkeypatch):
        """同步 POST /docx 与 SSE completed 的 file_path 语义一致（均为 URL）。

        直接比对两处源码契约，防止未来单侧回退。
        """
        import inspect

        from app.routers import export as export_router_mod

        src = inspect.getsource(export_router_mod.export_docx)
        assert "'file_path': f'/api/v1/projects/{project_id}/export/download/" in src
        stream_src = inspect.getsource(export_router_mod.export_docx_stream)
        assert "'file_path': download_url" in stream_src
        assert "'relative_path': rel_path" in stream_src


class TestExportRouteContract:
    def test_all_export_endpoints_registered(self):
        from app.routers.export import _presets_router, router

        paths = {(r.path, tuple(sorted(getattr(r, "methods", []) or []))) for r in router.routes}
        expected = [
            ("/projects/{project_id}/export/docx", ("POST",)),
            ("/projects/{project_id}/export/docx-stream", ("POST",)),
            ("/projects/{project_id}/export/check", ("GET",)),
            ("/projects/{project_id}/export/download/{export_id}", ("GET",)),
            ("/projects/{project_id}/export/latest", ("GET",)),
            ("/projects/{project_id}/export/history", ("GET",)),
        ]
        for path, methods in expected:
            assert (path, methods) in paths, f"导出端点缺失或方法不符: {path} {methods}"

        preset_paths = {r.path for r in _presets_router.routes}
        assert "/export/presets" in preset_paths, (
            "前端 getExportPresets() 依赖 GET /api/v1/export/presets，缺失则预设选择器整体失效"
        )

    def test_router_mounted_on_main_app(self):
        """main.py 必须同时挂载项目级导出路由与全局预设路由（跨模块调用链路）。

        注：当前 FastAPI 版本 include_router 为惰性挂载（app.routes 中出现
        _IncludedRouter 占位，无 path 属性），改用 OpenAPI schema 断言真实暴露路径。
        """
        from app.main import app

        paths = set(app.openapi()["paths"].keys())
        assert "/api/v1/projects/{project_id}/export/docx-stream" in paths
        assert "/api/v1/export/presets" in paths

    def test_celery_export_task_registered_with_group_lock(self):
        """后台任务挂接：Celery 导出任务存在且与同名源码契约一致。"""
        from app.tasks import generate_docx_export

        assert generate_docx_export.name == "tasks.export.generate_docx"
        import inspect

        from app import tasks as tasks_mod

        src = inspect.getsource(tasks_mod.generate_docx_export)
        assert "export_to_docx_recorded" in src, (
            "Celery 路径必须复用 build_docx 的导出记录，禁止自行补插残缺 exports 行"
        )
        assert "_task_lock.acquire" in src and "'export'" in src, (
            "Celery 路径必须持 export 组锁，与同步/SSE 路径互斥"
        )
