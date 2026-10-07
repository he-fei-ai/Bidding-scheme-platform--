"""DOCX 导出模块回归测试（2026-09-16）

覆盖本次修复的两类缺陷：
- F1 (export.py::_set_cached_export_path)：缓存索引绕过原子写入函数 `_save_export_cache_index`
      直接用 `open("w")` 截断写，多 worker / Celery 并发时会把 `_cache_index.json`
      截成半个 JSON，启动时孤儿清理还会据此误删仍在用的缓存 DOCX。
- F2 (export.py::_add_table_from_markup)：`col_count` 被 `_MAX_TABLE_COLS`(50) 截断后，
      `display_headers` 仍按原始列数占位，随后 `_style_table_header` 按 display_headers
      长度遍历会越界访问 `row.cells[i]`（i>=cols）抛 IndexError，导致超宽表格（>50 列）
      导出崩溃。

只测行为不测实现；所有用例不依赖真实 MinerU / 外部网络 / 数据库。
"""

from __future__ import annotations

import json

from docx import Document

import app.services.export as export_module
from app.models import ExportConfig
from app.services.export import _add_table_from_markup, _set_cached_export_path, _get_cached_export_path


# ==================== F2: 超宽表格（>50 列）不崩溃 ====================


def _make_wide_table_markup(cols: int) -> str:
    header = "| " + " | ".join(f"列{i}" for i in range(1, cols + 1)) + " |"
    row = "| " + " | ".join(f"值{i}" for i in range(1, cols + 1)) + " |"
    return f"<<表格：超宽示例\n{header}\n{row}\n>>"


def test_wide_table_over_max_cols_does_not_crash():
    """60 列表格必须被截断为 50 列且成功建表，不抛 IndexError"""
    doc = Document()
    config = ExportConfig()
    # 不应抛异常（原实现会越界访问 row.cells[i]）
    _add_table_from_markup(doc, _make_wide_table_markup(60), config)

    assert len(doc.tables) == 1
    table = doc.tables[-1]
    # 列数应被 _MAX_TABLE_COLS(50) 限制
    assert len(table.columns) == 50
    # 表头第 50 列（索引 49）必须存在且可访问
    assert table.rows[0].cells[49] is not None
    # 表头文本应被限制在前 50 列
    texts = [c.text for c in table.rows[0].cells]
    assert texts[0] == "列1"
    assert texts[49] == "列50"


def test_wide_table_header_text_truncated_consistently():
    """列数截断后，表头文本也应同步截断，避免出现空表头越界"""
    doc = Document()
    config = ExportConfig()
    _add_table_from_markup(doc, _make_wide_table_markup(60), config)
    table = doc.tables[-1]
    # 表头单元格数量必须等于实际列数（验证 _style_table_header 未越界）
    assert len(table.rows[0].cells) == len(table.columns)


# ==================== F1: 导出缓存索引原子写入 ====================


def _reset_cache_globals(monkeypatch, cache_dir):
    """把缓存相关模块全局指向临时目录，保证测试隔离"""
    index_path = cache_dir / "_cache_index.json"
    monkeypatch.setattr(export_module, "_EXPORT_RESULT_CACHE_DIR", cache_dir)
    monkeypatch.setattr(export_module, "_EXPORT_CACHE_INDEX_PATH", index_path)
    monkeypatch.setattr(export_module, "_EXPORT_CACHE", {})
    monkeypatch.setattr(export_module, "_EXPORT_CACHE_LOADED", False)
    return index_path


def test_set_cached_export_path_persists_valid_index(tmp_path, monkeypatch):
    """_set_cached_export_path 必须把映射原子持久化，可被 _get_cached_export_path 读回"""
    index_path = _reset_cache_globals(monkeypatch, tmp_path)
    key = "a" * 32
    out_path = str(tmp_path / "out.docx")

    _set_cached_export_path(key, out_path)

    # 磁盘索引必须是合法 JSON 且包含本次写入
    assert index_path.exists()
    data = json.loads(index_path.read_text(encoding="utf-8"))
    assert data.get(key) == out_path
    # 读回一致
    assert _get_cached_export_path(key) == out_path


def test_set_cached_export_path_eviction_keeps_valid_index(tmp_path, monkeypatch):
    """写入超过容量上限后索引仍合法 JSON，且条目数被限制（不会损坏）"""
    index_path = _reset_cache_globals(monkeypatch, tmp_path)
    monkeypatch.setattr(export_module, "_EXPORT_CACHE_MAX_ENTRIES", 3)

    for i in range(6):
        key = f"key_{i:028d}"  # 32 字符，模拟真实缓存 key
        _set_cached_export_path(key, str(tmp_path / f"o{i}.docx"))

    # 索引文件始终可读、合法
    data = json.loads(index_path.read_text(encoding="utf-8"))
    assert isinstance(data, dict)
    # 超过容量上限（3）后不应无限增长
    assert len(data) <= 3
