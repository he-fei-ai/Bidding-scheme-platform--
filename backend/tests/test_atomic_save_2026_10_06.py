"""_atomic_save_docx 原子写盘契约测试（2026-10-06）。

钉住 GAP-A 修复的不变量：
1. 正常路径：写满后原子替换为最终路径，且不留 .tmp 临时件；
2. save 中途失败：异常向上抛，最终路径不被污染，.tmp 被清理；
3. 零字节半成品：拒绝替换，最终路径不生成，.tmp 被清理。

不依赖 DB / 项目数据，直接用临时目录 + 真假 Document 双重验证。
"""
import os
from pathlib import Path

import pytest
from docx import Document

from app.services.export import _atomic_save_docx


def test_atomic_save_happy_path_no_tmp_left(tmp_path: Path):
    """正常保存：最终 .docx 生成、可读、同目录无 .tmp 残留"""
    out = tmp_path / "out.docx"
    doc = Document()
    doc.add_paragraph("hello atomic save")

    _atomic_save_docx(doc, str(out))

    assert out.exists(), "最终 .docx 应已生成"
    assert out.stat().st_size > 0
    # 回读校验文件完整（不是截断件）
    reloaded = Document(str(out))
    assert any("hello atomic save" in p.text for p in reloaded.paragraphs)
    # 无 .tmp 残留
    leftovers = list(tmp_path.glob("*.tmp.*"))
    assert leftovers == [], f"不应残留临时文件: {leftovers}"


def test_atomic_save_failure_does_not_pollute_final(tmp_path: Path):
    """save 抛错：异常上抛、最终路径不生成、.tmp 被清理"""
    out = tmp_path / "will_fail.docx"

    class BoomDoc:
        def save(self, _path):
            raise OSError("simulated disk full")

    with pytest.raises(OSError, match="simulated disk full"):
        _atomic_save_docx(BoomDoc(), str(out))

    assert not out.exists(), "save 失败时最终路径不应生成半成品"
    leftovers = list(tmp_path.glob("*.tmp.*"))
    assert leftovers == [], f"失败后 .tmp 必须被清理: {leftovers}"


def test_atomic_save_zero_byte_rejected(tmp_path: Path):
    """save 写出零字节：拒绝原子替换，最终路径不生成、.tmp 清理"""
    out = tmp_path / "zero.docx"

    class EmptyDoc:
        def save(self, path):
            Path(path).write_bytes(b"")

    with pytest.raises(RuntimeError, match="大小为 0"):
        _atomic_save_docx(EmptyDoc(), str(out))

    assert not out.exists(), "零字节半成品不得替换为最终路径"
    leftovers = list(tmp_path.glob("*.tmp.*"))
    assert leftovers == [], f"零字节拒绝后 .tmp 必须被清理: {leftovers}"
