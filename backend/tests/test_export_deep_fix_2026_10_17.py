"""【导出文档模块】第五轮深度审计与修复（2026-10-17）

基于第四轮（2026-10-16）审计后的残留项与新发现，钉住本轮修复行为：

| 编号 | 优先级 | 位置 | 修复要点 |
|------|--------|------|----------|
| FIX-A | P1 | services/export.py:`_add_image_placeholder` | 失败占位符剥离「图X-Y」前缀，避免与回滚复用的真图产生**重复图号** |
| FIX-B | P2 | services/export.py:`_EXPORT_ARTIFACT_SUFFIXES` | 孤儿清理白名单补齐 `.dotx`，与 routers/export.py 下载白名单收敛一致 |
| FIX-C | P2 | services/export.py:`_atomic_save_docx` | Windows 目标文件被占用时 `os.replace` 抛 PermissionError → 重试 + 独立文件名兜底，导出不因占用整体失败 |
| FIX-D | P3 | services/export.py:`build_docx` | 改为先写 exports 表、成功后再登记缓存索引，消除「缓存有记录、DB 无记录」幽灵缓存 |
| R-3   | P3 | services/export_chart_pipeline.py 图表 hash 缓存 | 补齐 LRU 容量淘汰路径测试（此前仅验证条目为 bytes 类型） |

全部先复现后修复，测试钉住修复后行为；运行：`py -3 -m pytest tests/test_export_deep_fix_2026_10_17.py -v`
"""
from __future__ import annotations

import io
import os
from pathlib import Path
from unittest.mock import MagicMock

import pytest
from docx import Document

from app.db import get_db, init_db
from app.models import ExportConfig
from app.services import export as export_mod
from app.services import export_chart_pipeline as chart_pipeline_mod


OWNER = 4401


def _new_project(name: str) -> int:
    init_db()
    with get_db() as conn:
        conn.execute(
            "INSERT INTO projects (name, industry, created_by) VALUES (?, ?, ?)",
            (name, "通用", OWNER),
        )
        return conn.execute(
            "SELECT id FROM projects ORDER BY id DESC LIMIT 1"
        ).fetchone()["id"]


# ============================================================================
# FIX-A —— 失败占位符不得携带图号（杜绝重复图号 / 图号虚跳）
# ============================================================================


class TestFailurePlaceholderStripsFigureNumber:
    """FIX-A：失败占位符（图片/图表插入失败兜底）必须剥离「图X-Y」前缀。

    背景：_add_inline_chart 取号后若 _render_commercial._emit 插入失败，
    旧实现会带着「图3-1」把占位说明写进文档，而图号随后被回滚复用，
    下一张真图重新取得「图3-1」→ 正文出现两张同号图。
    """

    def test_placeholder_strips_fig_no_prefix(self):
        """传入含图号前缀的 caption，占位符落盘文本不得包含「图3-1」。"""
        cfg = ExportConfig()
        doc = Document()
        export_mod._add_image_placeholder(
            doc, cfg, "图3-1 施工进度计划", reason="图表图片插入失败"
        )
        # 第一个段落是 reason（灰字），第二个段落才是图题
        cap_para = doc.paragraphs[-1]
        assert "图3-1" not in cap_para.text, (
            f"失败占位符仍携带图号，将导致重复图号：{cap_para.text!r}"
        )
        assert cap_para.text == "施工进度计划", (
            f"图号剥离后应仅剩标题文本，实际：{cap_para.text!r}"
        )

    def test_placeholder_without_fig_no_unchanged(self):
        """无图号前缀的 caption 行为不变（向后兼容）。"""
        cfg = ExportConfig()
        doc = Document()
        export_mod._add_image_placeholder(doc, cfg, "普通图片说明", reason="x")
        assert doc.paragraphs[-1].text == "普通图片说明"

    def test_reason_marker_still_present(self):
        """reason 灰字段落照常输出，仅图题剥离图号。"""
        cfg = ExportConfig()
        doc = Document()
        export_mod._add_image_placeholder(doc, cfg, "图2-3 组织架构", reason="图表图片插入失败")
        reason_para = doc.paragraphs[-2]
        assert "图表图片插入失败" in reason_para.text
        assert "图2-3" not in doc.paragraphs[-1].text


# ============================================================================
# FIX-B —— 孤儿清理白名单补齐 .dotx
# ============================================================================


class TestOrphanSuffixWhitelistDotx:
    """FIX-B：白名单与 routers/export.py 下载白名单 {docx,dotx,doc,html,pdf} 收敛。"""

    def test_is_safe_artifact_accepts_dotx(self, tmp_path, monkeypatch):
        monkeypatch.setattr(export_mod, "PROJECTS_DATA_DIR", tmp_path)
        (tmp_path / "1" / "exports").mkdir(parents=True)
        p = tmp_path / "1" / "exports" / "tpl.dotx"
        p.write_bytes(b"PK\x03\x04 fake dotx")
        assert export_mod._is_safe_export_artifact(p) is True

    def test_iter_finds_dotx(self, tmp_path, monkeypatch):
        monkeypatch.setattr(export_mod, "PROJECTS_DATA_DIR", tmp_path)
        exports_dir = tmp_path / "1" / "exports"
        exports_dir.mkdir(parents=True)
        (exports_dir / "a.docx").write_bytes(b"fake")
        (exports_dir / "b.dotx").write_bytes(b"PK fake")
        (exports_dir / "c.txt").write_text("skip", encoding="utf-8")
        found = export_mod._iter_export_artifacts()
        names = sorted(p.name for p in found)
        assert names == ["a.docx", "b.dotx"], f"实际: {names}"


# ============================================================================
# FIX-C —— os.replace 占用失败重试 + 独立文件名兜底
# ============================================================================


class TestAtomicSaveReplaceFallback:
    """FIX-C：Windows 目标被占用时 os.replace 抛 PermissionError，
    重试后仍失败则降级写入带时间戳的独立 .docx（仍可被孤儿清理识别）。
    """

    def test_returns_target_path_on_success(self, tmp_path):
        """正常场景：返回最终路径 == 传入 output_path（无重试、无改名）。"""
        out = str(tmp_path / "normal.docx")
        doc = Document()
        doc.add_paragraph("hello")
        final = export_mod._atomic_save_docx(doc, out)
        assert final == out
        assert Path(out).exists()
        reopened = Document(out)
        assert reopened.paragraphs[0].text == "hello"

    def test_fallback_when_target_locked(self, tmp_path, monkeypatch):
        """目标被持续占用：三次重试(占用)后降级写入 .gen<ts>.docx，返回独立路径。"""
        out = str(tmp_path / "locked.docx")
        real_replace = os.replace
        state = {"attempts": 0}

        def fake_replace(src, dst):
            # 凡指向被占用目标一律拒绝（对应重试循环中的 3 次 attempt）
            if dst == out:
                state["attempts"] += 1
                raise PermissionError("file locked by another process")
            # 降级到达独立文件名：真正落盘
            return real_replace(src, dst)

        monkeypatch.setattr(os, "replace", fake_replace)
        doc = Document()
        doc.add_paragraph("fallback content")
        final = export_mod._atomic_save_docx(doc, out)

        # 返回路径不同于被占用目标，且仍以 .docx 结尾（可被孤儿清理识别）
        assert final != out
        assert final.endswith(".docx")
        assert Path(final).exists()
        # 被占用目标本身不应被创建
        assert not Path(out).exists()
        # 重试 3 次 + 1 次兜底 = 至少 3 次命中占用分支
        assert state["attempts"] >= 3
        reopened = Document(final)
        assert reopened.paragraphs[0].text == "fallback content"

    def test_rejects_none_output_path(self):
        """X-5 既有兜底不被破坏：None/空串仍抛 TypeError。"""
        doc = Document()
        doc.add_paragraph("x")
        with pytest.raises(TypeError):
            export_mod._atomic_save_docx(doc, None)
        with pytest.raises(TypeError):
            export_mod._atomic_save_docx(doc, "")


# ============================================================================
# FIX-D —— 缓存索引与 DB 写入原子性（先 DB 后缓存）
# ============================================================================


class TestCacheDbWriteOrdering:
    """FIX-D：build_docx 先写 exports 表，成功后才登记缓存索引。

    旧实现先 set_cached 再 record，record 失败会留下「缓存有、DB 无」幽灵缓存项。
    """

    def _seed_section(self, pid: int):
        with get_db() as conn:
            conn.execute(
                "INSERT INTO sections (project_id, title, level, sort_order, content, status) "
                "VALUES (?, '第一节 施工组织设计', 1, 0, ?, 'completed')",
                (pid, "本项目施工组织设计依据国家标准编制，内容包括总体部署与进度计划安排。"),
            )

    def test_cache_set_only_when_record_succeeds(self, monkeypatch):
        """DB 写入成功（export_id 非空）→ 缓存索引被登记。"""
        pid = _new_project("FIXD-成功")
        self._seed_section(pid)
        cache_calls = []

        def _fake_set(cache_key, file_path):
            cache_calls.append((cache_key, file_path))

        monkeypatch.setattr(export_mod, "_set_cached_export_path", _fake_set)
        # _record_export 保持真实实现（写入 DB）
        out, export_id = export_mod.build_docx(pid, config=ExportConfig())
        assert out
        assert export_id is not None
        assert len(cache_calls) == 1, "DB 写入成功后必须登记缓存索引"

    def test_cache_skipped_when_record_fails(self, monkeypatch):
        """DB 写入失败（抛异常）→ 缓存索引不被登记（无幽灵缓存）。"""
        pid = _new_project("FIXD-失败")
        self._seed_section(pid)
        cache_calls = []

        def _fake_set(cache_key, file_path):
            cache_calls.append((cache_key, file_path))

        def _boom(*args, **kwargs):
            raise RuntimeError("db write failed")

        monkeypatch.setattr(export_mod, "_record_export", _boom)
        monkeypatch.setattr(export_mod, "_set_cached_export_path", _fake_set)
        # 导出文件仍应落盘成功（仅 DB 记录失败）
        out, export_id = export_mod.build_docx(pid, config=ExportConfig())
        assert out
        assert export_id is None
        assert cache_calls == [], "DB 写入失败时不得登记缓存索引（避免幽灵缓存）"


# ============================================================================
# R-3 —— 图表 hash 缓存 LRU 容量淘汰路径
# ============================================================================


class TestChartHashCacheLruEviction:
    """R-3：_CHART_HASH_CACHE_MAX_SIZE 限制下，超量写入时淘汰最旧条目。"""

    def test_lru_eviction_on_overflow(self, monkeypatch):
        chart_pipeline_mod.clear_chart_hash_cache()
        monkeypatch.setattr(chart_pipeline_mod, "_CHART_HASH_CACHE_MAX_SIZE", 2)
        codes = ["flowchart TD\n A-->B", "graph LR\n C-->D", "gantt\n title t"]
        for c in codes:
            chart_pipeline_mod._cache_chart_render(c, b"PNGDATA")
        # 容量上限 2：写入 3 条后应淘汰最旧 1 条，仅剩 2 条
        with chart_pipeline_mod._CHART_HASH_CACHE_LOCK:
            size = len(chart_pipeline_mod._CHART_HASH_CACHE)
        assert size == 2, f"LRU 淘汰后缓存大小应为 2，实际 {size}"
        # 最旧条目（第一条 code）应已被淘汰
        assert chart_pipeline_mod._get_cached_chart_render(codes[0]) is None
        # 最新两条仍在
        assert chart_pipeline_mod._get_cached_chart_render(codes[2]) is not None
        chart_pipeline_mod.clear_chart_hash_cache()
