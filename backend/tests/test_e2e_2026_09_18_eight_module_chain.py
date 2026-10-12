"""八模块端到端数据链贯通验证（2026-09-18）

链路：
  ① 文件导入 → ② 文件解析 → ③ 提取项目 → ④ 目录生成
  → ⑤ 全局事实 → ⑥ 正文生成 → ⑦ 审核预检 → ⑧ 导出文档

设计原则：
  1. **不依赖真实 AI 服务**：所有 AI 调用点注入替身，本文件验证的是模块之间的
     **数据契约**（字段名 / 类型 / 编码 / 落库列 / 参数是否真的传到了下一环），
     而不是模型输出质量；
  2. **不依赖外部中间件**：Redis 等可选依赖缺失时链路必须自行降级；
  3. 全程使用 conftest 提供的隔离临时数据目录，不触碰真实数据。

这个文件存在的意义：各模块自身的单测只能证明「模块内部是对的」，
无法证明「A 模块的产出确实被 B 模块消费」。历史 BUG 大多出现在接缝上
（例如 `tender_context` 一路传到 `generate_outline_v51` 却是死参数、
`scoring_points` 以 list 形态传进按 str 处理的函数）。
"""

from __future__ import annotations

import asyncio
import json
from pathlib import Path
from unittest.mock import MagicMock

import pytest


def _new_project(name: str) -> int:
    """在当前隔离测试库中创建项目并返回 id。"""
    from app.db import get_db, init_db

    init_db()
    with get_db() as conn:
        conn.execute("INSERT INTO projects (name, industry) VALUES (?, ?)", (name, "市政工程"))
        return conn.execute("SELECT id FROM projects ORDER BY id DESC LIMIT 1").fetchone()["id"]


TENDER_TEXT = (
    "第一章 招标公告\n"
    "本项目为市政道路工程，工期：450 天。\n"
    "投标人须具备市政公用工程施工总承包一级资质。\n"
    "第二章 评标办法\n"
    "评分标准：施工组织设计 30 分，质量保证措施 20 分，安全文明施工 15 分。\n"
    "技术要求：路面基层采用水泥稳定碎石，压实度不低于 98%。\n"
)

SCORING_ITEMS = [
    {"item": "施工组织设计", "score": 30},
    {"item": "质量保证措施", "score": 20},
    {"item": "安全文明施工", "score": 15},
]

GLOBAL_FACTS = [
    {"key": "duration", "label": "工期", "value": "450 日历天", "category": "schedule"},
    {"key": "qualification", "label": "资质要求", "value": "市政公用工程施工总承包一级", "category": "qualification"},
]

SECTION_BODY = (
    "本工程为市政道路工程，工期 450 日历天。\n\n"
    "施工准备阶段组建项目经理部，投入管理人员 12 人、作业人员 60 人。\n\n"
    "路面基层采用水泥稳定碎石，压实度不低于 98%，施工后按规范检测。\n\n"
    "质量保证体系执行三检制：自检、互检、专检，不合格工序不得进入下道工序。\n\n"
    "安全文明施工执行三级配电两级保护，临边洞口设置防护栏杆与警示标志。"
)


# ============================================================================
# 主链路：① → ⑧ 一气贯通
# ============================================================================


class TestEightModuleChain:
    def test_end_to_end_chain(self, tmp_path):
        from app.db import get_db
        from app.routers.sse import sse_db_ops, sse_utils
        from app.services import files as files_mod
        from app.services.ai import workflows_outline_v51 as v51
        from app.services.ai.workflows_quality import score_section_quality_local
        from app.services.ai.workflows_sections import build_chapter_content_messages
        from app.services.consistency_auditor import ConsistencyAuditor
        from app.services.export import build_docx

        pid = _new_project("端到端-八模块贯通")

        # ------------------------------------------------------------------
        # ① 文件导入
        # ------------------------------------------------------------------
        saved = files_mod.save_upload(
            pid, "招标文件.txt", TENDER_TEXT.encode("utf-8"), "main"
        )
        assert saved["status"] == "success", f"上传失败: {saved}"
        file_id = saved["id"]

        # ------------------------------------------------------------------
        # ② 文件解析
        # ------------------------------------------------------------------
        with get_db() as conn:
            task = conn.execute(
                "SELECT status, progress FROM file_parse_tasks "
                "WHERE project_id = ? AND filename = ?",
                (pid, "招标文件.txt"),
            ).fetchone()
        assert task is not None, "解析任务行必须存在（失败状态才有落点）"
        assert task["status"] == "completed"

        with get_db() as conn:
            file_rows = [
                dict(r)
                for r in conn.execute(
                    "SELECT * FROM uploaded_files WHERE project_id = ? ORDER BY id", (pid,)
                ).fetchall()
            ]
        main_file = files_mod._pick_main_file(file_rows)
        assert main_file is not None and main_file["id"] == file_id
        parsed_text = files_mod._read_text_auto(main_file["extracted_text_path"])
        assert "450" in parsed_text and "施工组织设计" in parsed_text, (
            "解析产出的文本必须可供下游读取"
        )

        # ------------------------------------------------------------------
        # ③ 提取项目（概述 + 评分点 + 全局事实落 analysis）
        # ------------------------------------------------------------------
        with get_db() as conn:
            conn.execute(
                "INSERT INTO analysis (project_id, summary_json, scoring_points_json, "
                "global_facts_json, confirmed_summary, confirmed_scoring) "
                "VALUES (?, ?, ?, ?, 1, 1) "
                "ON CONFLICT(project_id) DO UPDATE SET "
                "summary_json = excluded.summary_json, "
                "scoring_points_json = excluded.scoring_points_json, "
                "global_facts_json = excluded.global_facts_json",
                (
                    pid,
                    json.dumps({"summary": "市政道路工程，工期 450 天"}, ensure_ascii=False),
                    json.dumps(SCORING_ITEMS, ensure_ascii=False),
                    json.dumps(GLOBAL_FACTS, ensure_ascii=False),
                ),
            )

        project_data = asyncio.run(sse_utils._fetch_project_data_async(pid))
        assert project_data is not None, "项目数据必须可读（下游全链路依赖它）"
        # 类型契约：scoring_points 必须是 list，不能是字符串 "[]"
        assert isinstance(project_data["scoring_points"], list)
        assert project_data["scoring_points"], "评分点必须被解析出来"
        # 接缝关键：招标上下文必须由上传文件解析而来（非空）
        assert project_data["tender_context"], (
            "tender_context 为空 → 目录生成与招标文件脱节"
        )
        assert "450" in project_data["tender_context"]
        assert project_data["global_facts_items"], "全局事实必须随项目数据下发"

        # ------------------------------------------------------------------
        # ④ 目录生成（tender_context 必须真正进入 L1 提示词）
        # ------------------------------------------------------------------
        captured: dict = {}

        async def _fake_l1(provider, summary, sp_text, reference_json, **kwargs):
            captured["l1_user"] = kwargs
            captured["l1_args"] = (summary, sp_text)
            return [
                {
                    "id": "1",
                    "title": "施工组织设计",
                    "description": "总体组织",
                    "children": [],
                }
            ]

        async def _fake_l2(provider, level1, summary, sp_text, **kwargs):
            captured["l2_tender"] = kwargs.get("tender_context")
            return level1

        async def _fake_l3(provider, level1, summary, sp_text, **kwargs):
            return level1

        async def _fake_self_check(provider, level1, *args, **kwargs):
            return level1

        # 直接把提示词构造拉出来验证：tender_context 进入了 L1 的 user prompt
        tender_block = v51._format_tender_context_block(
            project_data["tender_context"], project_data["industry_profile"]
        )
        assert "<tender_context>" in tender_block and "450" in tender_block

        import unittest.mock as _mock

        orig_l1, orig_l2, orig_l3, orig_check = (
            v51._generate_level1_v51,
            v51._generate_level2_v51,
            v51._generate_level3_v51,
            v51._self_check_v51,
        )
        v51._generate_level1_v51 = _fake_l1
        v51._generate_level2_v51 = _fake_l2
        v51._generate_level3_v51 = _fake_l3
        v51._self_check_v51 = _fake_self_check
        try:
            outline = asyncio.run(
                v51.generate_outline_v51(
                    provider=MagicMock(),
                    summary="市政道路工程，工期 450 天",
                    scoring_points=json.dumps(SCORING_ITEMS, ensure_ascii=False),
                    industry_profile=project_data["industry_profile"],
                    tender_context=project_data["tender_context"],
                )
            )
        finally:
            v51._generate_level1_v51 = orig_l1
            v51._generate_level2_v51 = orig_l2
            v51._generate_level3_v51 = orig_l3
            v51._self_check_v51 = orig_check

        assert outline, "目录生成必须产出节点"
        # 接缝断言：招标文件内容确实被 L1/L2 拿到（修复前是死参数）
        assert captured["l1_user"].get("tender_context") == project_data["tender_context"]
        assert captured["l2_tender"] == project_data["tender_context"]

        saved_outline = sse_db_ops._clear_and_save_outline(pid, outline)
        assert saved_outline and saved_outline[0]["id"] > 0, "目录必须落库并带回整数 id"

        # 正文生成（`_build_leaf_dicts`）与导出渲染（`if not has_children`）共用
        # 同一份契约：**正文只挂在叶子节点上**。所以这里按契约取叶子，
        # 而不是随手取第一个节点 —— 取错会得到「导出无正文」的假阴性。
        with get_db() as conn:
            leaf = conn.execute(
                "SELECT s.id, s.title FROM sections s WHERE s.project_id = ? "
                "AND NOT EXISTS (SELECT 1 FROM sections c WHERE c.parent_id = s.id) "
                "ORDER BY s.sort_order, s.id LIMIT 1",
                (pid,),
            ).fetchone()
        assert leaf is not None, "目录必须至少有一个叶子节点"
        section_id = leaf["id"]
        section_title = leaf["title"]

        # ------------------------------------------------------------------
        # ⑤ 全局事实（随项目数据下发 → 进入正文提示词）
        # ------------------------------------------------------------------
        facts = project_data["global_facts_items"]
        assert any("450" in str(f.get("value", "")) for f in facts), (
            "全局事实必须包含工期事实，否则正文无法锚定关键数据"
        )

        # ------------------------------------------------------------------
        # ⑥ 正文生成（评分点 + 事实 → 提示词 → 质量打分 → 落库）
        # ------------------------------------------------------------------
        with get_db() as conn:
            conn.execute(
                "UPDATE sections SET target_words = ? WHERE id = ?", (3500, section_id)
            )
            sec_row = conn.execute(
                "SELECT * FROM sections WHERE id = ?", (section_id,)
            ).fetchone()

        messages = build_chapter_content_messages(
            dict(sec_row),
            facts=facts,
            scoring_context=project_data["scoring_points"],
            project_scope="市政道路工程施工",
        )
        user_prompt = "\n".join(m.content for m in messages if m.role == "user")
        assert section_title in user_prompt, "章节标题必须进入提示词"
        assert "施工组织设计" in user_prompt, "评分点必须进入提示词"
        assert "[{'item'" not in user_prompt, "提示词不得残留 Python list repr"
        assert "450 日历天" in user_prompt, "全局事实必须进入提示词"
        # 2026-10-05（P-22）契约更新：L2+ 章节目标字数由写作长度模式（缺省
        # standard）统一决定，显式 target_words 仅对 L1/异常层级生效。
        # 按章节层级断言对应的字数区间进入提示词（公式与
        # build_chapter_content_messages 保持一致）。
        from app.services.ai.writing_length_config import level_target, normalize_writing_length
        _ch_level = int(dict(sec_row).get("level") or 3)
        _mode = normalize_writing_length("")
        if _ch_level >= 2:
            _mid = level_target(_mode, _ch_level, 300)
            _expect_target = int(_mid)
            _expect_hi = int(_mid * 1.1)
            _expect_lo = max(50, int(_mid * 0.9))
            # BUG-FIX 2026-10-11：字数格式从「1080-1320」改为
            # 「约 {target} 字左右（硬性上限 {hi} 字，最低 {lo} 字）」
            # —— 模型面对区间字符串倾向超长，改为锚点值+硬性上限后偏差显著降低。
            assert f"约 {_expect_target} 字左右（硬性上限 {_expect_hi} 字，最低 {_expect_lo} 字）" in user_prompt, (
                f"章节字数预算（{_mode} 模式 L{_ch_level}）必须进入提示词"
            )
        else:
            assert "2975-4025" in user_prompt, "L1 章节显式 target_words 应生效"

        quality = score_section_quality_local(SECTION_BODY, section_title, 3500)
        assert 0 <= quality["score"] <= 100 and quality["evaluator"] == "local"

        with get_db() as conn:
            conn.execute(
                "UPDATE sections SET content = ?, status = 'completed' WHERE id = ?",
                (SECTION_BODY, section_id),
            )

        # ------------------------------------------------------------------
        # ⑦ 审核预检（一致性审计跑在真实落库的正文上）
        # ------------------------------------------------------------------
        audit = ConsistencyAuditor().run_full_audit(pid)
        assert audit.status.startswith("completed"), f"审计异常: {audit.status}"
        assert audit.total_sections >= 1, "审计必须能看到落库章节"
        # 正文里只出现「450 日历天」，不应产生工期冲突
        conflict_names = {c.fact_name for c in audit.conflicts}
        assert "工期" not in conflict_names, (
            f"正文与事实一致却报了工期冲突: {audit.conflicts}"
        )

        # ------------------------------------------------------------------
        # ⑧ 导出文档
        # ------------------------------------------------------------------
        out_path = tmp_path / "八模块贯通.docx"
        path, export_id = build_docx(pid, output_path=str(out_path))
        assert Path(path).exists(), "导出必须产出文件"
        assert Path(path).stat().st_size > 1000, "导出文件不应为空壳"

        # 正文真的写进了文档（用 docx 回读校验，而不是只看文件存在）
        from docx import Document

        doc = Document(path)
        full = "\n".join(p.text for p in doc.paragraphs)
        assert "施工组织设计" in full, "导出的文档必须包含章节标题"
        assert "水泥稳定碎石" in full, "导出的文档必须包含正文内容"

    def test_chain_survives_restart_of_section_regeneration(self, tmp_path):
        """接缝回归：重新生成目录不得摧毁已生成的正文（③→④→⑥ 的往返）。"""
        from app.db import get_db
        from app.routers.sse import sse_db_ops

        pid = _new_project("端到端-目录重建保全")
        outline = [
            {
                "id": "1",
                "title": "施工组织设计",
                "description": "",
                "target_words": 3000,
                "children": [],
            }
        ]
        saved = sse_db_ops._clear_and_save_outline(pid, outline)
        sid = saved[0]["id"]
        with get_db() as conn:
            conn.execute(
                "UPDATE sections SET content = ?, status = 'completed' WHERE id = ?",
                (SECTION_BODY, sid),
            )

        # 用户仅微调标题层级后重新生成目录
        saved2 = sse_db_ops._clear_and_save_outline(pid, outline)
        with get_db() as conn:
            row = conn.execute(
                "SELECT content, target_words FROM sections WHERE id = ?",
                (saved2[0]["id"],),
            ).fetchone()
        assert row["content"] == SECTION_BODY, "重新生成目录不得清空正文"
        assert row["target_words"] == 3000, "字数预算必须随章节一起保留"


# ============================================================================
# 全局事实链路专项（②→⑤→⑥→⑦→⑧）
# ENHANCE 2026-09-29（facts-ns）：上面主链路里的事实是“全为真值”的理想形态，
# 不经过占位过滤、策略透传、审计基准与失效重置。本轮修的全部接缝都在这些分支上，
# 因此单独走一条真实函数调用链（上传 → 提取 → 人工编辑 → 正文提示词 →
# 一致性审计 → 自动修复 → 导出），验证“模块产出确实被下一环消费”。
# ============================================================================


class TestFactsChainEndToEnd:
    def test_facts_chain_from_upload_to_export(self, tmp_path, monkeypatch):
        from app.db import get_db
        from app.models import FactsGenerateRequest, GlobalFact, GlobalFactList
        from app.routers import facts as facts_router
        from app.routers.sse import sse_db_ops, sse_utils
        from app.services import files as files_mod
        from app.services.ai import provider_factory, workflows
        from app.services.ai.workflows_sections import build_chapter_content_messages
        from app.services.consistency_auditor import ConsistencyAuditor
        from app.services.export import build_docx

        pid = _new_project("端到端-事实链路")

        # ② 文件导入与解析（真实函数，不替身）
        saved = files_mod.save_upload(pid, "招标文件.txt", TENDER_TEXT.encode("utf-8"), "main")
        assert saved["status"] == "success", f"上传失败: {saved}"

        # ⑤ AI 提取（placeholder 档）：人名给不出值必须留【待填写】，不得编造
        monkeypatch.setattr(provider_factory, "get_ai_config", lambda: object())
        monkeypatch.setattr(provider_factory, "create_provider", lambda _cfg: object())

        async def _fake_collect(_provider, _messages, **_kwargs):
            return [
                {"key": "scheduleDuration", "label": "建设工期", "value": "工期：450天", "category": "schedule"},
                {"key": "projectManager", "label": "项目经理", "value": "【待填写】", "category": "personnel"},
            ]

        monkeypatch.setattr(workflows, "collect_json_response", _fake_collect)
        generated = asyncio.run(
            facts_router.generate_facts(pid, FactsGenerateRequest(missing_value_policy="placeholder"))
        )
        assert generated["missing_value_policy"] == "placeholder"
        assert generated["placeholder_count"] >= 1, "占位项必须回传，前端才能出待填角标"
        with get_db() as conn:
            assert conn.execute(
                "SELECT facts_policy FROM analysis WHERE project_id = ?", (pid,)
            ).fetchone()["facts_policy"] == "placeholder", "策略必须落库，否则 Step02 重跑会覆盖回 fabricate"

        # GET 回看：占位项带 needs_manual_input，中文分类命名空间随体下发
        fetched = asyncio.run(facts_router.get_facts(pid))
        by_key = {i["key"]: i for i in fetched["items"]}
        assert by_key["projectManager"]["needs_manual_input"] is True
        assert "implementation" in fetched["categories"]

        # 埋一份旧审计 → 人工保存后必须失效（否则刷新后看到的还是上一批值的结论）
        with get_db() as conn:
            conn.execute(
                "UPDATE analysis SET facts_audit_json = ?, facts_audit_at = ? WHERE project_id = ?",
                ('{"items": [], "status": "ok"}', "2026-09-29 09:00:00", pid),
            )

        # 人工编辑：补齐项目经理真值（前端忘清角标，服务端以 value 为准）
        asyncio.run(
            facts_router.save_facts(
                pid,
                GlobalFactList(
                    items=[
                        GlobalFact(
                            key="projectManager", label="项目经理", value="张工", needs_manual_input=True
                        ),
                        GlobalFact(key="scheduleDuration", label="建设工期", value="工期：450天"),
                    ]
                ),
            )
        )
        asyncio.run(facts_router.confirm_facts(pid))
        refetched = asyncio.run(facts_router.get_facts(pid))
        assert refetched["last_audit"] is None, "保存后旧审计必须作废"
        assert not [i for i in refetched["items"] if i["key"] == "projectManager" and i.get("needs_manual_input")]
        assert refetched["confirmed"] is True

        # ④ 目录落库（正文与导出都需要章节）
        saved_outline = sse_db_ops._clear_and_save_outline(
            pid,
            [{"id": "1", "title": "施工组织设计", "description": "总体组织", "children": []}],
        )
        assert saved_outline and saved_outline[0]["id"] > 0
        section_id = saved_outline[0]["id"]

        # ⑥ 正文生成读取端：项目数据下发的事实集必须被占位过滤（与 GET 不同口径）
        project_data = asyncio.run(sse_utils._fetch_project_data_async(pid))
        facts = project_data["global_facts_items"]
        assert [f for f in facts if "【待填写" in str(f.get("value", ""))] == [], "占位不得泄漏给正文"
        assert any(str(f.get("key")) == "projectManager" for f in facts), "补齐后的真值必须可注入"

        with get_db() as conn:
            sec_row = conn.execute("SELECT * FROM sections WHERE id = ?", (section_id,)).fetchone()
        assert sec_row is not None, "目录必须落库为章节"
        messages = build_chapter_content_messages(
            dict(sec_row), facts=facts, scoring_context=project_data["scoring_points"],
            project_scope="市政道路工程施工",
        )
        user_prompt = "\n".join(m.content for m in messages if m.role == "user")
        assert "待填写" not in user_prompt, "提示词里不得残留占位串"
        assert "张工" in user_prompt and "450" in user_prompt, "已确认事实必须进入提示词"

        # ⑦ 一致性审计：正文把工期写成 500 天，已确认基准 450 天必须入局
        with get_db() as conn:
            conn.execute(
                "UPDATE sections SET content = ?, status = 'completed' WHERE id = ?",
                ("本章工期：500天，路面基层采用水泥稳定碎石。", section_id),
            )
        auditor = ConsistencyAuditor()
        conflicts = [c for c in auditor.detect_conflicts(pid) if c.fact_name == "工期"]
        assert conflicts, "正文与已确认基准不一致时必须报冲突（旧审计器根本不读全局事实）"
        assert conflicts[0].values[0] == "450", "基准值排最前，供修复建议优先取用"

        suggestion = auditor.suggest_fixes(conflicts[0].conflict_id, pid)
        assert suggestion and "统一为：450" in suggestion, "建议必须以基准为准，而非按章节票数"

        fixed = auditor.apply_autofix(conflicts[0].conflict_id, pid, "450")
        assert fixed["success"] and fixed["changed_sections"] >= 1, f"自动修复未落文: {fixed}"
        with get_db() as conn:
            new_content = conn.execute(
                "SELECT content FROM sections WHERE id = ?", (section_id,)
            ).fetchone()["content"]
        assert "工期：450天" in new_content and "500" not in new_content, "修复必须真写回正文"
        assert [c for c in auditor.detect_conflicts(pid) if c.fact_name == "工期"] == [], "修复后重扫不应再报冲突"

        # ⑧ 导出：修复后的正文进入文档
        out_path = tmp_path / "事实链路.docx"
        path, _export_id = build_docx(pid, output_path=str(out_path))
        from docx import Document

        full = "\n".join(p.text for p in Document(path).paragraphs)
        assert "工期：450天" in full and "500" not in full, "导出必须反映修复后的正文"
