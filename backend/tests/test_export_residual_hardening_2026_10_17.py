"""【导出文档模块】第五轮遗留项加固（2026-10-17，L-2..L-6）

钉住《报告_导出文档模块第五轮深度审计与修复_2026-10-17.md》遗留问题清单的
修复后行为：

| 编号 | 优先级 | 覆盖 |
|------|--------|------|
| L-2 | P3 | `_snapshot_figure_counter` 异常注入：整节点失败回滚不得二次异常（X-8 行为回归钉桩） |
| L-3 | P3 | 预渲染（preload_only=True）绕过共享渲染池（BUG-EXPORT-06 契约回归钉桩，杜绝嵌套双池复活） |
| L-4 | P3 | `_get_cached_export_path`：计数器持锁 + 存在性 + ZIP 完整性校验（损坏/缺失按 miss） |
| L-5 | P4 | 自定义 output_path 超长文件名与自动命名同口径截断 |
| L-6 | P4 | `figure_no_separator` 图题分隔符可配置（默认空格 = 历史口径），失败占位符对冒号图题同样剥离干净 |

运行：`py -3 -m pytest tests/test_export_residual_hardening_2026_10_17.py -v`
"""
from __future__ import annotations

import asyncio
import inspect
import io
import json
import logging
import zipfile
from types import SimpleNamespace

import pytest
from docx import Document
from fastapi import HTTPException

from app.db import get_db, init_db
from app.models import ExportConfig
from app.security.auth import User
from app.services import export as export_mod
from app.services import html_export as html_export_mod
from app.services.export_cache import _FIG_NO_PREFIX_RE


OWNER = 4402


def _new_project(name: str) -> int:
    init_db()
    with get_db() as conn:
        conn.execute(
            "INSERT INTO projects (name, industry, created_by) VALUES (?, ?, ?)",
            (name, "通用", OWNER),
        )
        return conn.execute(
            "SELECT id FROM projects ORDER BY id DESC LIMIT 1"
        ).fetchone()["id"]


def _seed_section(pid: int, title: str = "第一节 施工组织设计") -> None:
    with get_db() as conn:
        conn.execute(
            "INSERT INTO sections (project_id, title, level, sort_order, content, status) "
            "VALUES (?, ?, 1, 0, ?, 'completed')",
            (pid, title, "本项目施工组织设计依据国家标准编制，内容包括总体部署与进度计划安排。"),
        )


def _tiny_png() -> io.BytesIO:
    PILImage = pytest.importorskip("PIL.Image", reason="PIL 不可用时跳过")
    buf = io.BytesIO()
    PILImage.new("RGB", (4, 4), (200, 10, 10)).save(buf, format="PNG")
    buf.seek(0)
    return buf


# ============================================================================
# L-2 —— 快照异常注入：整节点失败回滚不二次异常（X-8 行为回归钉桩）
# ============================================================================


class TestFigureSnapshotExceptionInjection:
    def test_source_keeps_preinit_guard(self):
        """X-8 契约：`_fig_counter_snapshot` 必须在 try 之前兜底初始化。"""
        src = inspect.getsource(export_mod._write_tree_content_with_progress)
        assert "_fig_counter_snapshot: dict[str, int] = {}" in src
        # 兜底初始化必须出现在快照调用之前
        assert src.index("_fig_counter_snapshot: dict[str, int] = {}") < src.index(
            "_snapshot_figure_counter()"
        )

    def test_node_failure_with_raising_snapshot_no_unboundlocal(self, monkeypatch):
        """快照函数抛错 + 章节渲染抛错 → build_docx 正常完成（占位符兜底），
        不得出现 UnboundLocalError 二次异常。"""
        pid = _new_project("L2-快照异常")
        _seed_section(pid)

        def _boom():
            raise RuntimeError("contextvar snapshot boom")

        def _render_boom(*args, **kwargs):
            raise RuntimeError("section render boom")

        monkeypatch.setattr(export_mod, "_snapshot_figure_counter", _boom)
        monkeypatch.setattr(export_mod, "_render_section_charts", _render_boom)

        out, export_id = export_mod.build_docx(pid, config=ExportConfig())
        assert out
        reopened = Document(out)
        all_text = "\n".join(p.text for p in reopened.paragraphs)
        # 章节渲染失败被吞掉并写入占位符（节点级容错），而非 UnboundLocalError 逃逸
        assert "本小节内容渲染异常" in all_text or "施工组织设计" in all_text


# ============================================================================
# L-3 —— 预渲染绕过共享渲染池（BUG-EXPORT-06 契约钉桩）
# ============================================================================


class TestPreloadBypassesSharedRenderPool:
    def test_preload_never_touches_shared_executor(self, monkeypatch):
        """preload_only=True 必须直接调用渲染器；共享池被触碰即失败。
        该契约保证预加载池与渲染池不嵌套（否则 4 线程渲染池会被
        预加载任务占满，正文阶段级联超时——BUG-EXPORT-06 复发）。"""
        executor_touches: list[str] = []

        def _forbidden_executor():
            executor_touches.append("used")
            raise AssertionError("preload_only 路径不得使用共享渲染池")

        def _fake_render(payload, typ, skip_http=False):
            return _tiny_png()

        monkeypatch.setattr(export_mod, "_get_render_executor", _forbidden_executor)
        monkeypatch.setattr(export_mod, "render_mermaid_to_bytes", _fake_render)

        section = SimpleNamespace(
            id=1,
            project_id=None,
            title="施工进度计划",
            content="本节描述施工部署。",
            gantt_json=json.dumps({"type": "gantt", "code": "gantt\n    title t", "title": "t"}),
        )
        ok = export_mod._preload_chart_images_parallel(
            {1: section},
            ExportConfig(),
            max_workers=2,
            _http_service_available=False,
        )
        assert isinstance(ok, int)
        assert ok == 1, "预渲染结果必须如实计数（成功 1 张）"
        assert executor_touches == [], "共享渲染池在预加载路径被意外触碰"


# ============================================================================
# L-4 —— 缓存命中：计数持锁 + 存在性 + ZIP 完整性
# ============================================================================


class TestCacheHitValidation:
    def _reset_counters(self, monkeypatch):
        monkeypatch.setattr(export_mod, "_EXPORT_CACHE_HITS", 0)
        monkeypatch.setattr(export_mod, "_EXPORT_CACHE_MISSES", 0)

    def test_missing_file_counts_as_miss(self, tmp_path, monkeypatch):
        self._reset_counters(monkeypatch)
        monkeypatch.setattr(
            export_mod, "_load_export_cache_index", lambda: {"k": str(tmp_path / "gone.docx")}
        )
        assert export_mod._get_cached_export_path("k") is None
        assert export_mod._EXPORT_CACHE_MISSES == 1
        assert export_mod._EXPORT_CACHE_HITS == 0

    def test_corrupt_file_counts_as_miss(self, tmp_path, monkeypatch):
        """索引有键、文件存在但不是有效 ZIP → 按 miss，调用方将重建覆盖该键。"""
        self._reset_counters(monkeypatch)
        bad = tmp_path / "bad.docx"
        bad.write_bytes(b"PK\x03\x04 truncated-not-a-zip" + b"\x00" * 64)
        monkeypatch.setattr(
            export_mod, "_load_export_cache_index", lambda: {"k": str(bad)}
        )
        assert export_mod._get_cached_export_path("k") is None
        assert export_mod._EXPORT_CACHE_MISSES == 1

    def test_valid_zip_counts_as_hit(self, tmp_path, monkeypatch):
        self._reset_counters(monkeypatch)
        good = tmp_path / "good.docx"
        with zipfile.ZipFile(good, "w") as zf:
            zf.writestr("word/document.xml", "<xml/>")
        monkeypatch.setattr(
            export_mod, "_load_export_cache_index", lambda: {"k": str(good)}
        )
        assert export_mod._get_cached_export_path("k") == str(good)
        assert export_mod._EXPORT_CACHE_HITS == 1
        assert export_mod._EXPORT_CACHE_MISSES == 0

    def test_hit_counters_are_lock_guarded(self):
        """源码契约：计数自增必须发生在 _EXPORT_CACHE_LOCK 之内（L-4①）。"""
        src = inspect.getsource(export_mod._get_cached_export_path)
        assert src.count("_EXPORT_CACHE_LOCK") >= 2, "命中与未命中计数都必须持锁"


# ============================================================================
# L-5 —— 自定义 output_path 超长文件名截断
# ============================================================================


class TestCustomOutputPathTruncation:
    def test_overlong_custom_filename_truncated(self, tmp_path):
        pid = _new_project("L5-超长文件名")
        _seed_section(pid)
        long_stem = "超" * 200
        out = str(tmp_path / f"{long_stem}.docx")
        final_path, _export_id = export_mod.build_docx(pid, output_path=out)
        name = __import__("pathlib").Path(final_path).name
        # stem ≤80 + 扩展名（.docx=5）→ 总长 ≤85；叠加 .tmp.<pid>.<8hex>（约+25）
        # 临时名后缀仍 ~110，远低于 Windows 255 文件名上限
        assert len(name) <= 85, f"截断后文件名仍超长: {len(name)}"
        assert len(name.rsplit(".", 1)[0]) <= 80, "文件名 stem 必须截到 ≤80（与自动命名同口径）"
        assert name.endswith(".docx")
        assert __import__("pathlib").Path(final_path).exists()

    def test_normal_custom_filename_untouched(self, tmp_path):
        """向后兼容：正常长度自定义文件名不被改名。"""
        pid = _new_project("L5-正常文件名")
        _seed_section(pid)
        out = str(tmp_path / "我的标书方案.docx")
        final_path, _ = export_mod.build_docx(pid, output_path=out)
        assert __import__("pathlib").Path(final_path).name == "我的标书方案.docx"


# ============================================================================
# L-6 —— 图题分隔符可配置（figure_no_separator）
# ============================================================================


class TestFigureCaptionSeparator:
    def test_default_space_matches_legacy(self):
        cfg = ExportConfig()
        assert export_mod._figure_caption("图3-1", "施工进度计划", cfg) == "图3-1 施工进度计划"

    def test_colon_separator(self):
        cfg = ExportConfig(figure_no_separator="：")
        assert export_mod._figure_caption("图3-1", "施工进度计划", cfg) == "图3-1：施工进度计划"

    def test_empty_separator_allowed(self):
        cfg = ExportConfig(figure_no_separator="")
        assert export_mod._figure_caption("图3-1", "施工进度计划", cfg) == "图3-1施工进度计划"

    def test_invalid_separator_falls_back_to_space(self):
        cfg = ExportConfig(figure_no_separator="!!")
        assert export_mod._figure_caption("图3-1", "施工进度计划", cfg) == "图3-1 施工进度计划"

    def test_no_fig_no_omits_separator(self):
        cfg = ExportConfig(figure_no_separator="：")
        assert export_mod._figure_caption("", "施工进度计划", cfg) == "施工进度计划"
        assert export_mod._figure_caption("图3-1", "", cfg) == "图3-1"

    def test_prefix_regex_strips_colon_caption(self):
        """失败占位符/失败说明对冒号图题剥离后不残留分隔符。"""
        assert _FIG_NO_PREFIX_RE.sub("", "图3-1：施工进度计划") == "施工进度计划"
        assert _FIG_NO_PREFIX_RE.sub("", "图3-1 施工进度计划") == "施工进度计划"

    def test_placeholder_strips_colon_caption(self):
        """集成：含冒号图题的失败占位符不携带任何图号痕迹（FIX-A × L-6 交叉）。"""
        cfg = ExportConfig(figure_no_separator="：")
        doc = Document()
        export_mod._add_image_placeholder(doc, cfg, "图2-3：组织架构", reason="图表图片插入失败")
        assert doc.paragraphs[-1].text == "组织架构"


# ============================================================================
# R-2 —— 导出缓存索引淘汰必须产生告警日志（可观测性）
# ============================================================================


class TestCacheEvictionWarningLog:
    def test_eviction_emits_warning_log(self, tmp_path, monkeypatch, caplog):
        """索引条目超上限被挤掉时必须有 warning 日志（修复前只有注释没有日志，
        条目静默消失无法从日志侧感知）。磁盘文件仍不删（BUG-FIX 2026-09-18 契约）。"""
        monkeypatch.setattr(export_mod, "_EXPORT_RESULT_CACHE_DIR", tmp_path)
        monkeypatch.setattr(export_mod, "_EXPORT_CACHE_INDEX_PATH", tmp_path / "_cache_index.json")
        monkeypatch.setattr(export_mod, "_EXPORT_CACHE", {})
        monkeypatch.setattr(export_mod, "_EXPORT_CACHE_LOADED", False)
        monkeypatch.setattr(export_mod, "_EXPORT_CACHE_MAX_ENTRIES", 2)

        with caplog.at_level(logging.WARNING):
            for i in range(3):
                p = tmp_path / f"out{i}.docx"
                p.write_bytes(b"PK\x03\x04 placeholder")
                export_mod._set_cached_export_path("k" * 31 + str(i), str(p))

        eviction_logs = [r for r in caplog.records if "淘汰最旧条目" in r.getMessage()]
        assert eviction_logs, "缓存索引淘汰未产生告警日志"
        assert any("文件保留" in r.getMessage() for r in eviction_logs), (
            "日志必须说明磁盘文件保留（历史下载仍有效）"
        )


# ============================================================================
# R-5 —— HTML 产物 0 字节必须判失败并删除空壳（第四轮遗留测试缺口）
# ============================================================================


class TestHtmlExportZeroByteFailure:
    def test_zero_byte_html_is_failed_and_removed(self, tmp_path, monkeypatch):
        """build_html 产出 0 字节文件 → 路由 500 + 产物删除 + failed 记录，
        不把空壳 HTML 当 completed 交给前端下载（BUG-FIX 2026-10-16 行为钉桩）。"""
        pid = _new_project("R5-HTML零字节")
        _seed_section(pid)
        zero_path = tmp_path / "empty.html"

        def _build_zero_byte(project_id, config, output_path=None):
            zero_path.write_bytes(b"")
            return str(zero_path), "deadbeef"

        monkeypatch.setattr(html_export_mod, "build_html", _build_zero_byte)

        with pytest.raises(HTTPException) as ei:
            asyncio.run(
                export_mod_export_html(pid, ExportConfig(), identity=User(OWNER, "owner", "user"))
            )
        assert ei.value.status_code == 500
        assert "为空" in str(ei.value.detail) or "0 字节" in str(ei.value.detail)
        # 空壳产物必须被删除，不能留在 exports/ 目录供下载
        assert not zero_path.exists()


def export_mod_export_html(project_id, config, identity=None):
    """经 routers.export.export_html 路由函数（内联 import build_html，
    monkeypatch 打在 app.services.html_export.build_html 上即可生效）。"""
    from app.routers.export import export_html as _route

    return asyncio.run(_route(project_id, config, identity=identity))
