#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""图表生成与渲染 + 提示词模块 深度增强回归测试（2026-10-17）

本文件锁定本轮 4 类修复 + 提示词安全面，全部用例对应探查脚本发现的真实缺陷：

修复项：
  F1 xychart/journey 被 validate_mermaid 误杀 → 围栏整块静默删除
  F2 菱形节点花括号不配对：repair 不补、validate 不检 → 渲染层才炸、整图丢失
  F3 围栏路径图题「引导语」一级永远不可达 → 图题退化为类型通用名
  F4 chart_validators v3.1 缺口：timeline 日期逆序 / labor 维度不一致 /
     layout 几何重叠与坐标越界（CHART_DATA_FIX 提示词承诺检测但从未生效）

提示词安全面（既有能力锁定，防回归）：
  P1 render_prompt 单遍替换防跨变量注入
  P2 sanitize 元标记中和 + 脱敏（CJK 邻接）
  P3 残留占位符检测 + 变量大小写不敏感校验
"""

import io
import json
import sys
from pathlib import Path
from types import SimpleNamespace

_BACKEND_DIR = Path(__file__).parent.parent
sys.path.insert(0, str(_BACKEND_DIR))

import pytest

from app.services.content_cleaning import (
    parse_content,
    _fence_payload_to_marker_text,
    _extract_fence_caption_hint,
)
from app.services.ai.image_engine import (
    repair_mermaid,
    validate_mermaid,
)
from app.services.chart_validators import (
    validate_labor_data,
    validate_layout_data,
    validate_timeline_data,
)


# ============================================================
# F1: xychart / journey 不再被误杀（P1）
# ============================================================
class TestXyChartJourneySupport:
    def test_xychart_beta_valid(self):
        """xychart-beta（柱状/折线对比图）应通过 validate_mermaid"""
        code = (
            "xychart-beta\n"
            '    title "方案造价对比"\n'
            "    x [方案一, 方案二, 方案三]\n"
            '    y "万元" 100 -- 300\n'
            "    bar [120, 260, 210]\n"
            "    line [120, 260, 210]\n"
        )
        ok, msg = validate_mermaid(code)
        assert ok, f"xychart-beta 应通过校验，实际: {msg}"

    def test_xychart_missing_series_rejected(self):
        """缺 bar/line 数据系列 → 拒绝"""
        ok, msg = validate_mermaid("xychart-beta\n    x [甲, 乙]\n")
        assert not ok, "无数据系列的 xychart 应被拒绝"
        assert "xychart" in msg

    def test_journey_valid(self):
        """journey（用户旅程图）应通过 validate_mermaid"""
        code = (
            "journey\n"
            "    title 投标编制流程\n"
            "    section 准备阶段\n"
            "      获取招标文件: 5: 公告平台\n"
            "      研读评分办法: 4: 标前会\n"
        )
        ok, msg = validate_mermaid(code)
        assert ok, f"journey 应通过校验，实际: {msg}"

    def test_journey_missing_event_rejected(self):
        """无事件定义 → 拒绝"""
        ok, msg = validate_mermaid("journey\n    title 空流程\n    section 阶段\n")
        assert not ok, "无事件的 journey 应被拒绝"
        assert "journey" in msg

    def test_xychart_fence_block_not_deleted(self):
        """F1 核心：xychart 围栏块不得被判坏块删除，应归入 comparison 通道"""
        code = (
            "xychart-beta\n"
            '    title "方案造价对比"\n'
            "    x [方案一, 方案二]\n"
            "    bar [120, 260]\n"
        )
        text = f"```mermaid\n{code}\n```\n后续正文段落。"
        blocks = parse_content(text)
        kinds = [b.type for b in blocks]
        assert "chart_marker" in kinds, f"xychart 围栏不应被删除，实际块: {kinds}"
        marker = next(b for b in blocks if b.type == "chart_marker")
        payload = json.loads(marker.text)
        assert payload["chart_type"] == "comparison"
        assert any(
            b.type == "paragraph" and "后续正文段落" in b.text for b in blocks
        ), "图后正文应保留"

    def test_journey_fence_block_not_deleted(self):
        """F1 核心：journey 围栏块不得被判坏块删除，应归入 flowchart 通道"""
        code = "journey\n    title 流程\n    section S\n      事件一: 5: 描述\n"
        text = f"```mermaid\n{code}\n```\n"
        blocks = parse_content(text)
        kinds = [b.type for b in blocks]
        assert "chart_marker" in kinds, f"journey 围栏不应被删除，实际块: {kinds}"
        payload = json.loads(next(b for b in blocks if b.type == "chart_marker").text)
        assert payload["chart_type"] == "flowchart"

    def test_validate_output_format_consistent_with_validate_mermaid(self):
        """validate_output_format 列为合法的类型，validate_mermaid 必须同口径
        （BUG-FIX-PROMPT-O4 声称一致，此前实际不一致）"""
        from app.services.ai.prompts import validate_output_format

        cases = [
            'xychart-beta\n    x [甲, 乙]\n    bar [1, 2]\n',
            "journey\n    title T\n      事件: 5: 描述\n",
            "timeline\n    2026-01 : 开工\n",
        ]
        for code in cases:
            fmt_ok, _ = validate_output_format(code, "mermaid")
            m_ok, m_msg = validate_mermaid(code)
            if fmt_ok:
                assert m_ok, (
                    f"validate_output_format 放行但 validate_mermaid 拒绝"
                    f"（口径不一致）: {code!r} → {m_msg}"
                )


# ============================================================
# F2: 菱形节点花括号配对（repair + validate）
# ============================================================
class TestCurlyBracePairing:
    def test_repair_appends_missing_close_brace(self):
        """AI 截断产出 B{判断 → repair 在行尾补 }"""
        code = "flowchart TD\n    A[开始] --> B{判断\n    B -->|是| C[结束]"
        fixed = repair_mermaid(code)
        assert fixed.count("{") == fixed.count("}"), f"修复后应配对: {fixed!r}"
        ok, msg = validate_mermaid(fixed)
        assert ok, f"修复后应通过校验: {msg}"

    def test_repair_tracks_step(self):
        """修复步骤应被 diff 记录（供 chart_data_fix 上下文）"""
        code = "flowchart TD\n    A[开始] --> B{判断"
        _, diff = repair_mermaid(code, return_diff=True)
        assert "fix_unmatched_curly_braces" in diff["steps"], (
            f"步骤未记录: {diff['steps']}"
        )

    def test_validate_rejects_unpaired_curly_brace(self):
        """直接校验未修复的不配对代码 → 拦截"""
        code = "flowchart TD\n    A[开始] --> B{判断\n    B --> C[结束]"
        ok, msg = validate_mermaid(code)
        assert not ok, "花括号不配对的流程图应被拦截"
        assert "配对" in msg

    def test_quoted_label_curly_brace_not_false_positive(self):
        """引号内合法花括号不得误伤"""
        code = 'flowchart TD\n    A["{x}"] --> B[结束]\n    B --> C["{y}"]'
        ok, msg = validate_mermaid(code)
        assert ok, f"引号内花括号不应误报: {msg}"

    def test_repaired_truncated_fence_block_kept(self):
        """围栏路径：菱形截断块经 repair+validate 后应保留而非删除"""
        code = "flowchart TD\n    A[开工] --> B{质检\n    B --> C[竣工]"
        marker = _fence_payload_to_marker_text("mermaid", code)
        assert marker is not None, "截断菱形围栏块不应被判坏块删除"
        payload = json.loads(marker)
        assert payload["chart_type"] == "flowchart"
        inner = json.loads(payload["description"])
        assert inner["code"].count("{") == inner["code"].count("}")

    def test_unrepairable_code_still_deleted(self):
        """无法修复的纯文本坏块仍应删除（不放宽）"""
        marker = _fence_payload_to_marker_text("mermaid", "这不是任何 Mermaid 代码")
        assert marker is None, "纯文本坏块应删除"



# ============================================================
# F3: 图题引导语（第三优先级）——围栏路径补齐
# ============================================================
class TestCaptionHint:
    def test_hint_captured_from_leading_sentence(self):
        """围栏紧邻前行若是句尾标点收束的引语 → 捕获为 caption_hint"""
        text = (
            "本项目施工流程如下图所示，先平整场地再基础施工。\n"
            "```mermaid\nflowchart TD\n    A[场地平整] --> B[基础施工]\n```\n"
        )
        marker = next(b for b in parse_content(text) if b.type == "chart_marker")
        payload = json.loads(marker.text)
        assert payload.get("caption_hint", "").startswith("本项目施工流程"), (
            f"引导语未捕获: {payload}"
        )

    def test_hint_absent_after_blank_line(self):
        """不跨越空行抓取（避免抓到无关段落）"""
        text = "上文段落没有句号结尾的隔行内容。\n\n```mermaid\ngraph TD\nA[x] --> B[y]\n```\n"
        marker = next(b for b in parse_content(text) if b.type == "chart_marker")
        assert "caption_hint" not in json.loads(marker.text), "跨空行不应抓取"

    def test_hint_absent_after_heading(self):
        """标题行不作引导语"""
        text = "## 3.1 施工部署\n```mermaid\ngraph TD\nA[x] --> B[y]\n```\n"
        marker = next(b for b in parse_content(text) if b.type == "chart_marker")
        assert "caption_hint" not in json.loads(marker.text)

    def test_hint_absent_without_terminal_punctuation(self):
        """无句尾标点的裸行不作引导语"""
        text = "施工部署总体安排非常详细不再赘述\n```mermaid\ngraph TD\nA[x] --> B[y]\n```\n"
        marker = next(b for b in parse_content(text) if b.type == "chart_marker")
        assert "caption_hint" not in json.loads(marker.text)

    def test_hint_absent_for_overlong_line(self):
        """超过 50 字符的长段不作引导语（大概率是段落而非引语）"""
        long_line = (
            "本项目施工流程非常重要需要详细说明并严格遵守相关规范要求确保"
            "万无一失同时兼顾质量工期与成本三者的平衡关系做到统筹兼顾。"
        )
        assert len(long_line) > 50, "样例必须真实超过 50 字符"
        text = long_line + "\n```mermaid\ngraph TD\nA[x] --> B[y]\n```\n"
        marker = next(b for b in parse_content(text) if b.type == "chart_marker")
        assert "caption_hint" not in json.loads(marker.text)

    def test_hint_helper_rejects_structured_lines(self):
        """直接测 helper 的结构化行拒绝"""
        assert _extract_fence_caption_hint(["- 列表项。", "```mermaid"], 1) == ""
        assert _extract_fence_caption_hint(["正常引导语：", "```mermaid"], 1) == "正常引导语："
        assert _extract_fence_caption_hint(["孤立行。"], 1) == ""  # 越界保护



class TestCaptionHintDocxPriority:
    """真实 DOCX 级图题优先级验证（需求 8）：
    载荷 title > Mermaid title > 引导语 > 类型通用名"""

    def _render(self, content, chapter_no):
        from docx import Document
        from app.models import ExportConfig
        from app.services.export import _add_inline_chart, _reset_figure_counter

        _reset_figure_counter()
        marker = next(b for b in parse_content(content) if b.type == "chart_marker")
        payload = json.loads(marker.text)
        doc = Document()
        cfg = ExportConfig(image_max_width_cm=18.0, image_dpi=300)
        sd = SimpleNamespace(title="测试章", project_id=1)
        ok = _add_inline_chart(
            doc,
            payload["chart_type"],
            payload["description"],
            sd,
            cfg,
            parent_section_number=str(chapter_no),
            caption_hint=payload.get("caption_hint", ""),
        )
        texts = [p.text for p in doc.paragraphs]
        caption = next((t for t in texts if f"图{chapter_no}-1" in t), "")
        return ok, payload, caption

    def test_payload_title_beats_hint(self):
        """载荷 title 胜出，引导语不覆盖"""
        content = (
            "劳动力投入情况如下。\n"
            '```chart-json\n{"type":"labor","title":"劳动力投入计划表",'
            '"phases":["基础","主体"],"categories":["普工","木工"],'
            '"data":[[10,20],[30,40]]}\n```\n'
        )
        ok, payload, caption = self._render(content, "9")
        assert payload.get("caption_hint") == "劳动力投入情况如下。"
        assert ok, "载荷带 title 的劳动力图应渲染成功"
        assert "劳动力投入计划表" in caption, f"载荷 title 应胜出: {caption!r}"
        assert "劳动力投入情况如下" not in caption, f"引导语不应覆盖: {caption!r}"

    def test_hint_used_when_no_other_title(self):
        """载荷无 title 且 Mermaid 无 title 行 → 引导语胜出类型通用名"""
        content = (
            "本项目施工流程如下。\n"
            "```mermaid\ngraph TD\n    A[开工] --> B[竣工]\n```\n"
        )
        ok, _payload, caption = self._render(content, "8")
        assert ok, "无 title 流程图应渲染成功"
        assert "本项目施工流程如下" in caption, (
            f"引导语应作图题（此前退化为类型通用名）: {caption!r}"
        )

    def test_mermaid_title_beats_hint(self):
        """Mermaid title 行胜出引导语（第二优先级）"""
        content = (
            "施工工艺流程说明。\n"
            "```mermaid\nflowchart TD\n    title 钢筋绑扎工艺流程\n"
            "    A[下料] --> B[绑扎]\n```\n"
        )
        ok, _payload, caption = self._render(content, "7")
        assert ok, "带 title 行流程图应渲染成功"
        assert "钢筋绑扎工艺流程" in caption, f"Mermaid title 应胜出: {caption!r}"

    def test_default_type_name_when_nothing_available(self):
        """四级全空 → 类型通用名兜底"""
        content = "```mermaid\ngraph TD\n    A[开工] --> B[竣工]\n```\n"
        ok, payload, caption = self._render(content, "6")
        assert payload.get("caption_hint", "") == "", "前一行是围栏本身，无引导语"
        assert ok, "流程图应渲染成功"
        assert "施工流程图" in caption, f"应兜底为类型通用名: {caption!r}"

    def test_fig_no_format_chapter_sequence(self):
        """图号格式锁定：图{章号}-{序号}"""
        content = (
            "工艺流程如下。\n"
            "```mermaid\ngraph TD\n    A[开工] --> B[竣工]\n```\n"
        )
        ok, _payload, caption = self._render(content, "5")
        assert ok
        assert caption.startswith("图5-1 "), f"图号格式应为「图5-1 图题」: {caption!r}"



# ============================================================
# F4: chart_validators v3.1 规范补齐（提示词承诺的检测项落地）
# ============================================================
class TestTimelineDateOrder:
    def test_reverse_dates_rejected(self):
        """竣工早于开工 → 逆序报错（此前静默通过）"""
        data = {"milestones": [
            {"date": "2026-05-01", "title": "主体封顶", "status": "pending"},
            {"date": "2026-01-01", "title": "项目开工", "status": "pending"},
        ]}
        ok, errs = validate_timeline_data(data)
        assert not ok, "日期逆序应被拦截"
        assert any("逆序" in e for e in errs), f"应报逆序错误: {errs}"

    def test_ascending_dates_pass(self):
        """升序日期通过"""
        data = {"milestones": [
            {"date": "2026-01-01", "title": "开工", "status": "completed"},
            {"date": "2026-06-30", "title": "封顶", "status": "pending"},
        ]}
        ok, errs = validate_timeline_data(data)
        assert ok, f"升序日期应通过: {errs}"

    def test_invalid_dates_skipped_by_order_check(self):
        """非 YYYY-MM-DD 的脏日期不参与排序（避免与格式检查双重报错刷屏）"""
        data = {"milestones": [
            {"date": "2026-13-99", "title": "a", "status": "pending"},
            {"date": "bad-date", "title": "b", "status": "pending"},
        ]}
        ok, errs = validate_timeline_data(data)
        assert not ok
        assert not any("逆序" in e for e in errs), f"脏日期不应报逆序: {errs}"

    def test_fenced_timeline_json_reverse_dates_detected(self):
        """围栏路径（诊断层）也应拿到逆序错误"""
        from app.services.ai.image_engine import validate_chart_payload

        payload = {"type": "timeline", "milestones": [
            {"date": "2026-12-01", "title": "竣工", "status": "pending"},
            {"date": "2026-01-01", "title": "开工", "status": "pending"},
        ]}
        ok2, errs2 = validate_chart_payload("timeline", payload)
        assert not ok2 and any("逆序" in e for e in errs2)


class TestLaborDimensionConsistency:
    def test_row_shorter_than_categories_reported(self):
        """短行静默补 0 导致人数少报 → 现显式报维度不一致"""
        data = {"phases": ["基础"], "categories": ["普工", "木工"], "data": [[10]]}
        ok, errs = validate_labor_data(data)
        assert not ok, "维度不一致应被拦截"
        assert any("维度" in e for e in errs), f"应报维度错误: {errs}"

    def test_row_longer_than_categories_reported(self):
        """长行静默截断导致数据丢失 → 现显式报错"""
        data = {"phases": ["基础"], "categories": ["普工"], "data": [[10, 20, 30]]}
        ok, errs = validate_labor_data(data)
        assert not ok
        assert any("维度不一致" in e for e in errs)

    def test_row_count_mismatch_reported(self):
        """矩阵行数 ≠ 阶段数 → 报错（不抛 IndexError）"""
        data = {"phases": ["基础", "主体"], "categories": ["普工"], "data": [[10]]}
        ok, errs = validate_labor_data(data)
        assert not ok
        assert any("行" in e for e in errs), f"应报行数不一致: {errs}"

    def test_consistent_data_passes(self):
        """行列一致的合法数据通过"""
        data = {
            "phases": ["基础", "主体"],
            "categories": ["普工", "木工"],
            "data": [[10, 20], [30, 40]],
        }
        ok, errs = validate_labor_data(data)
        assert ok, f"行列一致应通过: {errs}"

    def test_all_zero_row_does_not_crash_when_matrix_longer(self):
        """矩阵行多于阶段数且全 0 → 不得 IndexError（旧实现 phases[i] 越界）"""
        data = {"phases": ["基础"], "categories": ["普工"], "data": [[10], [0]]}
        ok, errs = validate_labor_data(data)
        assert not ok, f"应报错（行数不一致），实际 errs={errs}"


class TestLayoutGeometry:
    def test_overlap_reported(self):
        """两区域几何重叠 → 报错（CHART_DATA_FIX 承诺检测项）"""
        data = {"zones": [
            {"id": "Z1", "name": "办公区", "x": 5, "y": 5, "w": 50, "h": 50},
            {"id": "Z2", "name": "堆场", "x": 40, "y": 40, "w": 50, "h": 50},
        ]}
        ok, errs = validate_layout_data(data)
        assert not ok, "重叠区域应被拦截"
        assert any("重叠" in e for e in errs), f"应报重叠错误: {errs}"

    def test_disjoint_zones_pass(self):
        """不重叠区域通过"""
        data = {"zones": [
            {"id": "Z1", "name": "办公区", "x": 5, "y": 5, "w": 25, "h": 20},
            {"id": "Z2", "name": "堆场", "x": 5, "y": 35, "w": 25, "h": 20},
        ]}
        ok, errs = validate_layout_data(data)
        assert ok, f"不重叠区域应通过: {errs}"

    def test_touching_zones_not_overlap(self):
        """边界相接（共享边）不算重叠"""
        data = {"zones": [
            {"id": "Z1", "name": "A", "x": 0, "y": 0, "w": 50, "h": 50},
            {"id": "Z2", "name": "B", "x": 50, "y": 0, "w": 50, "h": 50},
        ]}
        ok, errs = validate_layout_data(data)
        assert ok, f"边界相接不应判重叠: {errs}"

    def test_out_of_range_reported(self):
        """坐标越界（x+w > 100）→ 报错"""
        data = {"zones": [
            {"id": "Z1", "name": "主体区", "x": 60, "y": 10, "w": 55, "h": 30},
        ]}
        ok, errs = validate_layout_data(data)
        assert not ok
        assert any("越界" in e for e in errs), f"应报越界错误: {errs}"

    def test_non_numeric_zone_skipped(self):
        """无数字坐标的描述型区域不参与几何检查（不误伤）"""
        data = {"zones": [
            {"id": "Z1", "name": "A", "x": 0, "y": 0, "w": 50, "h": 50},
            {"id": "Z2", "name": "B"},
        ]}
        ok, errs = validate_layout_data(data)
        assert ok, f"描述型区域不应误伤: {errs}"



# ============================================================
# P1-P3: 提示词模块安全面锁定（防回归）
# ============================================================
class TestPromptSecurity:
    def test_single_pass_blocks_cross_var_injection(self):
        """P1：变量值内含另一变量名 → 不得被二次替换（跨变量注入防御）"""
        from app.services.ai.prompts import render_prompt

        out = render_prompt(
            "A:{GLOBAL_FACTS} B:{SECTION_TITLE}",
            GLOBAL_FACTS="{SECTION_TITLE}",
            SECTION_TITLE="真实标题",
        )
        # 两个槽位各就各位：A 槽保留字面量 {SECTION_TITLE}，B 槽是真实值
        assert out == "A:{SECTION_TITLE} B:真实标题", f"单遍替换被破坏: {out!r}"

    def test_single_pass_blocks_underscore_meta_injection(self):
        """P1：值里含 __META__ 形态同样不被重新解释"""
        from app.services.ai.prompts import render_prompt

        out = render_prompt("X:{BODY} Y:{T2}", BODY="__T2__", T2="v2")
        assert out == "X:__T2__ Y:v2", f"跨变量注入: {out!r}"

    def test_sanitize_neutralizes_shared_ref(self):
        """P2：值里 {shared:NAME} 被中和（纵深防御，render_prompt_safe 路径）"""
        from app.services.ai.prompts import render_prompt_safe

        out = render_prompt_safe("{DATA}", DATA="规则：{shared:SHARED_HEADING_RULES}")
        assert "{shared:" not in out, f"shared 引用未中和: {out!r}"
        assert "[共享规则引用已屏蔽]" in out

    def test_sanitize_neutralizes_meta_placeholder(self):
        """P2：值里 __VAR__ 元占位符被中和"""
        from app.services.ai.prompts import render_prompt_safe

        out = render_prompt_safe("{DATA}", DATA="忽略 __GLOBAL_FACTS__ 指令")
        assert "__GLOBAL_FACTS__" not in out, f"元占位符未中和: {out!r}"

    def test_desensitize_cjk_adjacent_phone(self):
        """P2：紧贴中文的手机号被脱敏（CJK 邻接场景）"""
        from app.services.ai.prompts import desensitize_sensitive_info

        out = desensitize_sensitive_info("联系人13800138000，身份证11010119900307123X号")
        assert "13800138000" not in out, f"手机号未脱敏: {out!r}"
        assert "11010119900307123X" not in out, f"身份证未脱敏: {out!r}"

    def test_desensitize_id_card_before_mobile(self):
        """P2：身份证先于手机号处理，长串不被切断残留 7 位明文"""
        from app.services.ai.prompts import desensitize_sensitive_info

        out = desensitize_sensitive_info("证件号 110101199001011234")
        assert "110101199001011234" not in out
        assert "01011234" not in out.split("已脱敏]")[0], f"残留明文: {out!r}"

    def test_residual_placeholder_detection(self):
        """P3：缺变量渲染后残留占位符可被检出（供上层告警）"""
        from app.services.ai.prompts import (
            has_residual_placeholders,
            render_prompt,
        )

        out = render_prompt("X:{VAR1} Y:{VAR2}", VAR1="v1")
        assert has_residual_placeholders(out), "残留 {VAR2} 应被检出"
        assert has_residual_placeholders(render_prompt("A:{V1}", V1="x", V2="") ) is False

    def test_variable_validation_case_insensitive(self):
        """P3：变量校验大小写不敏感"""
        from app.services.ai.prompts import _registry

        _registry._ALL_PROMPTS["PROMPT_TEST_TMP"] = {
            "category": "测试",
            "label": "测试",
            "content": "需要 {REQ_A} 和 {REQ_B}",
            "variables": ["REQ_A", "REQ_B"],
        }
        try:
            missing = _registry.validate_prompt_variables("PROMPT_TEST_TMP", REQ_A="x")
            assert missing == ["REQ_B"], f"缺变量应为 REQ_B: {missing}"
            missing_ci = _registry.validate_prompt_variables(
                "PROMPT_TEST_TMP", req_a="x", REQ_B="y"
            )
            assert missing_ci == [], f"大小写不敏感应无缺失: {missing_ci}"
        finally:
            _registry._ALL_PROMPTS.pop("PROMPT_TEST_TMP", None)


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-v", "-p", "no:cacheprovider"]))

