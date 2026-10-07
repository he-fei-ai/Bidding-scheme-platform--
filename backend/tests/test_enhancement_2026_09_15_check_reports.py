"""回归测试：检查结果 Excel 多 Sheet 导出 + 查重白名单分级 + 图片上下文定位。

ENHANCE 2026-09-15（第六轮）：移植自「易标投标工具箱」`checkResultExportService.cjs`
与 `duplicateCheckService.cjs` 的能力。

覆盖范围：
  A. Excel 报表数据层（纯函数）：Sheet 结构、误报剔除、概览统计
  B. Excel 落盘层：openpyxl 硬约束（Sheet 名非法字符/重名、单元格截断、``=`` 文本、
     百分比列格式、首行冻结、自动筛选）
  C. 招标原文白名单四级匹配：exact / strict / skeleton / near
  D. 图片重复上下文：docx 章节路径 / 序号 / 前一句；PDF 页码级
  E. 报告下载守卫：白名单根目录 + 路由 format 参数端到端
  F. 路由接线（源码级）：format 校验分支与安全守卫

设计原则：只测行为不测实现；不依赖真实 AI 与外部网络。
"""

from __future__ import annotations

import asyncio
import inspect
from pathlib import Path
import zipfile

import pytest

from app.config import (
    DUPLICATE_REPORTS_DIR,
    PROJECTS_DATA_DIR,
    REJECTION_REPORTS_DIR,
    is_downloadable_path,
    report_download_roots,
)
from app.services import check_report_excel as excel
from app.services import duplicate as dup


# ==================== 公共构造 ====================


def _rejection_task(**overrides) -> dict:
    task = {
        "id": 101,
        "project_id": 7,
        "status": "completed",
        "false_positives": [],
        "bid_documents": [{"name": "投标文件A.docx"}, {"name": "投标文件B.docx"}],
        "results": {
            "rejection": [
                {
                    "id": "i1",
                    "description": "未响应工期实质性条款",
                    "severity": "critical",
                    "quote": "工期为100日历天",
                    "location": "第1章 项目概述",
                    "check_item_id": "r001",
                    "bid_doc_id": 1,
                    "bid_doc_name": "投标文件A.docx",
                    "validated": True,
                    "line_hint": "第 12 行",
                    "suggestion": "补充工期承诺",
                },
                {
                    "id": "i2",
                    "description": "缺少资质证书扫描件",
                    "severity": "warning",
                    "validated": None,
                    "evidence_kind": "structural",
                },
            ],
            "spelling": [
                {
                    "id": "sp-001",
                    "original": "实施",
                    "corrected": "实施",
                    "quote": "本项目实施过程中",
                    "location": "第2段第1句",
                    "validated": True,
                    "line_hint": "第 30 行",
                },
            ],
            "logic": [
                {
                    "id": "lg-001",
                    "description": "总工期与阶段之和不一致",
                    "quote_a": "总工期180天",
                    "quote_b": "阶段合计250天",
                    "location_a": "第1章",
                    "location_b": "第3章",
                    "suggestion": "核对工期",
                },
            ],
        },
    }
    task.update(overrides)
    return task


def _duplicate_task(**overrides) -> dict:
    task = {
        "id": 202,
        "status": "completed",
        "tender_filename": "招标文件.docx",
        "results": {
            "metadata": {
                "metadata": [
                    {
                        "filename": "投标文件A.docx",
                        "author": "张三",
                        "company": "某某公司",
                        "created": "2026-01-01 10:00:00",
                        "modified": "2026-02-02 11:00:00",
                        "last_saved_by": "李四",
                        "title": "技术标",
                    },
                ],
                "anomalies": [{"description": "多个文件作者相同", "value": ["张三"]}],
                "fingerprints": [{"rsidRoot": "00AB12CD"}],
                "same_source_pairs": [
                    {
                        "file_a": "投标文件A.docx",
                        "file_b": "投标文件B.docx",
                        "severity": "high",
                        "evidence": ["共享 rsidRoot"],
                    },
                ],
            },
            "toc": {
                "pairwise": [
                    {
                        "file_a": "投标文件A.docx",
                        "file_b": "投标文件B.docx",
                        "score": 0.88,
                        "title_overlap": 0.9,
                        "path_overlap": 0.9,
                        "order_similarity": 0.8,
                        "shared_count": 12,
                        "risk": "high",
                    },
                ],
                "duplicate_groups": [
                    {
                        "file_a": "投标文件A.docx",
                        "file_b": "投标文件B.docx",
                        "similarity": 0.95,
                        "common_titles": ["施工组织设计"],
                    },
                ],
                "similar_groups": [
                    {
                        "file_a": "投标文件A.docx",
                        "file_b": "投标文件B.docx",
                        "title_a": "施工组织设计",
                        "title_b": "施工组织设计方案",
                        "similarity": 0.86,
                        "criterion": "dice",
                    },
                ],
            },
            "content": {
                "stats": {
                    "total_sentences": 100,
                    "duplicate_count": 1,
                    "similar_count": 1,
                    "tender_whitelist": {"exact": 5, "strict": 3, "skeleton": 1, "near": 2},
                    "tender_whitelist_count": 40,
                },
                "duplicate_groups": [
                    {
                        "sentence": "本项目采用先进的施工工艺",
                        "occurrences": [
                            {"filename": "投标文件A.docx", "sentence_idx": 3},
                            {"filename": "投标文件B.docx", "sentence_idx": 7},
                        ],
                    },
                ],
                "similar_groups": [
                    {
                        "similarity": 0.85,
                        "sentence_a": "甲句内容",
                        "sentence_b": "乙句内容",
                        "file_a": "投标文件A.docx",
                        "file_b": "投标文件B.docx",
                    },
                ],
            },
            "images": {
                "total_images": 4,
                "duplicate_groups": [
                    {
                        "hash": "abcdef123456...",
                        "size": 1234,
                        "occurrences": [
                            {
                                "filename": "投标文件A.docx",
                                "path": "word/media/image1.png",
                                "image_index": 2,
                                "directory": "第1章 / 施工方案",
                                "previous_sentence": "现场平面布置图如下",
                            },
                            {
                                "filename": "投标文件B.docx",
                                "image_index": 1,
                                "directory": "第1章",
                                "previous_sentence": "",
                            },
                        ],
                        "locations": ["投标文件A.docx · 第1章 / 施工方案 · 第2张"],
                    },
                ],
            },
        },
    }
    task.update(overrides)
    return task


# ==================== A. Excel 报表数据层 ====================


def test_rejection_workbook_has_four_expected_sheets():
    sheets = excel.build_rejection_workbook(_rejection_task())
    assert [s.title for s in sheets] == ["检查概览", "废标项", "错别字", "逻辑问题"]


def test_rejection_overview_counts_per_dimension():
    sheets = excel.build_rejection_workbook(_rejection_task())
    overview = sheets[0]
    assert overview.headers == ["检查类型", "状态", "问题数量", "更新时间", "检查对象", "说明"]
    assert len(overview.rows) == 3
    assert [row[0] for row in overview.rows] == ["废标项检查", "错别字检查", "逻辑谬误检查"]
    assert [row[2] for row in overview.rows] == [2, 1, 1]
    assert all(row[1] == "已完成" for row in overview.rows)


def test_rejection_overview_excludes_false_positives_and_notes_it():
    task = _rejection_task(false_positives=["i1", "sp-001"])
    sheets = excel.build_rejection_workbook(task)
    overview = sheets[0]
    assert [row[2] for row in overview.rows] == [1, 0, 1], "问题数量必须按剔除误报后的可见条数统计"
    assert "已剔除人工误报 1 项" in overview.rows[0][5]
    assert "已剔除人工误报 1 项" in overview.rows[1][5]
    assert any("已剔除人工误报：2 项" in note for note in overview.notes)
    # 明细 Sheet 同步剔除
    assert len(sheets[1].rows) == 1
    assert len(sheets[2].rows) == 0


def test_rejection_overview_reports_evidence_verification_breakdown():
    sheets = excel.build_rejection_workbook(_rejection_task())
    joined = "\n".join(sheets[0].notes)
    assert "已在原文定位 1 项" in joined
    assert "结构性线索（材料缺失类，无原文摘录）1 项" in joined


def test_rejection_detail_sheet_columns_and_verification_text():
    sheets = excel.build_rejection_workbook(_rejection_task())
    rows = sheets[1].rows
    assert sheets[1].headers[0] == "序号"
    assert rows[0][0] == 1
    assert rows[0][5] == "严重"
    assert "第 12 行" in rows[0][9]
    # 结构性线索（无原文摘录）要有独立文案，不能谎报"已定位"
    assert "结构性线索" in rows[1][9]
    assert rows[1][7] == ""


def test_logic_sheet_pairs_both_quotes():
    sheets = excel.build_rejection_workbook(_rejection_task())
    rows = sheets[3].rows
    assert rows[0][5] == "总工期180天"
    assert rows[0][7] == "阶段合计250天"
    assert rows[0][-1] == "核对工期"


def test_duplicate_workbook_has_six_expected_sheets():
    sheets = excel.build_duplicate_workbook(_duplicate_task())
    assert [s.title for s in sheets] == [
        "查重概览",
        "元数据",
        "目录重复",
        "文件相似度",
        "重复句子",
        "重复图片",
    ]


def test_duplicate_overview_counts_and_whitelist_note():
    sheets = excel.build_duplicate_workbook(_duplicate_task())
    overview = sheets[0]
    assert [row[0] for row in overview.rows] == ["元数据对比", "目录相似度", "正文重复", "图片重复"]
    assert [row[2] for row in overview.rows] == [1, 2, 2, 1]
    joined = "\n".join(overview.notes)
    assert "招标原文白名单命中" in joined and "exact" in joined
    assert "识别图片总数：4" in joined


def test_pairwise_sheet_declares_percent_columns():
    sheets = excel.build_duplicate_workbook(_duplicate_task())
    pairwise = sheets[3]
    assert pairwise.percent_columns == (3, 4, 5, 6)
    assert pairwise.rows[0][2] == pytest.approx(0.88)
    assert pairwise.rows[0][7] == "高"
    assert any("0.45" in note for note in pairwise.notes), "必须写明综合相似度的权重口径"


def test_toc_sheet_marks_duplicate_and_near_synonym_rows():
    sheets = excel.build_duplicate_workbook(_duplicate_task())
    rows = sheets[2].rows
    assert [row[1] for row in rows] == ["目录重复", "近义改写"]
    assert rows[0][3] == pytest.approx(0.95)
    assert rows[1][3] == pytest.approx(0.86)
    assert sheets[2].percent_columns == (4,)


def test_metadata_sheet_links_same_source_hint_to_both_files():
    sheets = excel.build_duplicate_workbook(_duplicate_task())
    metadata_sheet = sheets[1]
    assert metadata_sheet.rows[0][0] == "投标文件A.docx"
    assert "同源" in metadata_sheet.rows[0][7]
    assert any("共享 rsidRoot" in note for note in metadata_sheet.notes)


def test_content_sheet_separates_exact_and_similar():
    sheets = excel.build_duplicate_workbook(_duplicate_task())
    rows = sheets[4].rows
    assert rows[0][1] == "完全相同"
    assert rows[0][3] == 2
    assert rows[1][1].startswith("高度相似")


def test_image_sheet_uses_context_columns():
    sheets = excel.build_duplicate_workbook(_duplicate_task())
    rows = sheets[5].rows
    assert rows[0][1] == "abcdef123456..."
    assert "投标文件A.docx×1" in rows[0][3]
    assert "第1章 / 施工方案·第2张" in rows[0][4]
    assert "现场平面布置图如下" in rows[0][5]
    assert any("docx 的章节路径" in note for note in sheets[5].notes)


def test_builders_tolerate_empty_and_malformed_results():
    assert len(excel.build_rejection_workbook({})) == 4
    assert len(excel.build_duplicate_workbook({"results": None})) == 6
    # 概览 Sheet 是固定三/四行维度表，明细 Sheet 必须为空
    malformed = excel.build_rejection_workbook({"results": "oops"})
    assert len(malformed[0].rows) == 3
    assert all(sheet.rows == [] for sheet in malformed[1:])
    assert all(sheet.rows == [] for sheet in excel.build_duplicate_workbook({"results": "oops"})[1:])


def test_helpers_normalize_labels_and_ratios():
    assert excel.severity_label("high") == "高"
    assert excel.severity_label("critical") == "严重"
    assert excel.severity_label("unknown") == "警告"
    assert excel.risk_label("medium") == "中"
    assert excel.as_ratio(0.5) == pytest.approx(0.5)
    assert excel.as_ratio("not-a-number") == "not-a-number"
    assert excel.as_ratio(None) == ""
    assert excel.as_ratio(True) == ""


def test_cell_text_truncates_and_flattens():
    assert len(excel.cell_text("甲" * 40000)) == excel.MAX_CELL_CHARS
    assert excel.cell_text(["a", "b"]) == "a；b"
    assert excel.cell_text(True) == "是"
    assert excel.cell_text(None) == ""


# ==================== B. Excel 落盘层 ====================


def test_safe_sheet_name_filters_illegal_chars_and_truncates():
    used: set[str] = set()
    name = excel._safe_sheet_name("概览:项目[1]/2?*", used)
    assert not set(name) & set(":\\/?*[]")
    assert len(name) <= excel.MAX_SHEET_NAME_CHARS
    long_name = excel._safe_sheet_name("甲" * 50, used)
    assert len(long_name) <= excel.MAX_SHEET_NAME_CHARS


def test_safe_sheet_name_dedupes():
    used: set[str] = set()
    first = excel._safe_sheet_name("检查概览", used)
    second = excel._safe_sheet_name("检查概览", used)
    assert first != second
    assert second.startswith("检查概览")


def test_write_workbook_applies_freeze_filter_and_percent_format(tmp_path):
    from openpyxl import load_workbook

    specs = [
        excel.SheetSpec(
            title="测试表",
            headers=["名称", "相似度"],
            rows=[["甲", 0.875], ["乙", 0.5]],
            percent_columns=(2,),
        ),
    ]
    path = excel.write_workbook(specs, tmp_path / "out.xlsx")
    assert path.exists()

    ws = load_workbook(path)["测试表"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:B3"
    assert ws.cell(row=2, column=2).number_format == "0.00%"
    assert ws.cell(row=2, column=2).value == pytest.approx(0.875)
    assert ws.cell(row=1, column=1).font.bold is True


def test_write_workbook_keeps_formula_like_text_as_string(tmp_path):
    """AI 生成的正文可能以 = 开头，绝不能写成 Excel 公式（打开即报错）。"""
    from openpyxl import load_workbook

    specs = [excel.SheetSpec(title="公式", headers=["文本"], rows=[["=SUM(A1:A2)"], ["+1"], ["-1"]])]
    path = excel.write_workbook(specs, tmp_path / "formula.xlsx")

    ws = load_workbook(path)["公式"]
    for row_index in (2, 3, 4):
        cell = ws.cell(row=row_index, column=1)
        assert cell.data_type == "s", f"第{row_index}行被写成了公式"
        assert str(cell.value).startswith(("=", "+", "-"))


def test_write_workbook_requires_content(tmp_path):
    with pytest.raises(ValueError):
        excel.write_workbook([], tmp_path / "empty.xlsx")


def test_write_workbook_autowidth_caps_columns(tmp_path):
    from openpyxl import load_workbook

    specs = [excel.SheetSpec(title="宽", headers=["内容"], rows=[["甲" * 500]])]
    path = excel.write_workbook(specs, tmp_path / "width.xlsx")
    assert load_workbook(path)["宽"].column_dimensions["A"].width <= 60


# ==================== C. 招标原文白名单四级匹配 ====================


def test_skeleton_signature_replaces_numbers_and_units():
    skeleton = dup.skeleton_signature("投标保证金为人民币50万元")
    assert "〇" in skeleton
    assert not any(char.isdigit() for char in skeleton)


def test_skeleton_signature_rejects_too_generic_sentences():
    """骨架过于笼统（如「工期为〇天」）不得作为豁免依据，否则会放过真实重复。"""
    assert dup.skeleton_signature("工期为30天") == ""
    assert dup.skeleton_signature("") == ""
    assert dup.skeleton_signature("没有数字的一句话") == ""


def test_whitelist_exact_and_strict_levels():
    sentence = "投标人应当按照招标文件要求提交完整的资格证明文件"
    whitelist = dup.TenderWhitelist(f"{sentence}。")
    assert whitelist.match_level(sentence) == "exact"
    # 只改标点/空白 → strict
    assert whitelist.match_level("投标人应当按照招标文件要求，提交完整的资格证明文件") == "strict"


def test_whitelist_skeleton_level_allows_number_substitution():
    whitelist = dup.TenderWhitelist("投标保证金为人民币50万元。")
    assert whitelist.match_level("投标保证金为人民币80万元") == "skeleton"


def test_whitelist_near_level_matches_paraphrased_long_sentence():
    tender = "本项目施工过程中严格执行国家和地方现行的安全生产法律法规及标准规范确保施工安全。"
    bid = "本项目施工过程中严格执行国家和地方现行的安全生产法律法规及标准规范确保现场安全"
    whitelist = dup.TenderWhitelist(tender)
    normalized_tender = dup.normalize_sentence(tender)
    normalized_bid = dup.normalize_sentence(bid)
    # 用例前提：两者 Dice 落在长句阈值之上
    assert len(normalized_tender) >= dup._NEAR_LONG_LENGTH
    assert dup.bigram_dice_similarity(normalized_tender, normalized_bid) >= dup._NEAR_LONG_THRESHOLD
    assert whitelist.match_level(bid) == "near"


def test_whitelist_rejects_unrelated_and_short_loose_sentences():
    whitelist = dup.TenderWhitelist("本项目施工过程中严格执行安全生产法律法规及标准规范。")
    assert whitelist.match_level("我方具备建筑工程施工总承包一级资质") is None
    # 短句改写：Dice 低于短句阈值，不豁免（避免把"两家都这么写"当成引用招标原文）
    assert whitelist.match_level("施工组织设计应包括施工验收环节") is None


def test_whitelist_thresholds_are_stricter_for_short_sentences():
    assert dup._NEAR_SHORT_THRESHOLD > dup._NEAR_LONG_THRESHOLD
    assert dup._NEAR_LONG_LENGTH == 30


def test_whitelist_len_reports_indexed_sentence_count():
    whitelist = dup.TenderWhitelist("第一句足够长用来建立索引用于相似度比对判断。第二句也一样足够长的呀。")
    assert len(whitelist) == 2


def test_analyze_content_excludes_tender_quotes_and_reports_stats():
    """回归：招标原文引用不应被判为两家之间的正文重复。"""
    sentence = "投标人应当按照招标文件要求提交完整的资格证明文件并保证真实有效"
    file_infos = [
        {"file_id": 1, "filename": "A.docx", "text": f"{sentence}。"},
        {"file_id": 2, "filename": "B.docx", "text": f"{sentence}。"},
    ]

    result = dup.analyze_content(file_infos)
    assert len(result["duplicate_groups"]) == 1, "无招标原文时应正常判定重复（回归保护）"

    result_with_tender = dup.analyze_content(file_infos, tender_text=f"{sentence}。")
    assert result_with_tender["duplicate_groups"] == []
    assert result_with_tender["stats"]["tender_whitelist"]["exact"] >= 1


def test_analyze_content_still_reports_duplicates_when_tender_unrelated():
    files = [
        {"file_id": 1, "filename": "A.docx", "text": "我方承诺投入两台大型履带式起重机完成吊装作业任务。"},
        {"file_id": 2, "filename": "B.docx", "text": "我方承诺投入两台大型履带式起重机完成吊装作业任务。"},
    ]
    result = dup.analyze_content(files, tender_text="招标范围包括土建工程与安装工程两部分内容。")
    assert len(result["duplicate_groups"]) == 1
    assert result["stats"]["tender_whitelist_count"] == 1


# ==================== D. 图片重复上下文定位 ====================


_DOCUMENT_XML = """<?xml version="1.0" encoding="UTF-8"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
            xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
  <w:body>
    <w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>第1章 施工方案</w:t></w:r></w:p>
    <w:p><w:r><w:t>现场平面布置图如下</w:t></w:r>
      <w:r><w:drawing><a:blip r:embed="rId5"/></w:drawing></w:r></w:p>
    <w:p><w:pPr><w:pStyle w:val="Heading2"/></w:pPr><w:r><w:t>1.1 施工进度计划</w:t></w:r></w:p>
    <w:p><w:r><w:drawing><a:blip r:embed="rId6"/></w:drawing></w:r></w:p>
  </w:body>
</w:document>
"""

_RELS_XML = """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId5" Type="image" Target="media/image1.png"/>
  <Relationship Id="rId6" Type="image" Target="media/sub/image2.png"/>
</Relationships>
"""


def _make_docx(tmp_path: Path, name: str, image1: bytes, image2: bytes = b"img2") -> Path:
    """构造最小可用 docx（只含解析器需要的部件）。"""
    path = tmp_path / name
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("word/document.xml", _DOCUMENT_XML)
        z.writestr("word/_rels/document.xml.rels", _RELS_XML)
        z.writestr("word/media/image1.png", image1)
        z.writestr("word/media/sub/image2.png", image2)
    return path


def test_docx_image_contexts_extract_heading_path_and_previous_text(tmp_path):
    path = _make_docx(tmp_path, "a.docx", b"same-image-bytes")
    with zipfile.ZipFile(path) as z:
        contexts = dup._docx_image_contexts(z.read("word/document.xml"))

    assert [c["rid"] for c in contexts] == ["rId5", "rId6"]
    assert contexts[0]["image_index"] == 1
    assert contexts[0]["directory"] == "第1章 施工方案"
    assert contexts[0]["previous_sentence"] == "现场平面布置图如下"
    # 第二张图处于 1.1 小节内，路径应包含两级标题
    assert contexts[1]["image_index"] == 2
    assert contexts[1]["directory"] == "第1章 施工方案 / 1.1 施工进度计划"


def test_docx_relationship_targets_normalize_paths(tmp_path):
    path = _make_docx(tmp_path, "a.docx", b"x")
    with zipfile.ZipFile(path) as z:
        targets = dup._docx_relationship_targets(z.read("word/_rels/document.xml.rels"))
    assert targets["rId5"] == "word/media/image1.png"
    assert targets["rId6"] == "word/media/sub/image2.png"


def test_extract_docx_images_carries_context(tmp_path):
    path = _make_docx(tmp_path, "a.docx", b"same-image-bytes")
    images = dup._extract_docx_images(str(path))
    assert len(images) == 2
    first = next(img for img in images if img["path"] == "word/media/image1.png")
    assert first["image_index"] == 1
    assert first["directory"] == "第1章 施工方案"
    assert first["previous_sentence"] == "现场平面布置图如下"


def test_analyze_images_reports_locations_and_counts(tmp_path):
    first = _make_docx(tmp_path, "A.docx", b"shared-picture", b"only-a")
    second = _make_docx(tmp_path, "B.docx", b"shared-picture", b"only-b")
    file_infos = [
        {"file_id": 1, "filename": "A.docx", "path": str(first), "file_type": "docx"},
        {"file_id": 2, "filename": "B.docx", "path": str(second), "file_type": "docx"},
    ]

    result = dup.analyze_images(file_infos)
    assert result["total_images"] == 4
    assert len(result["duplicate_groups"]) == 1
    group = result["duplicate_groups"][0]
    assert len(group["occurrences"]) == 2
    assert result["duplicate_image_count"] == 2
    assert all("第1章 施工方案" in location for location in group["locations"])
    assert group["occurrences"][0]["previous_sentence"] == "现场平面布置图如下"


def test_analyze_images_ignores_single_file_matches(tmp_path):
    only = _make_docx(tmp_path, "A.docx", b"unique-a", b"unique-b")
    result = dup.analyze_images(
        [{"file_id": 1, "filename": "A.docx", "path": str(only), "file_type": "docx"}],
    )
    assert result["duplicate_groups"] == []
    assert result["duplicate_image_count"] == 0


# ==================== E. 报告下载守卫 ====================


def test_report_download_roots_cover_reports_and_project_exports():
    roots = report_download_roots()
    assert PROJECTS_DATA_DIR.resolve() in roots
    assert REJECTION_REPORTS_DIR.resolve() in roots
    assert DUPLICATE_REPORTS_DIR.resolve() in roots


def test_is_downloadable_path_accepts_reports_and_rejects_outside_paths():
    assert is_downloadable_path(REJECTION_REPORTS_DIR / "rejection_report_1.docx")
    assert is_downloadable_path(DUPLICATE_REPORTS_DIR / "duplicate_report_1.xlsx")
    assert is_downloadable_path(PROJECTS_DATA_DIR / "1" / "_exports" / "a.docx")
    assert not is_downloadable_path(Path("C:/Windows/System32/drivers/etc/hosts"))
    assert not is_downloadable_path(Path("/etc/passwd"))


def test_report_dirs_are_outside_projects_dir_which_caused_the_original_403():
    """回归说明：这个前提一旦变化，说明目录约定改了，需要同步调整守卫。"""
    assert not REJECTION_REPORTS_DIR.resolve().is_relative_to(PROJECTS_DATA_DIR.resolve())
    assert not DUPLICATE_REPORTS_DIR.resolve().is_relative_to(PROJECTS_DATA_DIR.resolve())


# ==================== F. 路由端到端 ====================


def test_rejection_export_route_serves_xlsx_and_docx():
    from app.routers import rejection as rejection_router
    from app.services import rejection as rejection_service

    task_id = rejection_service.create_task(
        project_id=None,
        tender_text="招标文件",
        bid_documents=[{"id": 1, "name": "技术标.docx", "text": "工期为100日历天。"}],
        check_items=[{"id": "r001", "description": "工期要求"}],
        enabled_dimensions=["rejection"],
    )
    rejection_service.update_task_status(
        task_id,
        "completed",
        results={
            "rejection": [
                {"id": "i1", "description": "工期超期", "severity": "critical", "quote": "工期为100日历天。"},
            ],
            "spelling": [],
            "logic": [],
        },
    )

    xlsx_response = asyncio.run(rejection_router.export_report(task_id=task_id, fmt="xlsx"))
    assert Path(xlsx_response.path).suffix == ".xlsx"
    assert xlsx_response.media_type.endswith("spreadsheetml.sheet")

    docx_response = asyncio.run(rejection_router.export_report(task_id=task_id, fmt="docx"))
    assert Path(docx_response.path).suffix == ".docx"
    assert docx_response.media_type.endswith("wordprocessingml.document")


def test_rejection_export_route_rejects_unknown_format():
    from fastapi import HTTPException

    from app.routers import rejection as rejection_router

    with pytest.raises(HTTPException) as excinfo:
        asyncio.run(rejection_router.export_report(task_id=1, fmt="pdf"))
    assert excinfo.value.status_code == 400


def test_rejection_export_route_404_for_missing_task():
    from fastapi import HTTPException

    from app.routers import rejection as rejection_router

    with pytest.raises(HTTPException) as excinfo:
        asyncio.run(rejection_router.export_report(task_id=99999999, fmt="xlsx"))
    assert excinfo.value.status_code == 404


def test_duplicate_export_route_serves_xlsx_and_docx():
    from app.routers import duplicate as duplicate_router
    from app.services import duplicate as duplicate_service

    task_id, _cached = duplicate_service.create_task([8001, 8002], None)
    duplicate_service.update_task_status(
        task_id,
        "completed",
        results={
            "metadata": {"metadata": [], "anomalies": [], "fingerprints": [], "same_source_pairs": []},
            "toc": {
                "pairwise": [
                    {
                        "file_a": "A.docx",
                        "file_b": "B.docx",
                        "score": 0.8,
                        "title_overlap": 0.8,
                        "path_overlap": 0.8,
                        "order_similarity": 0.8,
                        "shared_count": 5,
                        "risk": "high",
                    },
                ],
                "duplicate_groups": [],
                "similar_groups": [],
            },
            "content": {"stats": {}, "duplicate_groups": [], "similar_groups": []},
            "images": {"total_images": 0, "duplicate_groups": []},
        },
    )

    xlsx_response = asyncio.run(duplicate_router.export_report(task_id=task_id, fmt="xlsx"))
    assert Path(xlsx_response.path).suffix == ".xlsx"
    docx_response = asyncio.run(duplicate_router.export_report(task_id=task_id, fmt="docx"))
    assert Path(docx_response.path).suffix == ".docx"


def test_duplicate_export_route_rejects_unknown_format():
    from fastapi import HTTPException

    from app.routers import duplicate as duplicate_router

    with pytest.raises(HTTPException) as excinfo:
        asyncio.run(duplicate_router.export_report(task_id=1, fmt="csv"))
    assert excinfo.value.status_code == 400


def test_export_routes_use_shared_download_guard():
    """源码级接线检查：两个路由都不能再写死 PROJECTS_DATA_DIR 守卫。"""
    from app.routers import duplicate as duplicate_router
    from app.routers import rejection as rejection_router

    for module in (rejection_router, duplicate_router):
        source = inspect.getsource(module)
        assert "is_downloadable_path" in source
        assert "PROJECTS_DATA_DIR.resolve()" not in source, "下载守卫必须统一走 config.is_downloadable_path()"
        assert "_REPORT_FORMATS" in source


def test_excel_service_exports_are_registered():
    from app.services import duplicate as duplicate_service
    from app.services import rejection as rejection_service

    assert callable(rejection_service.export_report_excel)
    assert callable(duplicate_service.export_report_excel)
