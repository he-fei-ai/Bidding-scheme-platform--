"""图表编号（导出链路）组件级交互测试

覆盖 2026-09-29 修复项：图号占号/虚跳与并发导出串号
1. `_figure_counter_state` 按项目分桶：两个项目并发导出互不串号
2. `_reset_figure_counter(project_id)` 只清本项目桶，不再抹掉并发会话的计数
   （原实现进程级全局 + build_docx 无条件 clear → A 的号被 B 重置 → 同一文档重复图号）
3. `_rewind_figure_number` 回滚：渲染失败不得占号
   （原实现先占号再渲染，失败不回退 → 正文出现 图1-1 直接跳到 图1-3 的虚跳）
4. `_add_inline_chart` 端到端：失败说明文字不再携带已被回收的编号；
   失败图后的下一张图复用该号，正文图号连续
"""
import base64
import io
import json
from pathlib import Path
import sys
from unittest.mock import MagicMock, patch

from docx import Document


BACKEND_DIR = Path(__file__).resolve().parent.parent
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from app.models import ExportConfig, Section  # noqa: E402
from app.services import export as export_mod  # noqa: E402


# 1x1 合法 PNG，用于伪造"渲染成功"返回值
_PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def _mk_detail(project_id: int, section_id: int, chart_field: str, payload: dict) -> Section:
    kwargs = {
        "id": section_id,
        "project_id": project_id,
        "title": f"章节{section_id}",
        "level": 3,
        "content": "正文",
    }
    kwargs[chart_field] = json.dumps(payload, ensure_ascii=False)
    return Section(**kwargs)


def _render_calls(stream_values):
    """构造 render_mermaid_to_bytes 替身：按调用顺序返回给定值（None=渲染失败）"""
    seq = list(stream_values)

    def _fake(payload_str, chart_type, skip_http=False):
        return seq.pop(0) if seq else None

    return _fake


class TestFigureCounterScoping:
    """图号计数器的项目分桶与重置隔离"""

    def setup_method(self):
        export_mod._reset_figure_counter()

    def teardown_method(self):
        export_mod._reset_figure_counter()

    def test_same_project_same_parent_increments(self):
        assert export_mod._next_figure_number("3.2", 7) == "图3-1"
        assert export_mod._next_figure_number("3.2", 7) == "图3-2"

    def test_same_chapter_parents_share_continuous_seq(self):
        """同章不同子节共用一条连续计数器（图号「图{章号}-{序号}」按章连续）。

        BUG-FIX 2026-10-14：旧实现分桶键用完整 parent_section，导致 "3.2"→图3-1
        后 "3.5"→又图3-1（同章重号，即需求「编号冲突」）。现统一按截断章号分桶，
        同章恒连续：3.2→图3-1、3.5→图3-2。生产主路径传纯章号，行为不变。
        """
        assert export_mod._next_figure_number("3.2", 7) == "图3-1"
        assert export_mod._next_figure_number("3.5", 7) == "图3-2"

    def test_cross_project_isolation(self):
        """并发导出串号回归：项目 7 与项目 8 的计数互不影响"""
        assert export_mod._next_figure_number("3.2", 7) == "图3-1"
        assert export_mod._next_figure_number("3.2", 7) == "图3-2"
        assert export_mod._next_figure_number("3.2", 8) == "图3-1"
        assert export_mod._next_figure_number("3.2", 7) == "图3-3"

    def test_reset_only_clears_target_project(self):
        """原缺陷：build_docx 无条件 clear 全部 → 并发会话的号被抹掉后从头重算"""
        export_mod._next_figure_number("3.2", 7)
        export_mod._next_figure_number("3.2", 8)
        export_mod._reset_figure_counter(8)
        # 项目 7 的桶保留 → 继续 2，而不是被重置回 1
        assert export_mod._next_figure_number("3.2", 7) == "图3-2"
        # 项目 8 已被重置 → 从 1 开始
        assert export_mod._next_figure_number("3.2", 8) == "图3-1"

    def test_reset_without_project_clears_all(self):
        export_mod._next_figure_number("3.2", 7)
        export_mod._reset_figure_counter()
        assert export_mod._next_figure_number("3.2", 7) == "图3-1"


class TestFigureNumberRewind:
    """回滚语义：渲染失败不得占号"""

    def setup_method(self):
        export_mod._reset_figure_counter()

    def teardown_method(self):
        export_mod._reset_figure_counter()

    def test_rewind_restores_slot(self):
        assert export_mod._next_figure_number("2.1", 5) == "图2-1"
        assert export_mod._next_figure_number("2.1", 5) == "图2-2"
        export_mod._rewind_figure_number("2.1", 5)
        assert export_mod._next_figure_number("2.1", 5) == "图2-2"

    def test_rewind_on_first_number_removes_bucket(self):
        export_mod._next_figure_number("2.1", 5)
        export_mod._rewind_figure_number("2.1", 5)
        assert "5|2" not in export_mod._figure_counter_state
        assert export_mod._next_figure_number("2.1", 5) == "图2-1"

    def test_rewind_is_project_scoped(self):
        export_mod._next_figure_number("2.1", 5)
        export_mod._next_figure_number("2.1", 6)
        export_mod._rewind_figure_number("2.1", 5)
        # 项目 6 不受项目 5 的回滚影响
        assert export_mod._next_figure_number("2.1", 6) == "图2-2"


class TestInlineChartNumberingIntegration:
    """_add_inline_chart 的取号/回滚与失败文案"""

    def setup_method(self):
        export_mod._reset_figure_counter()

    def teardown_method(self):
        export_mod._reset_figure_counter()

    def _run(self, doc, detail, stream_values, project_id=7):
        cache_stub = MagicMock()
        cache_stub.get.return_value = None
        with patch.object(export_mod, "render_mermaid_to_bytes", _render_calls(stream_values)), \
             patch.object(export_mod, "_chart_cache", cache_stub), \
             patch.object(export_mod, "_get_cached_chart_render", return_value=None), \
             patch.object(export_mod, "_store_chart_render", return_value=None):
            return [
                export_mod._add_inline_chart(
                    doc,
                    "gantt",
                    "",
                    detail,
                    ExportConfig(),
                    _http_service_available=False,
                    parent_section_number="3.1",
                )
                for _ in range(len(stream_values))
            ]

    def test_failed_render_does_not_burn_figure_number(self, monkeypatch):
        """图1 渲染失败 → 号回收；图2 必须拿到 图3-1（正文不出现 图3-1→图3-2 漏号）

        CROSS-FIX 2026-10-04：失败说明插入受图表模块新增的 CHART_FAIL_PLACEHOLDER
        开关控制（默认 false=静默跳过，历史行为等价 true），本用例断言失败
        说明内容，显式开启；图号回滚不受开关影响。
        """
        monkeypatch.setattr(export_mod, "CHART_FAIL_PLACEHOLDER", True)
        doc = Document()
        detail = _mk_detail(7, 101, "gantt_json", {"type": "gantt", "title": "施工进度"})

        # 第一次渲染失败（None），第二次成功（合法 PNG）
        payload_stream = io.BytesIO(_PNG_1X1)
        results = self._run(doc, detail, [None, payload_stream])

        assert results[0] is False, "渲染失败必须返回 False"
        assert results[1] is True, "渲染成功必须返回 True"

        texts = [p.text for p in doc.paragraphs]
        fail_lines = [t for t in texts if "【图表渲染失败】" in t]
        assert fail_lines, texts
        # 失败说明里不得携带编号：号已回收给下一张图，留着会造成同号双图
        assert "图3-1" not in fail_lines[0], fail_lines

        # 成功那张图拿到 图3-1（未被失败图占掉）
        caption_lines = [t for t in texts if t.startswith("图3-")]
        assert caption_lines and caption_lines[0].startswith("图3-1"), caption_lines
        # 计数器最终只推进 1 格（失败图未占号）；分桶键按截断章号归一（图3-x→"3"）
        assert export_mod._figure_counter_state.get("7|3") == 1, (
            export_mod._figure_counter_state
        )

    def test_preload_only_never_consumes_number(self):
        """预渲染阶段（preload_only）不得消耗会话级计数器"""
        doc = Document()
        detail = _mk_detail(7, 102, "gantt_json", {"type": "gantt", "title": "预热"})
        cache_stub = MagicMock()
        cache_stub.get.return_value = None
        with patch.object(export_mod, "render_mermaid_to_bytes", _render_calls([None])), \
             patch.object(export_mod, "_chart_cache", cache_stub), \
             patch.object(export_mod, "_get_cached_chart_render", return_value=None):
            export_mod._add_inline_chart(
                doc,
                "gantt",
                "",
                detail,
                ExportConfig(),
                preload_only=True,
                _http_service_available=False,
                parent_section_number="3.1",
            )
        assert export_mod._figure_counter_state == {}
        # 预渲染不向文档写入任何内容（含失败说明）
        assert not [p for p in doc.paragraphs if "【图表渲染失败】" in p.text]
