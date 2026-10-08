"""【一致性校核模块】+【导出文档模块】—— 2026-10-09 深度审计修复回归测试。

原则：每条用例对应一个**实测确认**的缺陷；只断言可观察行为（HTTP 响应体 /
DB 落库 / 导出产物 / 渲染出的段落文本），不测内部实现细节；不依赖真实 AI
与网络。

覆盖缺陷：

一致性校核
  | C-1 | 金额单位被正则丢弃 → 「1200万元」vs「12000000元」被判冲突（误报），
        且采纳建议会把 12000000 改成 1200（**修复建议有害**）        |
  | C-2 | `SEVERITY_LOW`（轻微）是**永不产出的死档** —— 命名空间定义了一档，
        `_assess_severity` 只有 high/medium 两条分支                  |
  | C-3 | `POST autofix` 的 `new_value` 后端**零校验**：空值会把正文里的
        事实值替换成空串（「工期 30 天」→「工期 天」）                |
  | C-4 | 同一项目的「审计 / 自动修复 / AI 修复」并发无互斥，与
        `_save_conflicts` 的 DELETE+INSERT 交错 → 冲突状态错乱         |

导出文档
  | E-1 | 图号章号取**节点自身层级序号**（每章从 1 重数）→ 跨章重号      |
  | E-2 | L5+ 标题被硬拼一个顿号（"、施工组织设计"），且该节图表**无图号**|
  | E-3 | 带 info-string 的闭合行（```` ```mermaid ````）被判未闭合
        → 整块图表被当坏块静默删除                                     |
  | E-4 | 原子写盘的临时文件名固定 → 同进程并发导出互踩；只验大小不验
        完整性 → 截断的 docx 被提升为正式产物                          |
  | E-5 | 缓存键只在**顶层**剔除 `fig_no` → 嵌套层的图号仍击穿缓存       |
"""
from __future__ import annotations

import asyncio
import json
import os
import threading

import pytest


# ============================================================================
# 公共夹具
# ============================================================================

OWNER = 4100


def _new_project(name: str) -> int:
    from app.db import get_db, init_db

    init_db()
    with get_db() as conn:
        conn.execute(
            "INSERT INTO projects (name, industry, created_by) VALUES (?, ?, ?)",
            (name, "通用", OWNER),
        )
        return conn.execute("SELECT id FROM projects ORDER BY id DESC LIMIT 1").fetchone()["id"]


def _seed_section(project_id: int, *, title: str, content: str, sort_order: int = 0) -> int:
    from app.db import get_db

    with get_db() as conn:
        cur = conn.execute(
            "INSERT INTO sections (project_id, title, level, sort_order, content, status, section_type) "
            "VALUES (?, ?, 1, ?, ?, 'completed', 'normal')",
            (project_id, title, sort_order, content),
        )
        return cur.lastrowid


def _seed_conflict(
    project_id: int,
    *,
    conflict_id: str,
    fact_name: str = "工期",
    values: str = '["450","500"]',
    sources: str = '["section_1 (第一章)"]',
    resolved: int = 0,
) -> None:
    from app.db import get_db

    with get_db() as conn:
        conn.execute(
            "INSERT INTO consistency_conflicts "
            "(project_id, conflict_id, fact_name, values_json, sources_json, "
            " severity, description, resolved, created_at) "
            "VALUES (?, ?, ?, ?, ?, 'medium', '手工构造', ?, '2026-10-09 10:00:00')",
            (project_id, conflict_id, fact_name, values, sources, resolved),
        )


def _client(user_id: int = OWNER):
    from fastapi.testclient import TestClient

    from app.main import app
    from app.security.auth import User, get_current_user

    app.dependency_overrides[get_current_user] = lambda: User(user_id, "u", "user")
    return TestClient(app)


@pytest.fixture(autouse=True)
def _clean_overrides():
    yield
    try:
        from app.main import app

        app.dependency_overrides.clear()
    except Exception:
        pass


# ============================================================================
# C-1 —— 金额单位归一（消除误报 + 有害修复建议）
# ============================================================================


class TestAmountUnitNormalization:
    def test_amount_unit_is_converted_to_yuan(self):
        """1200万元 与 12000000元 是同一金额，必须归一到同一个比较键。"""
        from app.services import consistency_rules as rules

        assert rules.normalize_amount_to_yuan("1200万元") == "12000000"
        assert rules.normalize_amount_to_yuan("12000000元") == "12000000"
        assert rules.normalize_amount_to_yuan("12000000") == "12000000"
        assert rules.normalize_amount_to_yuan("1,200万元") == "12000000"
        assert rules.normalize_amount_to_yuan("100万") == "1000000"
        assert rules.normalize_amount_to_yuan("1亿元") == "100000000"
        # 不可换算 → None（调用方回退通用归一化，宁多报也不漏报）
        assert rules.normalize_amount_to_yuan("待定") is None
        assert rules.normalize_amount_to_yuan("") is None

    def test_amount_extraction_keeps_unit(self):
        """抽取必须把单位一起带上，否则比较与替换都失去语义。"""
        from app.services.consistency_auditor import ConsistencyAuditor

        facts = ConsistencyAuditor().extract_facts_from_text("合同价：1200万元")
        amounts = [f for f in facts if f.name == "金额"]
        assert amounts, [f.name for f in facts]
        assert amounts[0].value == "1200万元", amounts[0].value

    def test_same_amount_different_unit_is_not_a_conflict(self):
        """BUG C-1（P1·误报）回归：同一金额的两种写法不得判成冲突。"""
        from app.services.consistency_auditor import ConsistencyAuditor

        auditor = ConsistencyAuditor()
        # 旧实现：抽成 ["1200", "12000000"] → 误报一条严重冲突
        values = ["1200万元", "12000000元"]
        assert len(auditor._distinct_values(values, fact_name="金额")) == 1
        # 真冲突仍然是冲突
        assert len(auditor._distinct_values(["1200万元", "1500万元"], fact_name="金额")) == 2

    def test_amount_conflict_suggestion_carries_unit(self):
        """修复建议的 target_value 必须带单位，否则写回正文会改掉量级。"""
        pid = _new_project("C1-金额建议")
        _seed_section(pid, title="第一章", content="合同价：1200万元")
        _seed_section(pid, title="第二章", content="造价：1500万元", sort_order=1)
        _seed_conflict(
            pid,
            conflict_id=f"{pid}_amount",
            fact_name="金额",
            values=json.dumps(["1200万元", "1500万元"], ensure_ascii=False),
            sources=json.dumps(["section_1 (第一章)", "section_2 (第二章)"], ensure_ascii=False),
        )

        with _client() as client:
            res = client.get(
                f"/api/v1/consistency/conflicts/{pid}_amount/fix-suggestion",
                params={"project_id": pid},
            )
        assert res.status_code == 200, res.text
        body = res.json()
        assert body["target_value"] and "万元" in body["target_value"], body


# ============================================================================
# C-2 —— `low`（轻微）档必须可达
# ============================================================================


class TestSeverityLowIsReachable:
    def test_rules_assess_severity_tiers(self):
        from app.services import consistency_rules as rules

        # 关键事实 → high
        assert rules.assess_severity("工期", ["30", "45"]) == rules.SEVERITY_HIGH
        assert rules.assess_severity("金额", ["1", "2"]) == rules.SEVERITY_HIGH
        # 非关键事实 + ≥3 种取值 → high
        assert rules.assess_severity("质量标准", ["合格", "优良", "优秀"]) == rules.SEVERITY_HIGH
        # 非关键事实 + 2 种取值 + 正文内部互相矛盾（≥2 个章节）→ medium
        assert (
            rules.assess_severity("质量标准", ["合格", "优良"], section_count=2)
            == rules.SEVERITY_MEDIUM
        )
        # 非关键事实 + 2 种取值 + 只落在单个章节（正文自洽，仅与基准有偏差）→ low
        assert (
            rules.assess_severity("质量标准", ["合格", "优良"], section_count=1)
            == rules.SEVERITY_LOW
        )
        # section_count 未知时不判 low（保守回退 medium）
        assert rules.assess_severity("质量标准", ["合格", "优良"]) == rules.SEVERITY_MEDIUM

    def test_auditor_delegates_to_rules_and_can_emit_low(self):
        from app.services import consistency_rules as rules
        from app.services.consistency_auditor import ConsistencyAuditor

        auditor = ConsistencyAuditor()
        assert (
            auditor._assess_severity("质量标准", ["合格", "优良"], section_count=1)
            == rules.SEVERITY_LOW
        )
        assert (
            auditor._assess_severity("质量标准", ["合格", "优良"], section_count=3)
            == rules.SEVERITY_MEDIUM
        )
        # 既有契约不变（未提供 section_count 时行为与修复前一致）
        assert auditor._assess_severity("工期", ["30", "45"]) == rules.SEVERITY_HIGH
        assert auditor._assess_severity("质量标准", ["合格", "优良"]) == rules.SEVERITY_MEDIUM

    def test_rules_payload_exposes_low_rule(self):
        from app.services import consistency_rules as rules

        payload = rules.rules_payload()
        assert payload["severity"]["order"] == ["high", "medium", "low"]
        assert "low" in payload["severity"]["meta"]
        assert payload["facts"]["low_max_sections"] == rules.SEVERITY_LOW_MAX_SECTIONS


# ============================================================================
# C-3 —— autofix 的 new_value 必须校验
# ============================================================================


class TestAutoFixValueValidation:
    def _setup(self) -> int:
        pid = _new_project("C3-autofix校验")
        sid = _seed_section(pid, title="第一章", content="工期 30 天")
        _seed_conflict(
            pid,
            conflict_id=f"{pid}_c3",
            fact_name="工期",
            values='["30","45"]',
            sources=json.dumps([f"section_{sid} (第一章)"], ensure_ascii=False),
        )
        return pid

    def test_empty_new_value_is_rejected(self):
        """BUG C-3（P1）回归：空值会把正文里的「30」替换成空串并标记已解决。"""
        pid = self._setup()
        with _client() as client:
            res = client.post(
                f"/api/v1/consistency/conflicts/{pid}_c3/autofix",
                params={"project_id": pid, "new_value": "   "},
            )
        assert res.status_code == 400, res.text

        # 正文必须原封不动
        from app.db import get_db

        with get_db(readonly=True) as conn:
            content = conn.execute(
                "SELECT content FROM sections WHERE project_id = ?", (pid,)
            ).fetchone()["content"]
        assert content == "工期 30 天", content

    def test_overlong_new_value_is_rejected(self):
        pid = self._setup()
        with _client() as client:
            res = client.post(
                f"/api/v1/consistency/conflicts/{pid}_c3/autofix",
                params={"project_id": pid, "new_value": "X" * 500},
            )
        assert res.status_code == 400, res.text

    def test_valid_new_value_still_works(self):
        pid = self._setup()
        with _client() as client:
            res = client.post(
                f"/api/v1/consistency/conflicts/{pid}_c3/autofix",
                params={"project_id": pid, "new_value": "45"},
            )
        assert res.status_code == 200, res.text
        assert res.json()["status"] == "ok", res.json()

        from app.db import get_db

        with get_db(readonly=True) as conn:
            content = conn.execute(
                "SELECT content FROM sections WHERE project_id = ?", (pid,)
            ).fetchone()["content"]
        assert content == "工期 45 天", content


# ============================================================================
# C-4 —— 同项目写路径并发互斥
# ============================================================================


class TestConsistencyWritePathSerialization:
    def test_concurrent_audits_do_not_overlap(self, monkeypatch):
        """BUG C-4（P1·并发竞态）回归。

        `_save_conflicts` 是「DELETE 未解决冲突 + INSERT 本次结果」；两个审计
        并发时，后者的 DELETE 会清掉前者刚写入的记录（或反之），冲突状态错乱。
        现在三条写路径共用同一把项目锁，必须严格串行。
        """
        from app.security.auth import User
        from app.services.consistency_auditor import consistency_auditor as auditor
        from app.routers import consistency as consistency_router

        pid = _new_project("C4-并发")
        timeline: list[str] = []
        tl_lock = threading.Lock()

        def _slow_audit(project_id):
            from app.services.consistency_auditor import AuditResult

            with tl_lock:
                timeline.append("enter")
            import time

            time.sleep(0.05)
            with tl_lock:
                timeline.append("exit")
            return AuditResult(
                project_id=project_id,
                total_sections=0,
                facts_extracted=0,
                conflicts_found=0,
                status="completed",
            )

        monkeypatch.setattr(auditor, "run_full_audit", _slow_audit)

        async def _two_audits():
            await asyncio.gather(
                consistency_router.run_audit(project_id=pid, user=User(OWNER, "u", "user")),
                consistency_router.run_audit(project_id=pid, user=User(OWNER, "u", "user")),
            )

        asyncio.run(_two_audits())

        # 严格串行：enter → exit → enter → exit
        assert timeline == ["enter", "exit", "enter", "exit"], timeline

    def test_different_projects_are_not_blocked(self, monkeypatch):
        """锁按项目分桶：不同项目不得互相阻塞（否则多项目并发导出会被拖死）。"""
        from app.security.auth import User
        from app.services.consistency_auditor import consistency_auditor as auditor
        from app.routers import consistency as consistency_router

        pid_a = _new_project("C4-A")
        pid_b = _new_project("C4-B")
        events: list[str] = []
        ev_lock = threading.Lock()
        barrier = threading.Barrier(2, timeout=5)

        def _slow_audit(project_id):
            from app.services.consistency_auditor import AuditResult

            with ev_lock:
                events.append(f"enter:{project_id}")
            try:
                barrier.wait()  # 两个都进来了才继续 → 证明未互斥
            except Exception:
                pass
            with ev_lock:
                events.append(f"exit:{project_id}")
            return AuditResult(
                project_id=project_id,
                total_sections=0,
                facts_extracted=0,
                conflicts_found=0,
                status="completed",
            )

        monkeypatch.setattr(auditor, "run_full_audit", _slow_audit)

        async def _two_audits():
            await asyncio.gather(
                consistency_router.run_audit(project_id=pid_a, user=User(OWNER, "u", "user")),
                consistency_router.run_audit(project_id=pid_b, user=User(OWNER, "u", "user")),
            )

        asyncio.run(_two_audits())
        assert len([e for e in events if e.startswith("enter:")]) == 2, events


# ============================================================================
# E-1 / E-2 —— 图号章号与 L5+ 标题标点
# ============================================================================


class _FakeNode:
    def __init__(self, node_id: int, level: int, title: str, children=None):
        self.id = node_id
        self.level = level
        self.title = title
        self.children = children or []


class _FakeDetail:
    """最小章节详情替身：只提供 `_write_tree_content_with_progress` 需要的字段。"""

    def __init__(self, node_id: int, title: str):
        self.id = node_id
        self.title = title
        self.content = "正文占位"
        self.project_id = 1


class TestFigureChapterNumber:
    def test_inline_chart_prefers_chapter_number(self):
        """BUG E-1（P1·跨章重号）回归：显式章号优先于节点自身序号。"""
        from app.services import export as export_mod

        export_mod._reset_figure_counter()
        try:
            # 第 2 章的 L2 节点：自身编号是「一」→ 旧口径算出 图1-x（错）
            assert export_mod._next_figure_number("1", 9001) == "图1-1"
            assert export_mod._next_figure_number("1", 9001) == "图1-2"
            # 章号 2 应独立计数
            assert export_mod._next_figure_number("2", 9001) == "图2-1"
        finally:
            export_mod._reset_figure_counter()

    def test_tree_writer_passes_chapter_number_and_no_stray_punctuation(self, monkeypatch):
        """BUG E-1/E-2 回归：`_write_tree_content_with_progress` 的透传与标题标点。

        - 章号必须取当前 L1 序号（第 2 章下的节点也得到 "2"）；
        - L5+ 节点（heading_standard 不编号）不得被拼上孤立的「、」。
        """
        from docx import Document

        from app.models import ExportConfig
        from app.services import export as export_mod

        captured: list[dict] = []

        def _fake_add_parsed_content(doc, content, config, **kwargs):
            captured.append(kwargs)
            return None

        monkeypatch.setattr(export_mod, "_add_parsed_content", _fake_add_parsed_content)

        tree = [
            _FakeNode(1, 1, "第一章标题", [_FakeNode(2, 3, "子节点A")]),
            _FakeNode(3, 1, "第二章标题", [_FakeNode(4, 5, "五级节点B")]),
        ]
        doc = Document()
        detail_index = {
            1: _FakeDetail(1, "第一章标题"),
            2: _FakeDetail(2, "子节点A"),
            3: _FakeDetail(3, "第二章标题"),
            4: _FakeDetail(4, "五级节点B"),
        }
        export_mod._write_tree_content_with_progress(
            doc, tree, ExportConfig(), project_id=1, detail_index=detail_index
        )

        chapters = [c.get("chapter_number") for c in captured]
        assert chapters == ["1", "1", "2", "2"], chapters

        texts = [p.text for p in doc.paragraphs]
        assert len(texts) == 4, texts
        # 标题文本会被 `_strip_heading_numbering` 二次处理，故只断言编号前缀
        assert texts[0].startswith("第一节、"), texts
        assert texts[1].startswith("1、") and texts[1].endswith("子节点A"), texts
        assert texts[2].startswith("第二节、"), texts
        # L5+ 无编号 → **不得**出现孤立的顿号（旧实现渲染成「、五级节点B」）
        assert texts[3] == "五级节点B", texts

    def test_failed_section_retry_does_not_duplicate_content(self, monkeypatch):
        """BUG（P1·重试重复出图/重复正文）回归。

        旧实现在章节渲染抛异常时**保留已插入的元素**（可能已有若干张图）再追加
        占位符；随后重试对同一节点整段重跑一遍 → 文档里出现重复的图片与段落，
        且两轮各占一批图号（图号断档）。现在失败分支回滚本节点已写入的内容。
        """
        from docx import Document

        from app.models import ExportConfig
        from app.services import export as export_mod

        calls = {"n": 0}

        def _flaky(doc, content, config, **kwargs):
            calls["n"] += 1
            if calls["n"] == 1:
                # 模拟「插入了一部分（含图表）后才失败」
                doc.add_paragraph("内容片段")
                raise RuntimeError("模拟首次渲染中途失败")
            doc.add_paragraph("内容片段")
            return None

        monkeypatch.setattr(export_mod, "_add_parsed_content", _flaky)

        # 用真实章节：重试路径走 `_get_section_detail`，必须能取到正文才会重跑
        # `_add_parsed_content`（否则重试空转，本用例就失去意义）。
        pid = _new_project("E-重试重复")
        sid = _seed_section(pid, title="章节一", content="正文占位")

        doc = Document()
        tree = [_FakeNode(sid, 1, "章节一")]
        export_mod._write_tree_content_with_progress(
            doc, tree, ExportConfig(), project_id=pid
        )

        texts = [p.text for p in doc.paragraphs]
        # 首次的半截内容必须被回滚，重试只保留一份
        assert texts.count("内容片段") == 1, texts
        # 重试成功后失败占位说明不应残留
        assert not [t for t in texts if "渲染异常" in t], texts

    def test_insert_failure_does_not_burn_figure_number(self, monkeypatch):
        """BUG（P1·插入失败仍占号）回归。

        `_render_commercial` 里 `_emit` 旧实现无返回值、调用方一律 `return True`：
        渲染成功但 `_add_diagram_picture` 插入失败时也被算成功 → 文档里只剩一条
        灰色占位说明，却已消耗一个图号，后续图号整体前跳（虚跳）。
        """
        import io

        from docx import Document
        from unittest.mock import MagicMock, patch

        from app.models import ExportConfig
        from app.services import export as export_mod

        export_mod._reset_figure_counter()
        try:
            monkeypatch.setattr(export_mod, "_add_diagram_picture", lambda *a, **k: False)
            monkeypatch.setattr(export_mod, "_add_image_placeholder", lambda *a, **k: None)

            doc = Document()

            class _Detail:
                id = 555
                title = "插入失败"
                project_id = 7100
                gantt_json = '{"type": "gantt", "title": "T"}'

            cache_stub = MagicMock()
            cache_stub.get.return_value = None
            with patch.object(
                export_mod, "render_mermaid_to_bytes", lambda *a, **k: io.BytesIO(b"\x89PNG")
            ), patch.object(export_mod, "_chart_cache", cache_stub), patch.object(
                export_mod, "_get_cached_chart_render", return_value=None
            ), patch.object(export_mod, "_store_chart_render", return_value=None):
                ok = export_mod._add_inline_chart(
                    doc,
                    "gantt",
                    "",
                    _Detail(),
                    ExportConfig(),
                    _http_service_available=False,
                    parent_section_number="4",
                )

            assert ok is False, "插入失败必须返回 False（否则调用方不会回滚图号）"
            # 图号必须被回滚：计数桶不应残留
            assert export_mod._figure_counter_state.get("7100|4") is None, (
                export_mod._figure_counter_state
            )
        finally:
            export_mod._reset_figure_counter()

    def test_l5_charts_still_get_figure_number(self):
        """BUG E-2 连带：L5+ 章节的图表此前因章节号缺失而完全拿不到图号。"""
        from app.services import export as export_mod

        # 旧链路：`_section_number_to_arabic("、")` → "" → `_next_figure_number("")` → ""
        assert export_mod._next_figure_number("", 9002) == ""
        # 新链路：章号由 L1 计数器直接给出，与节点层级无关
        export_mod._reset_figure_counter()
        try:
            assert export_mod._next_figure_number("3", 9002) == "图3-1"
        finally:
            export_mod._reset_figure_counter()


# ============================================================================
# E-3 —— 带 info-string 的围栏闭合行
# ============================================================================


class TestFenceCloseWithInfoString:
    def test_close_line_with_language_is_accepted(self):
        """BUG E-3（P1·图表静默丢失）回归：```` ```mermaid ```` 闭合必须被识别。"""
        from app.services.content_cleaning import (
            _protect_chart_fences,
            _restore_chart_fences,
        )

        text = "引言\n```mermaid\ngraph TD\nA-->B\n```mermaid\n结语"
        protected = _protect_chart_fences(text)
        # 旧行为：整块被判未闭合 → 删除，正文里连源码都不剩
        assert "graph TD" not in protected
        assert "\x00CHARTFENCE" in protected, protected  # 被保护为占位符而非删除
        restored = _restore_chart_fences(protected)
        assert "graph TD" in restored, restored

    def test_iter_chart_fences_sees_the_block(self):
        from app.services.content_cleaning import iter_chart_fences

        text = "```mermaid\ngraph TD\nA-->B\n```mermaid\n"
        blocks = list(iter_chart_fences(text))
        assert len(blocks) == 1, blocks
        assert blocks[0][0] == "mermaid"
        assert "graph TD" in blocks[0][1]

    def test_plain_close_line_still_works(self):
        from app.services.content_cleaning import _protect_chart_fences

        text = "```mermaid\ngraph TD\n```\n"
        protected = _protect_chart_fences(text)
        assert "\x00CHARTFENCE" in protected, protected

    def test_truly_unclosed_block_is_still_dropped(self):
        """真正未闭合（到文末都没有闭合行）的块仍按坏块删除，不得回退。"""
        from app.services.content_cleaning import _protect_chart_fences

        text = "正文\n```mermaid\ngraph TD\nA-->B\n"
        protected = _protect_chart_fences(text)
        assert "graph TD" not in protected, protected
        assert "\x00CHARTFENCE" not in protected, protected


# ============================================================================
# E-4 —— 原子写盘：临时名唯一 + 完整性校验
# ============================================================================


class TestAtomicSaveHardening:
    def test_valid_docx_is_replaced_and_no_tmp_left(self, tmp_path):
        from docx import Document

        from app.services.export import _atomic_save_docx

        out = tmp_path / "a.docx"
        doc = Document()
        doc.add_paragraph("hello")
        _atomic_save_docx(doc, str(out))
        assert out.exists()
        assert not any(p.name.startswith("a.docx.tmp.") for p in tmp_path.iterdir())

    def test_truncated_docx_is_rejected(self, tmp_path):
        """BUG E-4（P2）回归：非零但被截断的半成品不得被提升为正式产物。"""
        from docx import Document

        from app.services.export import _atomic_save_docx

        good = tmp_path / "good.docx"
        doc = Document()
        doc.add_paragraph("x" * 200)
        doc.save(str(good))

        out = tmp_path / "b.docx"

        class _TruncatingDoc:
            def save(self, path):
                data = good.read_bytes()
                with open(path, "wb") as fh:
                    fh.write(data[: len(data) // 2])  # 非零但被截断

        with pytest.raises(RuntimeError, match="截断|ZIP|损坏"):
            _atomic_save_docx(_TruncatingDoc(), str(out))
        assert not out.exists(), "截断件不得进入最终路径"
        assert not any(p.name.startswith("b.docx.tmp.") for p in tmp_path.iterdir())

    def test_tmp_names_are_unique_per_call(self, tmp_path, monkeypatch):
        """BUG E-4（P2）回归：临时文件名必须每次不同，不能固定为 `<out>.tmp.<pid>`。

        旧实现把 pid 当唯一标识 —— 同一进程内写同一目标路径的两个导出会共用
        同一个临时名：后写完的 `os.replace` 顶掉前一个产物，失败清理时还会删掉
        对方正在写的临时件。

        注：不直接起两个线程并发 replace 同一目标路径 —— Windows 上第二个
        `os.replace` 会以 PermissionError 失败（目标被占用），那是**操作系统层**
        的限制（生产由导出组锁在更上层阻止并发写同一路径），与本用例要钉的
        「临时名唯一性」无关。改为拦截 `os.replace` 记录临时名，确定性断言。
        """
        from docx import Document

        from app.services import export as export_mod

        seen: list[str] = []
        real_replace = os.replace

        def _spy_replace(src, dst):
            seen.append(os.path.basename(src))
            return real_replace(src, dst)

        monkeypatch.setattr(export_mod.os, "replace", _spy_replace)

        for idx in range(2):
            out = tmp_path / f"c{idx}.docx"
            doc = Document()
            doc.add_paragraph("x")
            export_mod._atomic_save_docx(doc, str(out))
            assert out.exists()

        assert len(seen) == 2, seen
        assert seen[0] != seen[1], f"临时名重复（同进程会互踩）: {seen}"
        assert all(".tmp." in name for name in seen), seen


# ============================================================================
# E-5 —— 缓存键递归剔除 fig_no
# ============================================================================


class TestCacheKeyFigNoStripping:
    def test_nested_fig_no_is_stripped(self):
        """BUG E-5（P2）回归：嵌套层的 fig_no 也必须在缓存键里剔除。"""
        from app.services.export import _strip_fig_no_for_cache

        a = {"type": "gantt", "fig_no": "图1-1", "data": {"fig_no": "图1-1", "x": 1}}
        b = {"type": "gantt", "fig_no": "图2-5", "data": {"fig_no": "图2-5", "x": 1}}
        assert _strip_fig_no_for_cache(a) == _strip_fig_no_for_cache(b)

    def test_list_nested_fig_no_is_stripped(self):
        from app.services.export import _strip_fig_no_for_cache

        a = {"tasks": [{"fig_no": "图1-1", "n": 1}]}
        b = {"tasks": [{"fig_no": "图3-9", "n": 1}]}
        assert _strip_fig_no_for_cache(a) == _strip_fig_no_for_cache(b)

    def test_key_order_does_not_matter(self):
        from app.services.export import _strip_fig_no_for_cache

        assert _strip_fig_no_for_cache({"a": 1, "b": 2}) == _strip_fig_no_for_cache(
            {"b": 2, "a": 1}
        )
