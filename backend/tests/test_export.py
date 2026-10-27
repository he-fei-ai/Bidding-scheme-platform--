from app.models import ExportConfig
from app.services.export import (
    _add_parsed_content,
    _add_image_block,
    _extract_content_image_refs,
)
from app.services.content_cleaning import parse_content
from docx import Document

import base64
import io
import pytest


def _make_png_data_url() -> str | None:
    """生成一个 4x4 红色 PNG 的 data URL；PIL 不可用时返回 None"""
    try:
        from PIL import Image as PilImage
    except ImportError:
        return None
    buf = io.BytesIO()
    PilImage.new("RGB", (4, 4), (200, 30, 30)).save(buf, format="PNG")
    return "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()


def test_export_config_defaults():
    config = ExportConfig(
        font_title="黑体",
        font_body="宋体",
        size_title_1=22,
        size_title_2=18,
        size_title_3=16,
        size_body=12,
    )
    assert config.font_body == "宋体"
    assert config.font_title == "黑体"
    assert config.size_body == 12
    assert config.line_spacing == 1.5


def test_export_config_minimal():
    config = ExportConfig(font_body="宋体", size_body=12)
    assert config.font_body == "宋体"
    assert config.size_title_1 == 14  # 默认值


def test_add_parsed_content_skips_base_level_heading():
    """验证 _add_parsed_content 在 base_level=3 时跳过内容中同级（L3）标题，避免与树节点重复

    V3.0 标准：L3 = "1、"，base_level=3 时，内容中 "1、xxx" 标题应被跳过。
    """
    doc = Document()
    config = ExportConfig(font_body="宋体", size_body=12)

    # AI 生成的内容包含 L3 标题（"1、xxx"），导出时不应重复添加
    content = "1、设计与施工规范依据\n根据招标文件要求，所有施工材料须符合国家标准。\n\n2、材料标准要求\n材料须符合 GB/T 19001 标准。"

    _add_parsed_content(doc, content, config, base_level=3)

    # 获取所有段落文本
    texts = [p.text for p in doc.paragraphs]

    # L3 标题（"1、设计与施工规范依据" 和 "2、材料标准要求"）不应作为 heading 出现（已由树节点添加）
    l3_heading_count = sum(1 for t in texts if "设计与施工规范依据" in t or "材料标准要求" in t)
    # 这些标题应该被跳过，或者只以纯文本形式出现（不带编号）
    # 核心验证：不应带编号出现
    assert not any(t.strip().startswith("1、") or t.strip().startswith("1、") for t in texts if "设计与施工规范依据" in t), \
        f"L3 标题不应带编号出现在内容中（已由树节点提供）"


def test_add_parsed_content_keeps_sub_level_headings():
    """验证 _add_parsed_content 在 base_level=3 时仍正常渲染 L4 及以下子标题

    V3.0 标准：L4 = "1.1、"，base_level=3 时，内容中 "1.1、xxx" 子标题应正常渲染。
    """
    doc = Document()
    config = ExportConfig(font_body="宋体", size_body=12)

    # 内容包含 L3 和 L4 标题
    content = "1、设计与施工规范依据\n相关规范包括 GB50300-2013。\n\n1.1、国家强制性标准\n必须严格执行。\n\n1.2、行业推荐性标准\n参考执行。"

    _add_parsed_content(doc, content, config, base_level=3)

    texts = [p.text for p in doc.paragraphs]
    # L4 子标题应正常渲染（带 "1.1、" 格式编号）
    assert any("1.1、" in t and "国家强制性标准" in t for t in texts), f"L4 子标题应正常渲染，实际文本: {texts}"
    assert any("1.2、" in t and "行业推荐性标准" in t for t in texts), f"L4 子标题应正常渲染，实际文本: {texts}"
    # L3 标题不应带编号重复出现
    assert not any(t.strip().startswith("1、") for t in texts if "设计与施工规范依据" in t), \
        "L3 标题不应在内容中以编号形式重复出现"


# ===== BUG C 修复：data URL 图片 / HTML 内容导出 =====

def test_add_image_block_data_url_png():
    """RichTextEditor 插入的 data:image/png;base64 图片应真正写入 Word 而非占位符"""
    data_url = _make_png_data_url()
    if data_url is None:
        pytest.skip("PIL 不可用，无法生成测试 PNG")
    doc = Document()
    config = ExportConfig(font_body="宋体", size_body=12)

    _add_image_block(doc, data_url, project_id=0, config=config)

    texts = [p.text for p in doc.paragraphs]
    assert not any("[图片" in t for t in texts), f"不应出现占位符: {texts}"
    assert len(doc.inline_shapes) == 1, f"应插入 1 张图片，实际 {len(doc.inline_shapes)}"


def test_add_image_block_svg_data_url_placeholder():
    """SVG data URL（python-docx 不支持）应降级为占位符且不抛异常"""
    svg_url = "data:image/svg+xml;base64," + base64.b64encode(b'<svg xmlns="http://www.w3.org/2000/svg"/>').decode()
    doc = Document()
    config = ExportConfig(font_body="宋体", size_body=12)

    _add_image_block(doc, svg_url, project_id=0, config=config)

    texts = [p.text for p in doc.paragraphs]
    assert any("[图片" in t for t in texts), f"SVG 应降级为占位符: {texts}"
    assert len(doc.inline_shapes) == 0


def test_add_image_block_invalid_data_url_placeholder():
    """非法 base64 数据应降级为占位符而非崩溃"""
    # "AA" 长度不是 4 的倍数，b64decode 必然抛出 binascii.Error
    bad_url = "data:image/png;base64,AA"
    doc = Document()
    config = ExportConfig(font_body="宋体", size_body=12)

    _add_image_block(doc, bad_url, project_id=0, config=config)

    texts = [p.text for p in doc.paragraphs]
    assert any("[图片" in t for t in texts), f"非法数据应降级为占位符: {texts}"


def test_add_parsed_content_html_with_image():
    """HTML 内容（<img> data URL + 文本标签）经 _add_parsed_content 应插入图片且输出干净文本"""
    data_url = _make_png_data_url()
    if data_url is None:
        pytest.skip("PIL 不可用，无法生成测试 PNG")
    doc = Document()
    config = ExportConfig(font_body="宋体", size_body=12)

    content = (
        f'<p>本条依据<b>《建设工程质量管理条例》</b>执行，现场布置如下：</p>'
        f'<div style="text-align: center;"><img src="{data_url}" alt="布置图" /></div>'
        f'<p>以上为平面布置示意图。</p>'
    )

    _add_parsed_content(doc, content, config)

    texts = [p.text for p in doc.paragraphs]
    # 文本应剥离 HTML 标签（原始标签不得残留）
    assert any("本条依据《建设工程质量管理条例》执行" in t for t in texts), f"HTML 未剥离: {texts}"
    assert not any("<p>" in t or "<b>" in t or "<img" in t for t in texts), f"残留原始 HTML: {texts}"
    # 图片应真正插入
    assert len(doc.inline_shapes) == 1, f"应插入 1 张图片，实际 {len(doc.inline_shapes)}"


def test_extract_content_image_refs_html():
    """_extract_content_image_refs 应识别 HTML <img> 标签（含 data URL）"""
    data_url = "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAA=="
    content = f'<img src="{data_url}" alt="流程图" />\n![本地图](D:\\images\\a.png)'
    refs = _extract_content_image_refs(content)
    assert data_url in refs, f"data URL 未被识别: {refs}"
    assert any(r.endswith("a.png") for r in refs), f"本地路径未被识别: {refs}"


def test_export_cleans_residue_and_table_cells():
    """导出链路端到端：LaTeX/HTML/系统提示不得进入正文，表格单元格同步清洗。

    对应《招投标软件导出文档 — 分类问题报告》A-2 / A-3 / E-1。
    """
    doc = Document()
    config = ExportConfig(font_body="宋体", size_body=12)
    content = (
        r"合格率 \(100\%\)，单层厚度控制在 \(5\sim 7mm\)。"
        "\n\n北侧物料主入口<br>1.5 m×2.2m装配式大门"
        "\n\n请确认是否生成正式投标文本。"
        "\n\n<<表格：材料进场检验记录表\n"
        "| 检验项目 | 指标要求 |\n"
        r"| 合格率 | \(100\%\) |"
        "\n>>"
    )
    _add_parsed_content(doc, content, config)

    texts = "\n".join(p.text for p in doc.paragraphs)
    assert "\\(" not in texts and "\\)" not in texts, texts
    assert "<br>" not in texts, texts
    assert "请确认是否生成" not in texts, texts
    assert "合格率 100%" in texts, texts
    assert "北侧物料主入口" in texts and "1.5 m×2.2m装配式大门" in texts

    cell_texts = [c.text for t in doc.tables for r in t.rows for c in r.cells]
    assert any("100%" in c for c in cell_texts), cell_texts
    assert not any("\\%" in c or "<br>" in c for c in cell_texts), cell_texts


def test_collect_toc_entries_matches_body_numbering():
    """报告 G-5/I-5：静态目录条目必须与正文标题（程序编号 + 去旧编号）逐字一致。"""
    from types import SimpleNamespace

    from app.services.export import _collect_toc_entries

    tree = [
        SimpleNamespace(
            id=1,
            title="施工组织设计",
            level=1,
            children=[
                SimpleNamespace(
                    id=2,
                    title="一、编制依据与适用标准",
                    level=2,
                    children=[
                        SimpleNamespace(
                            id=3, title="1、设计与施工规范依据", level=3, children=[]
                        ),
                        SimpleNamespace(id=4, title="材料标准要求", level=3, children=[]),
                    ],
                ),
            ],
        ),
        SimpleNamespace(id=5, title="施工方案", level=1, children=[]),
    ]

    assert _collect_toc_entries(tree) == [
        (1, "第一节、施工组织设计"),
        (2, "一、编制依据与适用标准"),
        (3, "1、设计与施工规范依据"),
        (3, "2、材料标准要求"),  # 旧编号被剥离后按树位置重排
        (1, "第二节、施工方案"),
    ]


def test_claim_chart_title_dedupes_within_chapter():
    """报告 F-2/B-4：同一章同名图题只保留一张，跨章同名允许。

    图题去重必须在**消耗图号之前**判定（跳过时不占号），否则图号会断档。
    """
    from app.services.export_cache import (
        _claim_chart_title,
        _reset_chart_title_registry,
        _used_chart_titles_ctx_var,
    )

    token = _used_chart_titles_ctx_var.set({})
    try:
        _reset_chart_title_registry()
        assert _claim_chart_title("7", "环保与速干材料交叉施工协调") is True
        # 带图号/冒号/标点的同一图题归一化后重复 → 跳过
        assert _claim_chart_title("7", "图7-2：环保与速干材料交叉施工协调") is False
        # 跨章同名图合法（各章都可能有一张进度横道图）
        assert _claim_chart_title("8", "环保与速干材料交叉施工协调") is True
        # 过短标题不参与去重
        assert _claim_chart_title("9", "图9-1 示意") is True
        assert _claim_chart_title("9", "图9-2 示意") is True
    finally:
        _used_chart_titles_ctx_var.reset(token)

    # 会话上下文缺席（非导出链路直接调用）→ 一律放行，不做去重
    assert _claim_chart_title("7", "环保与速干材料交叉施工协调") is True
    assert _claim_chart_title("7", "环保与速干材料交叉施工协调") is True

