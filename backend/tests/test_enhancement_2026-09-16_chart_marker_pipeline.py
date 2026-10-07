"""回归测试：图表标记在「正文生成 → 落库 → 正文/导出渲染」全链路的一致性。

覆盖的缺陷与改动文件（CHART-MARKER-UNIFY 2026-09-16）：

1. `services/content_cleaning.py`
   - 新增 `CHART_TYPES` / `CHART_TYPE_KEYWORDS` / `infer_chart_type_from_text`
     / `iter_chart_markers` / `resolve_chart_marker`：图表标记的**唯一权威解析器**。
   - `CHART_MARKER_PATTERN` 放宽（允许空描述、DOTALL、IGNORECASE）。
   - `parse_content`：无内联数据的标记不再被直接丢弃，而是解析为
     「配图位置提示」block（`hint=True`），并支持多行标记聚合（防标记原文泄漏进正文）。

2. `services/ai/workflows_sections.py`
   - `extract_section_charts`：区分「真实数据条目」与「位置提示条目（`_hint`）」，
     提示条目不再伪造 data/chart_title 结构。
   - `build_chapter_content_messages`：配图标记指令改为精确语法
     （`<<CHART:类型>>`），消除「要求插入标记」与「不要输出图片标记」的自相矛盾。

3. `routers/sse/sse_charts.py`
   - 阶段一：真实内联数据落库（修复 `chart.get("data")` 取嵌套字段导致平铺数据
     从未落库）；提示条目转为编排强制意图；不再无条件 return 短路配图生成。

4. `services/ai/image_engine.py`
   - `generate_all_chart_arrangements(chart_hints=...)` + `_apply_chart_hints`：
     正文标注的配图意图强制 `needed=True`。

5. `services/export.py`
   - `_detect_inline_chart_types` 与 `parse_content` 统一口径（同一解析器），
     消除「检测到类型 → 章节末尾跳过；解析又丢弃标记 → 内联也不渲染」的静默丢图。
   - `_preload_chart_images_parallel` 对「位置提示型」标记一并预热（避免串行渲染）。
"""
import json

import pytest
from docx import Document


# ============================================================
# 1. 标记解析：resolve_chart_marker / iter_chart_markers / parse_content
# ============================================================
class TestChartMarkerParsing:
    def test_iter_chart_markers_supports_all_prefixes(self):
        from app.services.content_cleaning import iter_chart_markers

        content = (
            "<<CHART:GANTT>>\n"
            "<<FLOWCHART:施工工序>>\n"
            "[[TIMELINE:关键里程碑]]\n"
            '<<CHART:{"type":"labor","phases":["p1"]}>>\n'
        )
        markers = list(iter_chart_markers(content))
        assert [p for _, p, _ in markers] == ["chart", "flowchart", "timeline", "chart"]

    def test_resolve_explicit_type_prefix(self):
        from app.services.content_cleaning import resolve_chart_marker

        assert resolve_chart_marker("gantt", "") == ("gantt", "")
        assert resolve_chart_marker("gantt", "施工进度计划") == ("gantt", "施工进度计划")

    def test_resolve_type_token_with_chart_prefix(self):
        from app.services.content_cleaning import resolve_chart_marker

        assert resolve_chart_marker("chart", "GANTT") == ("gantt", "")
        assert resolve_chart_marker("chart", "{gantt}") == ("gantt", "")
        assert resolve_chart_marker("chart", "layout|总平面布置") == ("layout", "总平面布置")

    def test_resolve_plain_title_uses_keyword_inference(self):
        from app.services.content_cleaning import resolve_chart_marker

        resolved = resolve_chart_marker("chart", "{施工进度甘特图}")
        assert resolved is not None
        assert resolved[0] == "gantt"
        assert resolved[1] == "施工进度甘特图"

    def test_resolve_json_payload_prefers_inner_type(self):
        from app.services.content_cleaning import resolve_chart_marker

        payload = json.dumps(
            {"type": "flowchart", "title": "工序流程", "nodes": ["a", "b"]},
            ensure_ascii=False,
        )
        assert resolve_chart_marker("chart", payload) == ("flowchart", "工序流程")

    def test_resolve_returns_none_for_unrecognizable(self):
        from app.services.content_cleaning import resolve_chart_marker

        assert resolve_chart_marker("chart", "") is None
        assert resolve_chart_marker("chart", "这是一段超过六十个字的完全没有图表语义的纯描述文本" * 4) is None

    def test_parse_content_emits_hint_block_for_type_token(self):
        from app.services.content_cleaning import parse_content

        blocks = parse_content("前置说明\n<<CHART:GANTT>>\n后续说明")
        chart_blocks = [b for b in blocks if b.type == "chart_marker"]
        assert len(chart_blocks) == 1
        payload = json.loads(chart_blocks[0].text)
        assert payload["chart_type"] == "gantt"
        assert payload["hint"] is True
        # 不得泄漏标记原文
        assert all("<<" not in b.text for b in blocks if b.type != "chart_marker")

    def test_parse_content_emits_hint_block_for_legacy_plain_title(self):
        """`<<CHART:{施工进度甘特图}>>`（AI 实际输出主流形态）必须解析为 gantt 提示"""
        from app.services.content_cleaning import parse_content

        blocks = parse_content("正文\n<<CHART:{施工进度甘特图}>>\n正文")
        chart_blocks = [b for b in blocks if b.type == "chart_marker"]
        assert len(chart_blocks) == 1
        payload = json.loads(chart_blocks[0].text)
        assert payload["chart_type"] == "gantt"
        assert payload["hint"] is True

    def test_parse_content_keeps_real_inline_json_as_data(self):
        from app.services.content_cleaning import parse_content

        raw = json.dumps(
            {"type": "flowchart", "title": "工序", "nodes": ["a", "b", "c"]},
            ensure_ascii=False,
        )
        blocks = parse_content(f"正文\n<<CHART:{raw}>>\n正文")
        chart_blocks = [b for b in blocks if b.type == "chart_marker"]
        assert len(chart_blocks) == 1
        payload = json.loads(chart_blocks[0].text)
        assert payload["chart_type"] == "flowchart"
        assert payload.get("hint") is None
        assert json.loads(payload["description"])["nodes"] == ["a", "b", "c"]

    def test_parse_content_multiline_marker_is_not_leaked(self):
        """多行标记：不得把 `<<CHART:` 原文写进段落（导出后成乱码占位符）"""
        from app.services.content_cleaning import parse_content

        content = (
            "正文段落\n"
            "<<CHART:{\n"
            '  "type": "gantt",\n'
            '  "title": "进度计划",\n'
            '  "tasks": [{"name": "准备", "start": 1, "end": 5}]\n'
            "}>>\n"
            "尾部段落\n"
        )
        blocks = parse_content(content)
        assert all("<<" not in b.text for b in blocks), "标记原文泄漏进正文"
        chart_blocks = [b for b in blocks if b.type == "chart_marker"]
        assert len(chart_blocks) == 1
        assert json.loads(chart_blocks[0].text)["chart_type"] == "gantt"

    def test_parse_content_unterminated_marker_is_consumed(self):
        from app.services.content_cleaning import parse_content

        blocks = parse_content("正文\n<<CHART:GANTT\n尾部")
        assert all("<<" not in b.text for b in blocks), "未闭合标记原文泄漏进正文"


# ============================================================
# 2. extract_section_charts：真实数据 / 位置提示 必须区分
# ============================================================
class TestExtractSectionCharts:
    def test_hint_entries_carry_no_fabricated_data(self):
        from app.services.ai.workflows_sections import extract_section_charts

        charts = extract_section_charts("正文\n<<CHART:{施工进度甘特图}>>\n正文")
        assert len(charts) == 1
        chart = charts[0]
        assert chart["_hint"] is True
        assert chart["type"] == "gantt"
        # 绝不能带任何图表数据字段
        for forbidden in ("nodes", "tasks", "code", "tree", "phases", "milestones", "zones"):
            assert forbidden not in chart, f"提示条目不应包含 {forbidden}"

    def test_type_token_marker_becomes_hint(self):
        from app.services.ai.workflows_sections import extract_section_charts

        charts = extract_section_charts("正文\n<<CHART:GANTT>>\n正文")
        assert len(charts) == 1
        assert charts[0]["_hint"] is True
        assert charts[0]["type"] == "gantt"

    def test_real_json_payload_is_not_a_hint(self):
        from app.services.ai.workflows_sections import extract_section_charts

        payload = json.dumps(
            {"type": "flowchart", "title": "工序", "nodes": ["a", "b"]}, ensure_ascii=False
        )
        charts = extract_section_charts(f"正文\n<<CHART:{payload}>>\n正文")
        assert len(charts) == 1
        assert not charts[0].get("_hint")
        assert charts[0]["type"] == "flowchart"
        assert charts[0]["nodes"] == ["a", "b"]


# ============================================================
# 3. 检测口径一致性：检测结果 == parse_content 实际产出的类型
# ============================================================
class TestDetectionConsistency:
    @pytest.mark.parametrize(
        "content,expected",
        [
            ("正文\n<<CHART:GANTT>>\n正文", {"gantt"}),
            ("正文\n<<GANTT:>>\n正文", {"gantt"}),
            ("正文\n<<GANTT:施工进度计划>>\n正文", {"gantt"}),
            ("正文\n<<CHART:{施工进度甘特图}>>\n正文", {"gantt"}),
            ("正文\n<<CHART:{组织架构图}>>\n正文", {"architecture"}),
            ("正文\n[[FLOWCHART: 施工工序]]\n正文", {"flowchart"}),
            ("正文\n<<CHART:{劳动力配置}}>>\n正文", {"labor"}),
        ],
    )
    def test_detect_matches_expected_types(self, content, expected):
        from app.services.export import _detect_inline_chart_types

        assert _detect_inline_chart_types(content) == expected

    def test_detect_matches_parse_content_block_types(self):
        """核心不变量：检测到的类型集合 == parse_content 产出的 chart_marker 类型集合"""
        from app.services.content_cleaning import parse_content
        from app.services.export import _detect_inline_chart_types

        content = (
            "正文\n<<CHART:GANTT>>\n"
            "正文\n<<GANTT:>>\n"
            "正文\n<<CHART:{组织架构图}>>\n"
            "正文\n<<FLOWCHART: 施工工序>>\n"
        )
        detected = _detect_inline_chart_types(content)
        parsed_types = {
            json.loads(b.text)["chart_type"]
            for b in parse_content(content)
            if b.type == "chart_marker"
        }
        assert detected == parsed_types

    def test_hint_types_subset_reported(self):
        from app.services.export import _detect_inline_chart_markers

        content = (
            "正文\n<<CHART:GANTT>>\n"
            + "正文\n<<CHART:"
            + json.dumps({"type": "flowchart", "nodes": ["a", "b"]}, ensure_ascii=False)
            + ">>\n"
        )
        all_types, hint_types = _detect_inline_chart_markers(content)
        assert all_types == {"gantt", "flowchart"}
        # 标记内自带 JSON 数据的类型不属于「位置提示」→ 不预热（载荷不同）
        assert hint_types == {"gantt"}


# ============================================================
# 4. 编排阶段：配图提示强制 needed=True
# ============================================================
class TestChartHintsForcing:
    def test_apply_chart_hints_forces_needed(self):
        from app.services.ai.image_engine import _apply_chart_hints

        arrangements = {
            "flowchart": {"needed": False, "type": "none", "priority": 3, "reason": ""},
            "gantt": {"needed": False, "type": "none", "priority": 3, "reason": ""},
        }
        out = _apply_chart_hints(arrangements, {"gantt": "施工进度计划"})
        assert out["gantt"]["needed"] is True
        assert out["gantt"]["fromHint"] is True
        assert out["gantt"]["hintTitle"] == "施工进度计划"
        # 未命中的类型保持原判
        assert out["flowchart"]["needed"] is False

    def test_apply_chart_hints_creates_missing_entry(self):
        from app.services.ai.image_engine import _apply_chart_hints

        out = _apply_chart_hints({}, {"timeline": ""})
        assert out["timeline"]["needed"] is True
        assert "hintTitle" not in out["timeline"]

    def test_apply_chart_hints_ignores_unknown_type(self):
        from app.services.ai.image_engine import _apply_chart_hints

        out = _apply_chart_hints({}, {"sequence": "x"})
        assert "sequence" not in out

    @pytest.mark.asyncio
    async def test_batch_arrangement_injects_hints_and_forces_needed(self, monkeypatch):
        """批量编排：提示写入提示词，且模型判 false 也会被强制为 true"""
        from app.services.ai import image_engine

        captured = {}

        async def _fake_collect(provider, messages, max_retries=3):
            captured["system"] = messages[0].content
            captured["user"] = messages[1].content
            # 模型故意把 gantt 判为不需要
            return {
                "gantt": {"needed": False, "type": "none", "priority": 3, "reason": "模型拒绝"}
            }

        monkeypatch.setattr(image_engine, "collect_json_response", _fake_collect)

        result = await image_engine.generate_all_chart_arrangements(
            provider=object(),
            section_title="施工进度计划",
            section_content="正文内容",
            chart_hints={"gantt": "总进度计划"},
        )

        assert result["gantt"]["needed"] is True, "正文已标注配图位置却仍未生成"
        assert "gantt" in captured["system"], "提示未注入系统提示词"
        assert "gantt" in captured["user"], "提示未注入用户提示词"


# ============================================================
# 5. 第一优先级渲染路径：HTTP Service 可用性缓存必须带 TTL
# ============================================================
class TestHttpPriorityPath:
    def test_http_availability_cache_expires(self, monkeypatch):
        """可用性缓存不得永久生效：服务晚启动时必须能回到第一优先级渲染路径"""
        import time as _t

        from app.services import export
        from app.services.ai import mermaid_service, workflows_base

        state = {"available": True, "probes": 0}

        class _FakeClient:
            async def check_docker_available(self) -> bool:
                return state["available"]

            async def check_railway_available(self) -> bool:
                return False

        monkeypatch.setattr(
            mermaid_service, "get_mermaid_service_client", lambda *a, **k: _FakeClient()
        )

        def _fake_run(coro):
            # 真实探测 = 执行一次 health check；此处只计数并返回服务可用性
            coro.close()  # 避免 "coroutine was never awaited"
            state["probes"] += 1
            return state["available"]

        monkeypatch.setattr(workflows_base, "_safe_asyncio_run", _fake_run)
        export._HTTP_CHECK_CACHE.clear()
        try:
            assert export._check_http_service_available() is True
            assert state["probes"] == 1
            # TTL 内命中缓存，不再真实探测（避免每图 5s 健康检查开销）
            assert export._check_http_service_available() is True
            assert state["probes"] == 1
            # 缓存过期 → 重新探测（此时服务已不可用）
            state["available"] = False
            export._HTTP_CHECK_CACHE["available"] = (True, _t.time() - 10_000)
            assert export._check_http_service_available() is False
            assert state["probes"] > 1, "缓存过期后必须重新探测第一优先级渲染服务"
        finally:
            export._HTTP_CHECK_CACHE.clear()


# ============================================================
# 6. DOCX：提示标记就地渲染章节真实数据（无降级、不重复出图、不泄漏标记）
# ============================================================
def _make_section(section_id: int, content: str, **chart_fields):
    class _Section:
        pass

    section = _Section()
    section.id = section_id
    section.title = "施工进度计划"
    section.content = content
    for field in (
        "flowchart_json",
        "gantt_json",
        "architecture_json",
        "labor_json",
        "comparison_json",
        "layout_json",
        "timeline_json",
    ):
        setattr(section, field, None)
    for name, value in chart_fields.items():
        setattr(section, name, json.dumps(value, ensure_ascii=False))
    return section


_GANTT_DATA = {
    "type": "gantt",
    "title": "施工进度计划",
    "duration": 120,
    "tasks": [
        {"name": "施工准备", "start": 1, "end": 15},
        {"name": "基础施工", "start": 10, "end": 45},
        {"name": "主体施工", "start": 40, "end": 90},
        {"name": "竣工验收", "start": 90, "end": 120},
    ],
}


class TestDocxHintRendering:
    def test_hint_marker_renders_real_section_data_inline(self):
        from app.services.export import ExportConfig, _add_parsed_content

        doc = Document()
        section = _make_section(9101, "前置段落\n<<CHART:GANTT>>\n后续段落", gantt_json=_GANTT_DATA)

        _add_parsed_content(doc, section.content, ExportConfig(), section_detail=section)

        texts = "\n".join(p.text for p in doc.paragraphs)
        assert "<<CHART" not in texts, "标记原文泄漏进 DOCX"
        has_picture = any(p._p.xpath(".//pic:pic") for p in doc.paragraphs)
        assert has_picture, "正文标注了配图位置、章节也有真实数据，却未渲染图表"

    def test_hint_marker_without_data_renders_nothing(self):
        """无降级原则：章节没有真实图表数据时，不得伪造任何图表内容"""
        from app.services.export import ExportConfig, _add_parsed_content

        doc = Document()
        section = _make_section(9102, "前置段落\n<<CHART:GANTT>>\n后续段落")

        _add_parsed_content(doc, section.content, ExportConfig(), section_detail=section)

        texts = "\n".join(p.text for p in doc.paragraphs)
        assert "<<CHART" not in texts
        assert "GANTT" not in texts
        assert not any(p._p.xpath(".//pic:pic") for p in doc.paragraphs)

    def test_no_duplicate_chart_for_hint_marker(self):
        """提示标记 → 内联渲染；`_render_section_charts` 不能在同一章节再渲染一次"""
        from app.services.export import (
            ExportConfig,
            _add_parsed_content,
            _render_section_charts,
        )

        doc = Document()
        config = ExportConfig()
        section = _make_section(9103, "前置段落\n<<CHART:GANTT>>\n后续段落", gantt_json=_GANTT_DATA)

        _add_parsed_content(doc, section.content, config, section_detail=section)
        _render_section_charts(doc, section, "施工进度计划", config)

        picture_paragraphs = [p for p in doc.paragraphs if p._p.xpath(".//pic:pic")]
        assert len(picture_paragraphs) == 1, f"同一章节重复出图 {len(picture_paragraphs)} 次"

    def test_db_bound_chart_without_marker_still_rendered(self):
        """无标记时由 `_render_section_charts` 在章节末尾补图（原有能力不回归）"""
        from app.services.export import ExportConfig, _render_section_charts

        doc = Document()
        section = _make_section(9104, "只有正文，没有标记", gantt_json=_GANTT_DATA)

        _render_section_charts(doc, section, "施工进度计划", ExportConfig())

        assert any(p._p.xpath(".//pic:pic") for p in doc.paragraphs)

    def test_legacy_plain_title_marker_renders_real_data(self):
        """历史格式 `<<CHART:{施工进度甘特图}>>`：必须就地产出真实图表（原来会彻底丢图）"""
        from app.services.export import (
            ExportConfig,
            _add_parsed_content,
            _render_section_charts,
        )

        doc = Document()
        config = ExportConfig()
        section = _make_section(
            9105, "前置段落\n<<CHART:{施工进度甘特图}>>\n后续段落", gantt_json=_GANTT_DATA
        )

        _add_parsed_content(doc, section.content, config, section_detail=section)
        _render_section_charts(doc, section, "施工进度计划", config)

        texts = "\n".join(p.text for p in doc.paragraphs)
        assert "<<CHART" not in texts
        assert sum(1 for p in doc.paragraphs if p._p.xpath(".//pic:pic")) == 1
