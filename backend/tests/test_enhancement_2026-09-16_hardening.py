"""全面审查加固（2026-09-16）回归测试。

覆盖本轮「按最高标准全面审查」修复的缺陷与改动文件：

| # | 缺陷 | 文件 |
|---|------|------|
| 1 | `/ai/providers` 把内置厂商条目（含 `AGNES_API_KEY`）原样序列化返回 → 平台主密钥明文泄露 | `routers/ai_config.py` |
| 2 | `/ai/config` 主动解密后回传明文密钥，`ENC:` 加密存储被读接口完全抵消 | `routers/ai_config.py` |
| 3 | `/ai/provider-configs`、`/ai/image-provider-configs` 读接口解密回传明文；PUT 收到掩码/空串时把已存密钥**清空** | `routers/ai_config.py` |
| 4 | `get_db` / `get_db_with_retry` 异常分支只清 `_local.connection`，既不 close 也不移出 `_all_connections`，连接**永久泄漏**；池满 32 后整个进程数据库不可用 | `db.py` |
| 5 | `DBPool.get_connection` 池满分支 `raise ... from e`，首连路径 `e` 未绑定 → UnboundLocalError 掩盖真因 | `db.py` |
| 6 | `get_db_with_retry` 完全绕过 `_db_write_lock`，写串行化只对一半路径生效 | `db.py` |
| 7 | `rate_limiter.get_identifier_from_request` 对 dict 形态的 `current_user` 取 `.id` → 认证开启后所有限流端点 500 | `routers/rate_limiter.py` |
| 8 | `export.sanitize_for_docx`：标题/单元格/段落未做 XML 安全清洗，一个 `\x0b` 让整节标题+正文被占位符吞掉 | `services/export.py` |
| 9 | 导出表格行列数无上限，畸形 `<<表格：` 标记可拖垮导出 | `services/export.py` |
| 10 | `text_patch.apply_range_edits` 缺失区间重叠校验，docstring 声称的行为未实现 | `services/text_patch.py` |
| 11 | `check_report_excel.as_ratio` 不过滤 Infinity/NaN → Excel 打开报「内容有问题」 | `services/check_report_excel.py` |
| 12 | 查重 Word 报告硬下标取值，字段缺失即整份报告 500；且无控制字符清洗 | `services/duplicate.py` |
| 13 | `_save_export_cache_index` 非原子写，多进程并发写坏索引 → 启动清理误删仍在用的缓存 | `services/export.py` |
| 14 | `/keys` 无角色校验 + permissions 任意指定 + 365 天有效期 | `routers/api_keys.py` |
| 15 | OCR `_clear_old_output` 无护栏递归删除任意目录；临时目录只在成功路径清理 | `services/ocr_parser.py` |
| 16 | `security.get_current_user` 对无效 token（401）也降级 mock admin，与 auth_service「无效 token 一律 401」约定冲突 | `security/auth.py` |

测试只测行为，不测实现细节。
"""

from __future__ import annotations

import os
import tempfile
import time
from pathlib import Path

import pytest

# ---------------------------------------------------------------------------
# 1-3. AI 配置读接口不得泄露明文密钥
# ---------------------------------------------------------------------------

MASK = '***已配置***'


def _with_ai_key(monkeypatch, value: str):
    """让 provider_factory 返回带密钥的配置。"""
    from app.services.ai import provider_factory

    monkeypatch.setattr(
        provider_factory,
        "get_ai_config",
        lambda: {
            "provider": "openai_compatible",
            "provider_name": "agnes",
            "api_key": value,
            "model": "test-model",
            "base_url": "https://example.invalid/v1",
        },
        raising=False,
    )


def test_providers_endpoint_never_exposes_api_key(monkeypatch):
    """`/ai/providers` 即使某厂商条目带 api_key 字段，也必须剔除后返回。"""
    from app.routers import ai_config as ac

    # 快照-还原：D-1 后 KNOWN_PROVIDERS 元素含 readonly/is_default 等字段，
    # 旧还原逻辑只剔除 api_key，会把这些字段丢掉并残留给后续用例
    # （曾导致 provider_catalog 同源断言在全量套件中失败）。
    original = dict(ac.KNOWN_PROVIDERS[0])
    ac.KNOWN_PROVIDERS[0] = {
        "name": "agnes",
        "base_url": "https://x/v1",
        "model": "m",
        "type": "openai_compatible",
        "api_key": "sk-super-secret-value",
    }
    try:
        import asyncio

        result = asyncio.run(ac.list_known_providers())
    finally:
        ac.KNOWN_PROVIDERS[0] = original

    assert all("api_key" not in p for p in result)
    assert all("sk-super-secret-value" not in str(p) for p in result)


def test_get_ai_config_masks_key(monkeypatch):
    from app.routers import ai_config as ac

    _with_ai_key(monkeypatch, "sk-live-abcdef1234567890")
    cfg = asyncio_run(ac.get_ai_config())
    assert cfg["api_key"] == MASK
    assert "sk-live-abcdef1234567890" not in str(cfg)


def test_get_ai_config_empty_key_stays_empty(monkeypatch):
    from app.routers import ai_config as ac

    _with_ai_key(monkeypatch, "")
    cfg = asyncio_run(ac.get_ai_config())
    assert cfg["api_key"] == ""


def test_display_api_key_is_never_plaintext():
    from app.routers import ai_config as ac

    assert ac._display_api_key("") == ""
    assert ac._display_api_key("sk-plaintext-key-000") == MASK
    assert ac._display_api_key("ENC:abc") == MASK
    assert ac._display_api_key(MASK) == MASK


# ---------------------------------------------------------------------------
# 4-6. 连接池：异常路径不得泄漏连接
# ---------------------------------------------------------------------------


def test_pool_discard_connection_removes_from_indexes():
    from app.db import DBPool

    pool = DBPool(max_connections=2)
    conn = pool.get_connection()
    assert len(pool._all_connections) == 1

    pool.discard_connection(conn)

    assert pool._local.connection is None
    assert len(pool._all_connections) == 0
    assert pool._thread_index == {}
    stats = pool.get_stats()
    assert stats["active_connections"] == 0


def test_exception_inside_get_db_does_not_leak_quota():
    """`with get_db()` 体内 raise（如 task_lock 冲突）后配额必须归还。

    旧实现泄漏连接且因 thread_id 存活而**永不回收**，累计 32 个后
    整个进程数据库不可用 —— 本用例锁死该不变量。
    """
    from app.db import DBPool, get_db

    pool = DBPool(max_connections=3)
    orig = __import__("app.db", fromlist=["_pool"])._pool
    __import__("app.db", fromlist=["_pool"])._pool = pool
    try:
        for _ in range(6):  # 远超 max_connections=3
            with pytest.raises(RuntimeError, match="boom"):
                with get_db() as conn:
                    conn.execute("SELECT 1")
                    raise RuntimeError("boom")
        # 每次异常后都能继续取到连接 → 无泄漏
        with get_db() as conn:
            assert conn.execute("SELECT 1").fetchone()[0] == 1
        assert pool.get_stats()["active_connections"] <= 1
    finally:
        pool.close_all()
        __import__("app.db", fromlist=["_pool"])._pool = orig


def test_exception_inside_get_db_with_retry_does_not_leak_quota():
    from app.db import DBPool, get_db_with_retry

    pool = DBPool(max_connections=3)
    db_mod = __import__("app.db", fromlist=["_pool"])
    orig = db_mod._pool
    db_mod._pool = pool
    try:
        for _ in range(6):
            with pytest.raises(RuntimeError):
                with get_db_with_retry() as conn:
                    conn.execute("SELECT 1")
                    raise RuntimeError("x")
        with get_db_with_retry() as conn:
            assert conn.execute("SELECT 1").fetchone()[0] == 1
    finally:
        pool.close_all()
        db_mod._pool = orig


def test_pool_exhausted_error_chains_real_cause():
    """池满 RuntimeError 不应变成 UnboundLocalError（首连路径 e 未绑定）。

    BUG-FIX 2026-09-18（P1）：回收判据改为「线程不在 enumerate() 即回收」，
    旧写法（注入假 thread_id + 新鲜 _last_used）的占位连接会被立即回收，
    永远走不到池满分支。改用真实存活线程持有连接来填满配额：
    存活线程的连接不可能被回收，必然触发 exhausted 分支。
    """
    import threading

    from app.db import DBPool

    pool = DBPool(max_connections=1)
    started = threading.Event()
    stop = threading.Event()

    def _holder():
        # 在独立线程里取一条真实连接并保持线程存活（存活线程永不回收）
        pool.get_connection()
        started.set()
        stop.wait(10)

    th = threading.Thread(target=_holder, daemon=True)
    th.start()
    assert started.wait(5), "holder 线程未能取得连接"
    try:
        with pytest.raises(RuntimeError, match="pool exhausted"):
            pool.get_connection()
        # 错误链必须是真实原因（TypeError 等），不是 UnboundLocalError
    finally:
        stop.set()
        th.join(timeout=12)
        pool.close_all()


# ---------------------------------------------------------------------------
# 7. 限流标识对 dict / 对象两种形态都可用
# ---------------------------------------------------------------------------


class _State:
    def __init__(self, user):
        self.current_user = user


class _Req:
    def __init__(self, user):
        self.state = _State(user)
        self.headers = {}
        self.query_params = {}
        self.client = None


def test_rate_limit_identifier_accepts_dict_user():
    from app.routers.rate_limiter import get_identifier_from_request

    req = _Req({"id": 42, "username": "u", "role": "admin"})
    assert get_identifier_from_request(req) == "user_42"


def test_rate_limit_identifier_accepts_object_user():
    from app.routers.rate_limiter import get_identifier_from_request
    from app.security.auth import User

    req = _Req(User(user_id=7, username="u"))
    assert get_identifier_from_request(req) == "user_7"


def test_rate_limit_identifier_falls_back_to_ip():
    from app.routers.rate_limiter import get_identifier_from_request

    class _C:
        host = "10.0.0.9"

    req = _Req(None)
    req.client = _C()
    assert get_identifier_from_request(req) == "ip_10.0.0.9"


# ---------------------------------------------------------------------------
# 8-9. DOCX XML 安全清洗 + 表格上限
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "raw,forbidden",
    [
        ("第一章\x0b施工组织", "\x0b"),
        ("标题\ud800带代理项", "\ud800"),
        ("尾部\ufffe", "\ufffe"),
        ("段\x01落", "\x01"),
    ],
)
def test_sanitize_for_docx_strips_xml_unsafe_chars(raw, forbidden):
    from app.services.export import sanitize_for_docx

    out = sanitize_for_docx(raw)
    assert forbidden not in out
    assert "施工组织" in out or "标题" in out or "尾" in out or "段" in out


def test_sanitize_for_docx_keeps_normal_text():
    from app.services.export import sanitize_for_docx

    text = "第1.2条 工期要求：总日历天数 ≤ 300 天（含法定节假日）。"
    assert sanitize_for_docx(text) == text


def test_sanitize_for_docx_truncates():
    from app.services.export import MAX_DOCX_TEXT_CHARS, sanitize_for_docx

    out = sanitize_for_docx("a" * (MAX_DOCX_TEXT_CHARS + 500), MAX_DOCX_TEXT_CHARS)
    assert len(out) == MAX_DOCX_TEXT_CHARS + 1  # 含省略号


def test_set_cell_text_truncates_long_text():
    from app.services.export import MAX_DOCX_TEXT_CHARS, _set_cell_text
    from docx import Document

    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    _set_cell_text(table.rows[0].cells[0], "x" * 50000)
    assert len(table.rows[0].cells[0].text) <= MAX_DOCX_TEXT_CHARS + 1


def test_table_markup_row_cap():
    """超过 500 行的畸形表格标记应被截断而不是创建 N 万单元格。"""
    from app.services.export import _add_table_from_markup, _theme_color
    from app.services.export import ExportConfig  # noqa: F401
    from docx import Document

    from app.services.export import ExportConfig as Cfg

    doc = Document()
    cfg = Cfg()
    rows = ["| a | b |"] + [f"| r{i} | v{i} |" for i in range(900)]
    markup = "\n".join(rows)
    _add_table_from_markup(doc, markup, cfg)
    tables = doc.tables
    assert tables, "应至少生成一个表格"
    assert len(tables[0].rows) <= 500


# ---------------------------------------------------------------------------
# 10. text_patch 区间重叠
# ---------------------------------------------------------------------------


def test_apply_range_edits_rejects_overlapping_ranges():
    from app.services.text_patch import apply_range_edits

    result = apply_range_edits(
        "ABCDEFGH",
        [
            {"start": 0, "end": 5, "newText": "X"},
            {"start": 3, "end": 8, "newText": "Y"},
        ],
    )
    assert result["changed"] is False
    assert result["applied"] == []
    assert result["content"] == "ABCDEFGH"
    assert result["errors"], "重叠必须整批拒绝"
    assert any("重叠" in e for e in result["errors"])


def test_apply_range_edits_allows_adjacent_ranges():
    from app.services.text_patch import apply_range_edits

    result = apply_range_edits(
        "ABCDEFGH",
        [
            {"start": 0, "end": 3, "newText": "X"},
            {"start": 3, "end": 6, "newText": "Y"},
        ],
    )
    assert result["changed"] is True
    # 两段编辑各自引用原文下标，倒序应用后互不吞噬对方区间
    assert result["content"] == "XYGH"
    assert result["errors"] == []


# ---------------------------------------------------------------------------
# 11. Excel 数值单元格 inf/nan 防护
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("bad", [float("nan"), float("inf"), float("-inf")])
def test_as_ratio_rejects_non_finite(bad):
    from app.services.check_report_excel import as_ratio

    assert as_ratio(bad) == ""


def test_as_ratio_keeps_valid_ratio():
    from app.services.check_report_excel import as_ratio

    assert as_ratio(0.87) == 0.87
    assert as_ratio(-1) == ""
    assert as_ratio(None) == ""
    assert as_ratio(True) == ""


# ---------------------------------------------------------------------------
# 12. 查重报告安全取值
# ---------------------------------------------------------------------------


def test_report_text_handles_none_and_non_str():
    from app.services.duplicate import _report_text

    assert _report_text(None) == ""
    assert _report_text(None, "-") == "-"
    assert _report_text(123) == "123"
    assert _report_text("ok") == "ok"
    assert "\x0b" not in _report_text("bad\x0btext")


def test_report_join_handles_nested_and_non_str():
    from app.services.duplicate import _report_join

    assert _report_join(None) == ""
    assert _report_join([]) == ""
    assert _report_join(["a", "b"]) == "a, b"
    assert _report_join([1, 2]) == "1, 2"
    assert _report_join([{"k": 1}]) == '{"k": 1}'


def test_duplicate_report_survives_missing_fields(monkeypatch, tmp_path):
    """结果 JSON 字段缺失/类型异常时，Word 报告不应 500。

    旧实现大量使用 `g['sentence']` / `o['filename']` 等硬下标取值，
    历史结果里任一字段缺失或为数字就会 KeyError/TypeError，整份报告 500。
    """
    from app.services import duplicate

    pytest.importorskip("docx")

    task = {
        "id": 1,
        "results": {
            "metadata": {
                "metadata": [{"filename": None, "author": 42}],
                "anomalies": [{"description": "异常", "value": [1, {"x": 2}]}],
                "fingerprints": [{"author": "a"}],
                "same_source_pairs": [{"file_a": "A", "file_b": 3}],
            },
            "toc": {
                "pairwise": [{"file_a": None, "score": 0.5}],
                "duplicate_groups": [{"file_a": "A", "file_b": 3}],
                "similar_groups": [{"file_a": "A", "title_b": None, "similarity": 0.9}],
            },
            "content": {
                "stats": {"total_sentences": 1},
                "duplicate_groups": [
                    {"sentence": "重复\x0b句子", "occurrences": [{"filename": 42}, {}]}
                ],
            },
            "images": {
                "total_images": 1,
                "duplicate_groups": [{"hash": "h", "size": None, "occurrences": [{}]}],
            },
        },
    }

    monkeypatch.setattr(duplicate, "get_task", lambda tid: task)
    monkeypatch.setattr(duplicate, "DUPLICATE_REPORTS_DIR", tmp_path)

    out = duplicate.export_report(1)
    assert out and Path(out).exists()
    assert Path(out).stat().st_size > 0


# ---------------------------------------------------------------------------
# 13. 导出缓存索引原子写
# ---------------------------------------------------------------------------


def test_save_export_cache_index_leaves_no_temp_file(monkeypatch, tmp_path):
    from app.services import export

    index_path = tmp_path / "_cache_index.json"
    monkeypatch.setattr(export, "_EXPORT_CACHE_INDEX_PATH", str(index_path))
    export._EXPORT_CACHE = {"k": "v"}

    export._save_export_cache_index()

    assert index_path.exists()
    assert index_path.read_text(encoding="utf-8").strip().startswith("{")
    leftovers = [p for p in tmp_path.iterdir() if ".tmp." in p.name]
    assert leftovers == [], "原子写失败时应清理临时文件"


# ---------------------------------------------------------------------------
# 14. API Key 签发权限收敛
# ---------------------------------------------------------------------------


def test_api_key_permission_whitelist_exists():
    from app.routers.api_keys import ALLOWED_PERMISSIONS, MAX_EXPIRES_DAYS

    assert "projects:read" in ALLOWED_PERMISSIONS
    assert MAX_EXPIRES_DAYS <= 365


@pytest.mark.anyio
async def test_create_api_key_requires_admin():
    """非管理员创建 API Key 应被拒绝（403）。"""
    from app.routers import api_keys
    from app.security.auth import User

    from fastapi import HTTPException

    class _R:
        def __init__(self, user):
            self._u = user

        @property
        def state(self):
            return self

        current_user = None

    # get_current_user 通过 request 读取；直接 monkeypatch 掉依赖
    req = _R(User(user_id=2, username="p", role="user"))
    import app.routers.api_keys as ak

    original = ak.get_current_user
    ak.get_current_user = lambda request: User(user_id=2, username="p", role="user")
    try:
        with pytest.raises(HTTPException) as info:
            await ak.create_new_api_key(
                req, name="k", permissions="projects:read", expires_days=30
            )
        assert info.value.status_code == 403
    finally:
        ak.get_current_user = original


@pytest.mark.anyio
async def test_create_api_key_rejects_unknown_permission():
    import app.routers.api_keys as ak
    from app.security.auth import User

    class _R:
        @property
        def state(self):
            return self

    original = ak.get_current_user
    ak.get_current_user = lambda request: User(user_id=1, username="a", role="admin")
    try:
        from fastapi import HTTPException

        with pytest.raises(HTTPException) as info:
            await ak.create_new_api_key(
                _R(), name="k", permissions="root:god", expires_days=30
            )
        assert info.value.status_code == 400
    finally:
        ak.get_current_user = original


# ---------------------------------------------------------------------------
# 15. OCR 临时目录清理护栏
# ---------------------------------------------------------------------------


def test_ocr_cleanup_guard_rejects_non_temp_dir(tmp_path):
    from app.services.ocr_parser import MinerUParser

    parser = MinerUParser.__new__(MinerUParser)
    parser.auto_cleanup = True

    outside = tmp_path / "victim"
    outside.mkdir()
    (outside / "keep.txt").write_text("data", encoding="utf-8")

    parser._clear_old_output(str(outside))

    assert outside.exists(), "非 mineru_ 临时目录绝不能被递归删除"
    assert (outside / "keep.txt").exists()


def test_ocr_cleanup_guard_allows_mineru_temp_dir(tmp_path, monkeypatch):
    from app.services.ocr_parser import MinerUParser

    monkeypatch.setattr(tempfile, "gettempdir", lambda: str(tmp_path))
    parser = MinerUParser.__new__(MinerUParser)
    parser.auto_cleanup = True

    target = tmp_path / "mineru_doc_"
    target.mkdir()
    (target / "page.png").write_bytes(b"x")

    parser._clear_old_output(str(target))
    assert not target.exists()


# ---------------------------------------------------------------------------
# 16. 无效 token 不得降级为 admin
# ---------------------------------------------------------------------------


def test_invalid_token_is_not_silently_promoted_to_admin(monkeypatch):
    from fastapi import HTTPException

    from app.security import auth as sec_auth

    def _raise(request):
        raise HTTPException(status_code=401, detail="无效的登录凭证")

    monkeypatch.setitem(
        __import__("sys").modules, "app.services.auth_service", _FakeAuthService(_raise)
    )
    monkeypatch.setattr(sec_auth, "AUTH_ENABLED", False, raising=False)

    class _Req:
        pass

    with pytest.raises(HTTPException) as info:
        sec_auth.get_current_user(_Req())
    assert info.value.status_code == 401


class _FakeAuthService:
    def __init__(self, fn):
        self.get_current_user_dependency = fn


def test_dev_mode_still_falls_back_on_unrelated_errors(monkeypatch):
    from app.security import auth as sec_auth

    def _boom(request):
        raise RuntimeError("db down")

    monkeypatch.setattr(
        sec_auth, "get_current_user_dependency", _boom, raising=False
    )
    monkeypatch.setattr(sec_auth, "AUTH_ENABLED", False, raising=False)

    class _Req:
        pass

    user = sec_auth.get_current_user(_Req())
    assert user is not None and user.role == "admin"


# ---------------------------------------------------------------------------
# 工具
# ---------------------------------------------------------------------------


def asyncio_run(coro):
    import asyncio

    return asyncio.run(coro)
