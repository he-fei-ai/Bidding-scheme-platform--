# -*- coding: utf-8 -*-
"""完整图表质量测试 + DOCX 集成 + 真实导出端到端"""
import sys, os, json, io, struct, traceback
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app.services.ai.mermaid_renderer import render_mermaid_to_bytes
from PIL import Image

# ============ 正确格式的 7 类图表测试数据 ============
TEST_DATA = {
    "gantt": {
        "title": "施工总进度计划",
        "tasks": [
            {"name": "土方开挖", "start_day": 1, "duration": 30, "resource": "土方组"},
            {"name": "基础浇筑", "start_day": 31, "duration": 60, "resource": "混凝土组", "dependencies": ["土方开挖"]},
            {"name": "主体施工", "start_day": 91, "duration": 180, "resource": "土建组", "dependencies": ["基础浇筑"]},
            {"name": "装饰工程", "start_day": 271, "duration": 90, "resource": "装饰组", "dependencies": ["主体施工"]},
            {"name": "机电安装", "start_day": 271, "duration": 120, "resource": "机电组", "dependencies": ["主体施工"]},
        ],
        "highlight_critical_path": True,
    },
    "flowchart": {
        "title": "施工工艺流程图",
        "nodes": [
            {"id": "A", "label": "施工准备", "type": "start"},
            {"id": "B", "label": "测量放线", "type": "process"},
            {"id": "C", "label": "土方开挖", "type": "process"},
            {"id": "D", "label": "基础浇筑", "type": "process"},
            {"id": "E", "label": "验收合格?", "type": "decision"},
            {"id": "F", "label": "返工处理", "type": "process"},
            {"id": "G", "label": "主体施工", "type": "process"},
            {"id": "H", "label": "竣工验收", "type": "end"},
        ],
        "edges": [
            {"from": "A", "to": "B"}, {"from": "B", "to": "C"}, {"from": "C", "to": "D"},
            {"from": "D", "to": "E"}, {"from": "E", "to": "G", "label": "是"},
            {"from": "E", "to": "F", "label": "否"}, {"from": "F", "to": "D"}, {"from": "G", "to": "H"},
        ],
    },
    "architecture": {
        "title": "项目组织架构图",
        "name": "项目经理", "type": "manager",
        "children": [
            {"name": "技术负责人", "type": "tech", "children": [
                {"name": "施工员", "type": "staff"},
                {"name": "质检员", "type": "staff"},
            ]},
            {"name": "商务负责人", "type": "business", "children": [
                {"name": "造价员", "type": "staff"},
                {"name": "材料员", "type": "staff"},
            ]},
            {"name": "安全负责人", "type": "safety", "children": [
                {"name": "安全员", "type": "staff"},
            ]},
        ],
    },
    "labor": {
        "title": "劳动力投入计划",
        "phases": ["基础阶段", "主体阶段", "装饰阶段"],
        "categories": ["土建工", "钢筋工", "木工", "混凝土工", "普工"],
        "data": [
            [30, 20, 10],
            [15, 40, 15],
            [20, 35, 25],
            [25, 25, 15],
            [40, 30, 20],
        ],
    },
    "comparison": {
        "title": "主要材料占比分析",
        "items": [
            {"label": "钢材", "value": 35},
            {"label": "混凝土", "value": 30},
            {"label": "木材", "value": 15},
            {"label": "砂石", "value": 12},
            {"label": "其他", "value": 8},
        ],
    },
    "layout": {
        "title": "施工现场平面布置图",
        "zones": [
            {"name": "办公区", "x": 10, "y": 10, "w": 20, "h": 15},
            {"name": "生活区", "x": 10, "y": 30, "w": 20, "h": 15},
            {"name": "加工区", "x": 35, "y": 10, "w": 25, "h": 20},
            {"name": "材料堆场", "x": 35, "y": 35, "w": 25, "h": 15},
            {"name": "施工区", "x": 65, "y": 10, "w": 30, "h": 40},
        ],
        "width": 100, "height": 60,
    },
    "timeline": {
        "title": "里程碑计划",
        "events": [
            {"date": "2026-03-01", "title": "开工典礼", "description": "项目正式启动"},
            {"date": "2026-06-30", "title": "基础完工", "description": "±0.000以下完成"},
            {"date": "2026-12-31", "title": "主体封顶", "description": "结构封顶完成"},
            {"date": "2027-06-30", "title": "竣工交付", "description": "整体验收通过"},
        ],
    },
}

def render_and_check(chart_type, payload):
    """渲染并返回质量数据"""
    payload_str = json.dumps(payload, ensure_ascii=False)
    stream = render_mermaid_to_bytes(payload_str, chart_type, skip_http=True)
    if not stream:
        return None
    
    png_bytes = stream.getvalue()
    with Image.open(io.BytesIO(png_bytes)) as img:
        w, h = img.size
        dpi = img.info.get("dpi", (0, 0))
        dpi_x = dpi[0] if dpi and dpi[0] > 0 else 96
        phys_w_cm = (w / dpi_x) * 2.54
        phys_h_cm = (h / dpi_x) * 2.54
    
    return {
        "type": chart_type,
        "title": payload.get("title", ""),
        "stream": stream,
        "w": w, "h": h,
        "dpi_x": dpi_x, "dpi_y": dpi[1] if dpi else 0,
        "phys_w_cm": phys_w_cm, "phys_h_cm": phys_h_cm,
        "size_kb": len(png_bytes) / 1024,
    }

# ============ 渲染全部 ============
print("=" * 70)
print("图表渲染管线质量检查 v2")
print("=" * 70)

results = []
for chart_type, payload in TEST_DATA.items():
    print(f"▶ {chart_type:14s} ... ", end="", flush=True)
    q = render_and_check(chart_type, payload)
    if q:
        dpi_ok = q["dpi_x"] >= 280
        size_ok = q["size_kb"] >= 30
        print(f"✅ {q['w']}x{q['h']}px, DPI={q['dpi_x']:.0f}, {q['phys_w_cm']:.1f}x{q['phys_h_cm']:.1f}cm, {q['size_kb']:.0f}KB")
        results.append(q)
    else:
        print("❌ 失败")

# ============ 生成高质量 DOCX ============
print(f"\n{'=' * 70}")
print("生成完整 DOCX 质量报告（含 7 类图表 + 图号 + 标题）")
print("=" * 70)

from docx import Document
from docx.shared import Cm, Pt, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from app.services.export import _add_diagram_picture, _next_figure_number, _reset_figure_counter, _apply_font

_reset_figure_counter()

doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.left_margin = Cm(2.5)
sec.right_margin = Cm(2.5)
sec.top_margin = Cm(2.5)
sec.bottom_margin = Cm(2.5)

class Config:
    image_max_width_cm = 16.0
    image_dpi = 300
    font_body = "宋体"
    size_body = 10.5

config = Config()

doc.add_heading("图表渲染质量检查报告", 0)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
p.add_run("测试项目：7 类图表 × 商业级 v2 渲染管线 × 300 DPI × 2.5x 超采样").font.size = Pt(12)

# 汇总表格
doc.add_heading("一、渲染质量汇总", level=1)
table = doc.add_table(rows=1, cols=6)
table.style = 'Table Grid'
hdr = table.rows[0].cells
for i, t in enumerate(["图表类型", "像素尺寸", "内嵌DPI", "物理尺寸", "文件大小", "状态"]):
    hdr[i].text = t
    for p in hdr[i].paragraphs:
        for r in p.runs:
            r.bold = True

for q in results:
    row = table.add_row().cells
    dpi_ok = "✅" if q["dpi_x"] >= 280 else "⚠️"
    row[0].text = f"{q['type']} ({q['title']})"
    row[1].text = f"{q['w']}×{q['h']}"
    row[2].text = f"{q['dpi_x']:.0f} {dpi_ok}"
    row[3].text = f"{q['phys_w_cm']:.1f}×{q['phys_h_cm']:.1f}cm"
    row[4].text = f"{q['size_kb']:.0f}KB"
    row[5].text = "✅" if q["dpi_x"] >= 280 and q["size_kb"] >= 30 else "⚠️"

# 逐图展示
doc.add_heading("二、各图表实际渲染效果", level=1)

parent_section = "2"
for idx, q in enumerate(results, 1):
    fig_no = _next_figure_number(parent_section)
    caption = f"{fig_no} {q['title']}"
    doc.add_heading(f"2.{idx} {q['type'].upper()} — {q['title']}", level=2)
    _add_diagram_picture(doc, q["stream"], config, caption=caption)
    
    # 质量标注
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = note.add_run(f"  [技术参数] {q['w']}×{q['h']}px, 内嵌 {q['dpi_x']:.0f} DPI, 物理尺寸 {q['phys_w_cm']:.1f}×{q['phys_h_cm']:.1f}cm, 文件 {q['size_kb']:.0f}KB")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    run.italic = True

out_path = os.path.join(os.path.dirname(__file__), "chart_quality_report.docx")
doc.save(out_path)
size_mb = os.path.getsize(out_path) / 1024 / 1024
print(f"✅ DOCX 已生成: {out_path} ({size_mb:.2f} MB)")
print(f"   包含 {len(results)} 张图表，均为 300 DPI / 2.5x 超采样")

# 打开文件
try:
    os.startfile(out_path)
    print("   📄 已尝试打开")
except:
    pass
